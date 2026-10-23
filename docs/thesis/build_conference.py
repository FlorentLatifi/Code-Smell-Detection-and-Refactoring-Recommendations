"""Prezantimi në anglisht për Konferencën Ndërkombëtare të UBT-së 2026, dhe fleta e pyetjeve.

    PYTHONPATH=backend python docs/thesis/build_conference.py

Shkruan te `docs/konferenca/`:
- `UBT2026_Florent_Latifi.pptx`: dymbëdhjetë sllajde që thuhen dhe një me referencat IEEE,
  me tekstin folës anglisht te shënimet (rreth dhjetë minuta me ritëm të qetë, tetë kur
  flitet shpejt), dhe sllajde rezervë për pyetjet;
- `UBT2026_Florent_Latifi_9min.pptx`: të njëjtat sllajde, me një tekst folës më të gjatë
  (rreth nëntë minuta kur flitet shpejt) për rastin kur kryesuesi jep vërtet dhjetë
  minuta;
- `UBT2026_pergatitja.docx`: koha për sllajd për të dy versionet, pyetjet e mundshme me
  përgjigje, shqiptimi i termave dhe lista për ditën e konferencës;
- `UBT2026_cka_te_flas.docx`: udhëzuesi i fjalimit, sllajd për sllajd: çka shihet në
  sllajd, ku të tregohet, dhe çdo fjali angleze me kuptimin e saj shqip;
- `Punimi_nga_e_para.docx`: gjithë punimi i shpjeguar si një libër, për një lexues që nuk
  di asgjë për programimin, me shembuj të matur nga vetë mjeti dhe me të njëjtat shifra si
  sllajdet;
- `UBT2026_teksti_qe_flas.docx` dhe `UBT2026_teksti_qe_flas_9min.docx`: vetëm fjalitë
  angleze që thuhen, sllajd për sllajd, për secilin version.

Dy versionet ndërtohen nga i njëjti tekst: i gjati është i shkurtri plus disa fjali
shtesë, që pjesa e provuar të mbetet e njëjtë dhe shtesat të mund të lihen jashtë në
sallë pa humbur fillin.

Si `build_slides.py`, çdo shifër lexohet nga `data/results/` me të njëjtët ngarkues si
punimi, që prezantimi, abstrakti i konferencës dhe punimi të mos thonë tre numra të
ndryshëm për të njëjtën gjë. Shembulli i kodit prodhohet nga motori përmes
`build_code_example.py`, ndaj kërkon korpusin; pa të, sllajdi e thotë.

Grafikët janë grafikë të PowerPoint-it e jo figurat e punimit: ato kanë etiketa shqip,
dhe një grafik vendas mund të ndryshohet në PowerPoint pa e rivizatuar.
"""

from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass

from build_slides import (
    ACCENT,
    CRITICAL,
    FIGURES,
    HEIGHT,
    INK,
    MARGIN,
    SOFT,
    WIDTH,
    _blank,
    _mean,
    _note,
)
from chapters import (
    MADEYSKI_2023_DS1,
    REPOSITORY,
    SMELL_SQ,
    _load,
    _repositories_in_dataset,
    _rows_if_present,
    _significance,
)
from docx import Document
from docx.shared import Cm as DocCm
from docx.shared import Pt as DocPt
from docx.shared import RGBColor as DocRGB
from javasmell import glossary
from javasmell.analysis import analyze_source
from javasmell.detectors.rules import detect_all
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(__file__)
OUT_DIR = os.path.join(HERE, "..", "konferenca")
DECK = os.path.join(OUT_DIR, "UBT2026_Florent_Latifi.pptx")
DECK_LONG = os.path.join(OUT_DIR, "UBT2026_Florent_Latifi_9min.pptx")
SHEET = os.path.join(OUT_DIR, "UBT2026_pergatitja.docx")
SPEECH = os.path.join(OUT_DIR, "UBT2026_cka_te_flas.docx")
BOOK = os.path.join(OUT_DIR, "Punimi_nga_e_para.docx")
QUESTIONS = os.path.join(OUT_DIR, "UBT2026_pyetjet_e_shpjeguara.docx")
SPOKEN = os.path.join(OUT_DIR, "UBT2026_teksti_qe_flas.docx")
SPOKEN_LONG = os.path.join(OUT_DIR, "UBT2026_teksti_qe_flas_9min.docx")

TITLE = (
    "Automated Code Smell Detection and Refactoring Recommendations "
    "Using Software Metrics and Machine Learning"
)
AUTHORS = "Florent Latifi  ·  Altina Salihu"
AFFILIATION = "UBT – Higher Education Institution, Kosovo"
# Emri siç e shkruan agjenda e konferencës: email-i dhe faqja e saj përdorin shkurtesa
# të ndryshme, dhe agjenda është dokumenti që ka salla.
EVENT = "15th UBT Annual International Conference on Computer Science and Engineering  ·  2026"
# Nga agjenda paraprake; salla nuk shkruhet, sepse agjenda përfundimtare mund ta ndryshojë.
SESSION = "Session 1  ·  Software Engineering  ·  24 October 2026"

# Gri për rregullat, blu për modelin: e njëjta ndarje si te figurat e punimit, që
# lexohet edhe kur projektori i zbeh ngjyrat. Gria është e errët sa të kalojë kontrastin
# 3:1 që i duhet një grafiku kundrejt të bardhës.
MUTED = RGBColor(0x87, 0x91, 0x9C)
ADDED = RGBColor(0x2F, 0x6F, 0x45)
MONO = "Consolas"

MODEL_EN = {
    "majority": "Majority",
    "logistic": "Logistic regression",
    "random_forest": "Random forest",
    "gradient_boosting": "Gradient boosting",
}

# Ritmi i një folësi jo-amtar që flet qetë. Koha e çdo sllajdi llogaritet nga fjalët e
# shënimit, që fleta e kohës të ndjekë tekstin e jo një vlerësim me sy.
WORDS_PER_MINUTE = 115
# Autori flet shpejt, dhe nervozizmi në sallë e shpejton më tej. Ky është një supozim,
# jo një matje: fleta i thotë ta zëvendësojë me ritmin e vet, të matur me kronometër.
FAST_WORDS_PER_MINUTE = 145

# Rreshtat e krahasimit me PMD-në, siç i emërton sllajdi rezervë dhe dokumenti i njohurive.
PMD_NAMES = {
    "blob/strategy": "Blob (strategy)",
    "blob/with_size": "Blob (with size)",
    "data class/strategy": "Data Class",
    "long method/strategy": "Long Method",
    "feature envy/law_of_demeter": "Feature Envy (LawOfDemeter)",
}

# Emrat e refaktorimeve siç thuhen me gojë, jo siç janë në kod.
SPOKEN_REFACTORING = {"ReplaceNestedConditionalWithGuardClauses": "Guard Clauses"}

# Burimet që citon prezantimi, në stilin IEEE që kërkoi mentorja për punimin (VD-152). Teksti
# është ai i listës së punimit fjalë për fjalë (`references.py` në degën
# claude/hungry-lederberg-3b72ba); vetëm numri ndryshon, sepse IEEE numëron sipas radhës së
# citimit të parë, dhe këtu ajo radhë është e sllajdeve, jo e kapitujve.
DECK_SOURCES = [
    ("Fowler 2018",
     "M. Fowler, *Refactoring: Improving the Design of Existing Code*, 2nd ed. Boston, MA, USA: "
     "Addison-Wesley, 2018."),
    ("Lanza & Marinescu 2006",
     "M. Lanza and R. Marinescu, *Object-Oriented Metrics in Practice: Using Software Metrics to "
     "Characterize, Evaluate, and Improve the Design of Object-Oriented Systems*. Berlin, Germany: "
     "Springer, 2006, doi: 10.1007/3-540-39538-5."),
    ("Madeyski & Lewowski 2020",
     "L. Madeyski and T. Lewowski, “MLCQ: Industry-relevant code smell data set,” in *Proceedings "
     "of the 24th International Conference on Evaluation and Assessment in Software Engineering "
     "(EASE ’20)*, 2020, pp. 342–347, doi: 10.1145/3383219.3383264."),
    ("Arcelli Fontana et al. 2016",
     "F. Arcelli Fontana, M. V. Mäntylä, M. Zanoni, and A. Marino, “Comparing and experimenting "
     "machine learning techniques for code smell detection,” *Empirical Software Engineering*, "
     "vol. 21, no. 3, pp. 1143–1191, 2016, doi: 10.1007/s10664-015-9378-4."),
    ("Di Nucci et al. 2018",
     "D. Di Nucci, F. Palomba, D. A. Tamburri, A. Serebrenik, and A. De Lucia, “Detecting code "
     "smells using machine learning techniques: Are we there yet?” in *Proceedings of the 25th "
     "IEEE International Conference on Software Analysis, Evolution and Reengineering (SANER "
     "’18)*, 2018, pp. 612–621, doi: 10.1109/SANER.2018.8330266."),
    ("Tsantalis & Chatzigeorgiou 2009",
     "N. Tsantalis and A. Chatzigeorgiou, “Identification of Move Method refactoring "
     "opportunities,” *IEEE Transactions on Software Engineering*, vol. 35, no. 3, pp. 347–367, "
     "2009, doi: 10.1109/TSE.2009.1."),
    ("Opdyke 1992",
     "W. F. Opdyke, “Refactoring object-oriented frameworks,” Ph.D. dissertation, University of "
     "Illinois at Urbana-Champaign, Urbana, IL, USA, 1992."),
    ("Chidamber & Kemerer 1994",
     "S. R. Chidamber and C. F. Kemerer, “A metrics suite for object oriented design,” *IEEE "
     "Transactions on Software Engineering*, vol. 20, no. 6, pp. 476–493, 1994, doi: "
     "10.1109/32.295895."),
    ("Chicco & Jurman 2020",
     "D. Chicco and G. Jurman, “The advantages of the Matthews correlation coefficient (MCC) over "
     "F1 score and accuracy in binary classification evaluation,” *BMC Genomics*, vol. 21, no. 1, "
     "Art. no. 6, 2020, doi: 10.1186/s12864-019-6413-7."),
    ("PMD Team 2026",
     "PMD Team, *PMD: An Extensible Cross-Language Static Code Analyzer*, version 7.27.0. (2026). "
     "[Online]. Available: https://pmd.github.io/"),
    ("Madeyski & Lewowski 2023",
     "L. Madeyski and T. Lewowski, “Detecting code smells using industry-relevant data,” "
     "*Information and Software Technology*, vol. 155, Art. no. 107112, 2023, doi: "
     "10.1016/j.infsof.2022.107112."),
]  # fmt: skip


def _cite(*labels: str) -> str:
    """`_cite("Opdyke 1992")` → «[7]»: numri vjen nga radha e listës, jo i shkruar me dorë."""
    order = [label for label, _ in DECK_SOURCES]
    return ", ".join(f"[{order.index(label) + 1}]" for label in labels)


class Line(str):
    """Një fjali e tekstit folës, me kuptimin e saj shqip.

    Përkthimi nuk thuhet në sallë. Udhëzuesi i fjalimit e vë poshtë fjalisë angleze, që
    folësi ta dijë çka po thotë dhe të mos mësojë përmendësh tinguj pa kuptim. Fjalia dhe
    përkthimi shkruhen bashkë, që një shifër e ndryshuar të mos lërë njërën pas.
    """

    sq: str

    def __new__(cls, text: str, sq: str) -> Line:
        line = super().__new__(cls, text)
        line.sq = sq
        return line


class Extra(Line):
    """Një fjali që thuhet vetëm në versionin nëntë-minutësh.

    Shtesat janë pjesë e të njëjtit tekst e jo një tekst i dytë, që versioni i gjatë të
    mos i ndryshojë fjalitë e provuara: në sallë, kur koha mbaron, ato thjesht lihen.
    """


def _spoken(parts: list[str], long: bool) -> list[str]:
    return [part for part in parts if long or not isinstance(part, Extra)]


def _word_count(parts: list[str]) -> int:
    return len(" ".join(parts).split())


def _clock(seconds: int) -> str:
    return f"{seconds // 60}:{seconds % 60:02d}"


def _words(name: str) -> str:
    """`LongParameterList` → `Long Parameter List`, siç e thotë folësi."""
    return SPOKEN_REFACTORING.get(name, re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name))


# ----------------------------------------------------------------------
# Të dhënat
# ----------------------------------------------------------------------
def _facts() -> dict:
    """Çdo shifër që thotë prezantimi, e lexuar një herë."""
    rules = _load("rules_evaluation.json")
    ml = _load("ml_evaluation.json")
    refactoring = _load("refactoring_evaluation.json")
    quality = _load("rewrite_quality.json")
    agreement = _load("reviewer_agreement.json")
    dataset = _load("mlcq_dataset.json")
    reference = _load("system_reference.json")
    pmd = _load("pmd_comparison.json")
    smells = sorted(ml["per_smell"])
    unsure, _ = _significance(smells)
    best = {s: ml["per_smell"][s]["models"][ml["per_smell"][s]["best_model"]] for s in smells}
    strategy = {s: _mean(rules, s) for s in smells}
    reviewers = {s: agreement["per_smell"][s]["mcc"] for s in smells}
    applied, detected = refactoring["applied"], refactoring["detected"]
    # Krahasimi me PMD-në: ndryshimi është yni minus i PMD-së, dhe vetëm një interval që
    # e përjashton zeron lejon të thuhet se njëri është përpara.
    differences = {
        key: entry["difference"] for key, entry in pmd["by_smell"].items() if "difference" in entry
    }
    by_refactoring = quality["by_refactoring"]
    return {
        "rules": rules,
        "ml": ml,
        "refactoring": refactoring,
        "quality": quality,
        "agreement": agreement,
        "dataset": dataset,
        "reference": reference,
        "pmd": pmd,
        "smells": smells,
        "names": [SMELL_SQ[s] for s in smells],
        "strategy": strategy,
        "best": best,
        "sure": [SMELL_SQ[s] for s in smells if SMELL_SQ[s] not in unsure],
        "unsure": unsure,
        "samples": dataset["rows"],
        "repositories": _repositories_in_dataset(),
        "metrics": len(reference["metrics"]["class"]) + len(reference["metrics"]["method"]),
        "strategies": len(reference["strategies"]),
        "reviewers": reviewers,
        "recall": [strategy[s]["recall"] for s in smells],
        "model_mcc": [best[s]["mcc"] for s in smells],
        "rule_mcc": [strategy[s]["mcc"] for s in smells],
        "applied": applied,
        "detected": detected,
        "resolved": refactoring["resolution"]["resolved"] / applied,
        "broken": refactoring["verdicts"].get("new_errors", 0) / applied,
        "refused": sum(refactoring["refused_by_reason"].values()) / detected,
        "moderate": [s for s in smells if reviewers[s] > 0.1],
        "chance": [s for s in smells if reviewers[s] <= 0.1],
        "pmd_ahead": [SMELL_SQ[key.split("/")[0]] for key, d in differences.items()
                      if d["excludes_zero"] and d["median"] < 0],
        "pmd_behind": [SMELL_SQ[key.split("/")[0]] for key, d in differences.items()
                       if key.endswith("/strategy") and d["excludes_zero"] and d["median"] > 0],
        "introduced": refactoring["introduced_smells"],
        # Sa herë u rillogarit ndryshimi model − rregulla; e njëjtë për çdo erë.
        "resamples": _load("bootstrap_intervals.json")["per_smell"][smells[0]]["intervals"][
            "difference_model_minus_rules"
        ]["resamples_used"],
        # Sa erëra fiton edhe algoritmi më i dobët kundër rregullave: përgjigjja kur dikush
        # pyet nëse zgjedhja e më të mirit e fryn rezultatin.
        "weakest": sum(
            min(e["mcc"] for e in ml["per_smell"][s]["models"].values() if e["mcc"] is not None)
            > strategy[s]["mcc"]
            for s in smells
        ),
        "always": [_words(name) for name, entry in by_refactoring.items() if entry["acceptable"] == 1.0],
        "never": [_words(name) for name, entry in by_refactoring.items() if entry["acceptable"] == 0.0],
    }  # fmt: skip


def _and(items: list[str], word: str = "and") -> str:
    if len(items) < 2:
        return "".join(items)
    return ", ".join(items[:-1]) + f" {word} " + items[-1]


def _dhe(items: list[str]) -> str:
    return _and(items, "dhe")


def _pct(value: float, digits: int = 1) -> str:
    return f"{value * 100:.{digits}f}%"


# ----------------------------------------------------------------------
# Elementet e sllajdeve që `build_slides.py` nuk i ka
# ----------------------------------------------------------------------
def _speech(slide, parts: list[str]) -> None:
    """Shënimi i folësit, me fjalitë shtesë të pjerrëta.

    Pamja e folësit në PowerPoint e ruan stilin, ndaj në sallë dallohet me një shikim cila
    fjali mund të lihet kur koha nuk mjafton.
    """
    paragraph = slide.notes_slide.notes_text_frame.paragraphs[0]
    for index, part in enumerate(parts):
        run = paragraph.add_run()
        run.text = part if index == 0 else " " + part
        run.font.italic = isinstance(part, Extra)


def _bar_chart(slide, categories, series, top, height, minimum=0.0, maximum=1.0, left=MARGIN,
               width=None):  # fmt: skip
    """Shtylla të grupuara me vlerën mbi secilën, pa rrjetë dhe pa bosht të tepërt."""
    data = CategoryChartData()
    data.categories = categories
    for name, values, _ in series:
        data.add_series(name, values, number_format="0.00")
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width or WIDTH - 2 * MARGIN, height, data
    )
    chart = frame.chart
    # Me një seri të vetme PowerPoint-i i vë grafikut emrin e serisë si titull.
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(16)
    chart.font.color.rgb = INK
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(16)
    axis = chart.value_axis
    axis.minimum_scale = minimum
    axis.maximum_scale = maximum
    axis.has_major_gridlines = False
    axis.visible = False
    chart.category_axis.tick_labels.font.size = Pt(16)
    chart.category_axis.has_major_gridlines = False
    chart.category_axis.format.line.color.rgb = FRAME
    # Emrat e erërave poshtë grafikut, jo te boshti zero, ku një shtyllë negative i mbulon.
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = "0.00"
    labels.number_format_is_linked = False
    labels.font.size = Pt(16)
    labels.font.name = FONT_STRONG
    for index, (_, _, colour) in enumerate(series):
        plot.series[index].invert_if_negative = False
        fill = plot.series[index].format.fill
        fill.solid()
        fill.fore_color.rgb = colour
    return chart


def _elided(lines: list[str], marked: set[int]) -> tuple[list[str], set[int]]:
    """Blloku synchronized mblidhet në një rresht me një koment që thotë sa rreshta fsheh.

    Ai pjesë nuk ndryshon nga rishkrimi, dhe e plotë nuk hyn në një sllajd me shkronja
    që lexohen nga salla. Shkurtimi shënohet në vetë kodin, që askush të mos e lexojë
    si kodin e plotë.
    """
    start = next(i for i, line in enumerate(lines) if line.strip().startswith("synchronized"))
    depth, end = 0, start
    for end in range(start, len(lines)):
        depth += lines[end].count("{") - lines[end].count("}")
        if depth == 0:
            break
    inner = lines[start + 1 : end]
    folded = f"{lines[start].rstrip()} /* … {len(inner)} unchanged lines */ {lines[end].strip()}"
    shown = lines[:start] + [folded] + lines[end + 1 :]
    kept = {i if i < start else start if i <= end else i - (end - start) for i in marked}
    return shown, kept


def _example() -> tuple[list[str], set[int], list[str], set[int]] | None:
    """Rishkrimi i vërtetë i figurës së Nënkapitullit 4.7, ose None pa korpus."""
    import build_code_example as example

    if not example.SOURCE.exists():
        return None
    before, after = example._rewritten()
    removed, added = example._changed(before, after)
    before, removed = _elided(before, removed)
    after, added = _elided(after, added)
    return before, removed, after, added


def _script(f: dict) -> list[tuple]:
    """Teksti folës i sllajdeve kryesore, me radhë: emri i brendshëm dhe fjalitë.

    I ndarë nga paraqitja e sllajdeve, që një sllajd të rindërtohet pa e prekur tekstin
    që folësi e ka provuar. Fjalitë janë të shkurtra dhe me fjalë të përditshme: salla ka
    studentë e pedagogë nga fusha të tjera, dhe folësi nuk është anglishtfolës amtar.
    """
    refactoring, quality = f["refactoring"], f["quality"]
    applied, detected = f["applied"], f["detected"]
    moderate, chance = f["moderate"], f["chance"]
    pmd_ahead = f["pmd_ahead"]
    lowest, highest = min(f["recall"]) * 100, max(f["recall"]) * 100
    found_low, found_high = (bound(f["best"][s]["recall"] for s in f["smells"]) * 100
                             for bound in (min, max))  # fmt: skip
    shape = refactoring["refused_by_reason"]["shape_not_matched"] / detected * 100
    majority = f["ml"]["per_smell"]["feature envy"]["models"]["majority"]["accuracy"] * 100
    limit = f["reference"]["thresholds"]["long_method_loc"]
    resolved = f["resolved"] * 100
    acceptable = quality["acceptable_reweighted"] * 100
    refused, missing = refactoring["refused"], refactoring["unlocatable"]
    verdicts = refactoring["verdicts"]
    new_errors = verdicts.get("new_errors", 0)
    passed = (verdicts.get("compiles", 0) + verdicts.get("no_new_errors", 0)) / applied
    verify = _load("verify_with_project.json")
    in_project = verify["compiled_in_project"]
    # Ato që kalonin të vetme dhe dështojnë brenda projektit: gabimi që izolimi e fshihte.
    hidden = verify["moved"].get("no_new_errors -> new_errors", 0)
    changed = sum(e["behaviour"]["changed"] for e in quality["by_refactoring"].values())
    pmd_sentence = (
        "It was never better than my rules" if not pmd_ahead
        else f"It was better for {_and(pmd_ahead)}"
    )  # fmt: skip
    pmd_sq = "Nuk doli kurrë më i mirë se rregullat e mia" if not pmd_ahead else (
        f"Doli më i mirë për {_dhe(pmd_ahead)}")  # fmt: skip
    # Çdo sllajd nis me një fjali që e lidh me të mëparshmin, që salla ta ndjekë historinë
    # e jo një listë sllajdesh; fjalitë mbeten të shkurtra, për një folës jo-amtar.
    return [
        (
            "Title",
            Line("Good morning, everyone.", "Mirëmëngjes të gjithëve."),
            Line("My name is Florent Latifi. This is my bachelor thesis at UBT, with my supervisor "
                 "Altina Salihu.",
                 "Unë quhem Florent Latifi. Kjo është teza ime e bachelorit në UBT, me mentoren "
                 "time Altina Salihu."),
            Line("Today I want to show you three things: how a tool can find design problems in "
                 "Java code, how accurate it is, and whether it can also fix them.",
                 "Sot dua t'ju tregoj tri gjëra: si mund t'i gjejë një mjet problemet e ndërtimit "
                 "në kodin Java, sa i saktë është, dhe a mund t'i ndreqë edhe vetë."),
        ),
        (
            "The problem",
            Line("Let me start with a problem that every developer knows.",
                 "Po filloj me një problem që e njeh çdo programues."),
            Line("There is always one class that nobody wants to touch.",
                 "Gjithmonë ka një klasë që askush nuk do ta prekë."),
            Line("The program works, but every small change takes too long.",
                 "Programi punon, por çdo ndryshim i vogël zgjat shumë."),
            Line("We call these design problems code smells.",
                 "Këto probleme të ndërtimit i quajmë erëra të kodit."),
            Line("They are not bugs. They just make the code expensive to change.",
                 "Nuk janë gabime. Thjesht e bëjnë kodin të shtrenjtë për t'u ndryshuar."),
            Line(f"Most tools find them with fixed numbers. For example: if a method has more than "
                 f"{limit:g} lines, it is a Long Method.",
                 f"Shumica e mjeteve i gjejnë me numra fiksë. Për shembull: nëse një metodë ka më "
                 f"shumë se {limit:g} rreshta, ajo është Long Method."),
            Line("But these numbers come from other projects, and the tools only point at the "
                 "problem. They do not fix it.",
                 "Por këta numra vijnë nga projekte të tjera, dhe mjetet vetëm e tregojnë "
                 "problemin. Nuk e ndreqin."),
            Line("So I asked three questions.", "Prandaj bëra tri pyetje."),
            Line("One: how accurate are the published rules?",
                 "Një: sa të sakta janë rregullat e botuara?"),
            Line("Two: can machine learning do better?",
                 "Dy: a mund të bëjë më mirë mësimi makinerik?"),
            Line("Three: can a tool fix these smells automatically, without breaking the code?",
                 "Tre: a mund t'i ndreqë një mjet këto erëra automatikisht, pa e prishur kodin?"),
        ),
        (
            "Contribution",
            Line("Other researchers have worked on these questions before. So what is new in my "
                 "work? Three things.",
                 "Studiues të tjerë janë marrë me këto pyetje edhe më parë. Atëherë, çka është e re "
                 "në punën time? Tri gjëra."),
            Line("First, earlier studies take the published rules as correct. I measure them "
                 "against professional developers.",
                 "E para, studimet e mëparshme i marrin rregullat e botuara si të sakta. Unë i mat "
                 "kundrejt programuesve profesionistë."),
            Line("Second, I compare rules and machine learning in a fair way: on the same samples, "
                 "with the same score, and with a split by project.",
                 "E dyta, i krahasoj rregullat dhe mësimin makinerik në mënyrë të drejtë: mbi të "
                 "njëjtat mostra, me të njëjtën notë, dhe me ndarje sipas projektit."),
            Line("Third, earlier tools only suggest a fix. Mine finds a smell, fixes it, checks it "
                 "with the compiler, and measures it again.",
                 "E treta, mjetet e mëparshme vetëm sugjerojnë një ndreqje. Imi e gjen erën, e "
                 "ndreq, e kontrollon me kompajlerin, dhe e mat sërish."),
        ),
        (
            "Four smells",
            Line("Before the results, let me show you the four smells I studied.",
                 "Para rezultateve, po ju tregoj katër erërat që studiova."),
            Line("A Blob is a class that does too much.", "Blob është një klasë që bën shumë gjëra."),
            Line("A Data Class only holds data, and does almost nothing with it.",
                 "Data Class vetëm mban të dhëna, dhe nuk bën pothuajse asgjë me to."),
            Line("A Long Method is a method that has grown too long.",
                 "Long Method është një metodë që është zgjatur shumë."),
            Line("And Feature Envy is a method that works more with another class than with its "
                 "own.",
                 "Dhe Feature Envy është një metodë që punon më shumë me një klasë tjetër se me "
                 "klasën e vet."),
        ),
        (
            "What we built",
            Line("To find these smells, and to fix them, I built a tool called JavaSmell.",
                 "Për t'i gjetur këto erëra, dhe për t'i ndrequr, ndërtova një mjet që quhet "
                 "JavaSmell."),
            Line(f"First, it reads the Java code and measures {f['metrics']} numbers for every "
                 "class and method, like the number of lines, or how complex the code is.",
                 f"Së pari, e lexon kodin Java dhe mat {f['metrics']} numra për çdo klasë dhe "
                 "metodë, si numri i rreshtave, ose sa i ndërlikuar është kodi."),
            Line("Then it uses these numbers in three ways.",
                 "Pastaj këta numra i përdor në tri mënyra."),
            Line(f"Part A uses {f['strategies']} rules from well-known books and papers.",
                 f"Pjesa A përdor {f['strategies']} rregulla nga libra dhe artikuj të njohur."),
            Line("Part B uses machine learning on the same numbers.",
                 "Pjesa B përdor mësimin makinerik mbi të njëjtët numra."),
            Line("And Part C fixes the code, but only when it can prove that the change is safe.",
                 "Dhe pjesa C e ndreq kodin, por vetëm kur mund ta vërtetojë që ndryshimi është i "
                 "sigurt."),
            Extra(f"Every finding also says why. For example: this method has 53 lines, and the "
                  f"limit is {limit:g}.",
                  f"Çdo gjetje thotë edhe pse. Për shembull: kjo metodë ka 53 rreshta, dhe kufiri "
                  f"është {limit:g}."),
            Line("On the right you see the web interface. It explains everything in plain "
                 "Albanian, so you do not need to be an expert.",
                 "Djathtas shihni ndërfaqen web. Gjithçka e shpjegon në shqip të thjeshtë, ndaj "
                 "nuk ke nevojë të jesh ekspert."),
        ),
        (
            "Data and evaluation",
            Line("But a tool is only useful if its answers are right. To check that, I need "
                 "answers from real people.",
                 "Por një mjet vlen vetëm nëse përgjigjet e tij janë të sakta. Për ta kontrolluar "
                 "këtë, më duhen përgjigje nga njerëz të vërtetë."),
            Line("So I used MLCQ, a public dataset where professional developers judged real Java "
                 "code.",
                 "Prandaj përdora MLCQ, një bazë publike të dhënash ku programues profesionistë "
                 "gjykuan kod të vërtetë Java."),
            Line(f"In total, I have {f['samples']:,} pieces of code from {f['repositories']} "
                 "open-source projects.",
                 f"Gjithsej kam {f['samples']:,} copa kodi nga {f['repositories']} projekte me "
                 "burim të hapur."),
            Line("And the test is fair: a model is never tested on a project it has already seen. "
                 "It learns from four groups of projects, and it is tested on the fifth.",
                 "Dhe prova është e drejtë: një model nuk provohet kurrë mbi një projekt që e ka "
                 "parë. Mëson nga katër grupe projektesh, dhe provohet te i pesti."),
            Extra("This matters, because code from the same project looks very similar. Testing "
                  "on a project it has seen would be like an exam where you already know the "
                  "answers.",
                  "Kjo ka rëndësi, sepse kodi i të njëjtit projekt ngjan shumë. Provimi mbi një "
                  "projekt të parë do të ishte si një provim ku i di përgjigjet që më parë."),
            Line("To score the results, I use MCC, not simple accuracy. Here is why.",
                 "Për t'i vlerësuar rezultatet, përdor MCC-në, jo saktësinë e thjeshtë. Ja pse."),
            Line(f"On Feature Envy, a detector that always says “clean” is {majority:.0f} percent "
                 "accurate, but it is useless. Its MCC is zero.",
                 f"Te Feature Envy, një detektor që thotë gjithmonë «i pastër» është {majority:.0f} "
                 "për qind i saktë, por është i padobishëm. MCC-ja e tij është zero."),
        ),
        (
            "Rules vs. machine learning",
            Line("So, how good are the rules, and can machine learning do better? Here are the "
                 "main results.",
                 "Pra, sa të mira janë rregullat, dhe a bën më mirë mësimi makinerik? Ja "
                 "rezultatet kryesore."),
            Line("Gray is the published rules, blue is machine learning.",
                 "Gria janë rregullat e botuara, bluja është mësimi makinerik."),
            Line("Higher is better: zero means guessing, and one means perfect.",
                 "Sa më lart, aq më mirë: zero do të thotë hamendje, dhe një do të thotë "
                 "përsosmëri."),
            Line(f"The rules are usually right when they raise an alarm. But out of every 100 real "
                 f"smells, they find only {lowest:.0f} to {highest:.0f}.",
                 f"Rregullat zakonisht kanë të drejtë kur alarmojnë. Por nga çdo 100 erëra të "
                 f"vërteta, gjejnë vetëm {lowest:.0f} deri {highest:.0f}."),
            Line(f"Machine learning is better on all four smells. It finds {found_low:.0f} to "
                 f"{found_high:.0f} out of 100, but usually with more false alarms.",
                 f"Mësimi makinerik është më i mirë te të katër erërat. Gjen {found_low:.0f} deri "
                 f"{found_high:.0f} nga 100, por zakonisht me më shumë alarme të rreme."),
            Line(f"For {_and(f['sure'])}, the improvement is real, not luck.",
                 f"Për {_dhe(f['sure'])}, përmirësimi është i vërtetë, jo fat."),
            Line(f"For {_and(f['unsure'])}, it could be luck.",
                 f"Për {_dhe(f['unsure'])}, mund të jetë fat."),
            Extra(f"I checked this by repeating the calculation {f['resamples']:,} times on random "
                  "samples of the projects.",
                  f"Këtë e kontrollova duke e përsëritur llogaritjen {f['resamples']:,} herë mbi "
                  "mostra të rastësishme të projekteve."),
            Extra(f"I also compared with PMD, a popular free tool: {pmd_sentence[0].lower()}"
                  f"{pmd_sentence[1:]}.",
                  f"E krahasova edhe me PMD-në, një mjet i njohur falas: {pmd_sq[0].lower()}"
                  f"{pmd_sq[1:]}."),
        ),
        (
            "Human agreement",
            Line(f"You may ask: why is {_and([SMELL_SQ[s] for s in chance])} so hard, even for "
                 "machine learning?",
                 f"Mund të pyesni: pse është {_dhe([SMELL_SQ[s] for s in chance])} kaq e vështirë, "
                 "edhe për mësimin makinerik?"),
            Line("To find out, I looked at code that two developers judged, and checked how often "
                 "they agree.",
                 "Për ta gjetur, shikova kodin që e gjykuan dy programues, dhe kontrollova sa shpesh "
                 "pajtohen."),
            Line(f"For {_and([SMELL_SQ[s] for s in moderate])}, they agree only partly.",
                 f"Për {_dhe([SMELL_SQ[s] for s in moderate])}, pajtohen vetëm pjesërisht."),
            Line(f"For {_and([SMELL_SQ[s] for s in chance])}, they agree no more than a coin flip.",
                 f"Për {_dhe([SMELL_SQ[s] for s in chance])}, pajtohen jo më shumë se një monedhë "
                 "e hedhur."),
            *[
                Extra(f"Their score is {f['reviewers'][s]:.2f}, practically zero.",
                      f"Nota e tyre është {f['reviewers'][s]:.2f}, praktikisht zero.")
                for s in chance
            ],
            Line("So the problem is not only the algorithm.",
                 "Pra problemi nuk është vetëm algoritmi."),
            Line("If experts do not agree on what a smell is, no tool can learn it well.",
                 "Nëse ekspertët nuk pajtohen se çka është një erë, asnjë mjet nuk mund ta mësojë "
                 "mirë."),
        ),
        (
            "Automatic refactoring",
            Line("So much for finding smells. Now the third question: once we find a smell, can "
                 "the tool fix it?",
                 "Kaq për gjetjen e erërave. Tani pyetja e tretë: kur e gjejmë një erë, a mund ta "
                 "ndreqë mjeti?"),
            Line("The tool knows three fixes: it can split a long method, turn a big if into an "
                 "early exit, or group many parameters into one object.",
                 "Mjeti di tri ndreqje: mund ta ndajë një metodë të gjatë, ta kthejë një if të "
                 "madh në një dalje të hershme, ose t'i mbledhë shumë parametra në një objekt."),
            Line(f"In all the projects, it found {detected:,} places where a fix could apply.",
                 f"Në të gjitha projektet, gjeti {detected:,} vende ku mund të zbatohej një "
                 "ndreqje."),
            Line(f"It changed {applied:,} of them, about {applied / detected * 100:.0f} percent.",
                 f"Ndryshoi {applied:,} prej tyre, rreth {applied / detected * 100:.0f} për qind."),
            Line(f"It refused {refused:,}, because it could not prove that the change was safe. "
                 "And it always says why.",
                 f"Refuzoi {refused:,}, sepse nuk mundi ta vërtetojë që ndryshimi ishte i sigurt. "
                 "Dhe gjithmonë thotë pse."),
            Line(f"The last {missing} could not be found again in the file. Together, that makes "
                 f"{detected:,}.",
                 f"{missing} të fundit nuk u gjetën dot sërish në skedar. Bashkë, bëjnë "
                 f"{detected:,}."),
            Line("For me, a refusal is a good result, not a failure. It is better to refuse than "
                 "to break someone's code.",
                 "Për mua, refuzimi është rezultat i mirë, jo dështim. Më mirë të refuzosh, se sa "
                 "ta prishësh kodin e dikujt."),
            Extra(f"In {shape:.0f} percent of the places, the code simply does not have the right "
                  "shape for the fix.",
                  f"Në {shape:.0f} për qind të vendeve, kodi thjesht nuk ka formën e duhur për "
                  "ndreqjen."),
        ),
        (
            "Validation",
            Line("But a rewrite is only useful if it is correct. So how do I check that? Step by "
                 "step.",
                 "Por një rishkrim vlen vetëm nëse është i saktë. Atëherë, si e kontrolloj? Hap pas "
                 "hapi."),
            Line("First, before the change, the tool must prove its conditions on the syntax "
                 "tree. If it cannot, it refuses.",
                 "Së pari, para ndryshimit, mjeti duhet t'i vërtetojë kushtet e veta mbi pemën e "
                 "sintaksës. Nëse nuk mundet, refuzon."),
            Line(f"Second, after the change, the compiler checks the file. {passed * 100:.0f} "
                 f"percent pass, and the {new_errors} that fail are never written.",
                 f"Së dyti, pas ndryshimit, kompajleri e kontrollon skedarin. {passed * 100:.0f} për "
                 f"qind kalojnë, dhe ato {new_errors} që dështojnë nuk shkruhen kurrë."),
            Line(f"Third, the smell is measured again. It is gone after {resolved:.0f} percent of "
                 "the rewrites.",
                 f"Së treti, era matet sërish. Zhduket pas {resolved:.0f} për qind të "
                 "rishkrimeve."),
            Extra(f"I also took a sample of {verify['rewrites']} rewrites and compiled the whole "
                  f"project: {in_project['compiles']} compiled fully, and {hidden} had an error "
                  "that the single file had hidden.",
                  f"Mora edhe një mostër prej {verify['rewrites']} rishkrimesh dhe kompilova gjithë "
                  f"projektin: {in_project['compiles']} u kompiluan plotësisht, dhe {hidden} "
                  "kishin një gabim që skedari i vetëm e kishte fshehur."),
            Line(f"Finally, I read {quality['reviewed']} rewrites by hand. About {acceptable:.0f} "
                 f"percent were acceptable, and {_count_word(changed, quality['reviewed'])} changed "
                 "the behavior.",
                 f"Në fund, lexova {quality['reviewed']} rishkrime me dorë. Rreth "
                 f"{acceptable:.0f} për qind ishin të pranueshme, dhe {changed} e ndryshonte "
                 "sjelljen."),
            Line("So, to be clear: compiling does not guarantee the same behavior, because I could "
                 "not run the projects' own tests.",
                 "Pra, që të jetë e qartë: kompilimi nuk e garanton sjelljen e njëjtë, sepse nuk "
                 "munda t'i ekzekutoj testet e vetë projekteve."),
        ),
        (
            "Example",
            Line("Let me show you what one of these rewrites looks like. This is real code, from "
                 "the Apache Geode project.",
                 "Po ju tregoj si duket një nga këto rishkrime. Ky është kod i vërtetë, nga "
                 "projekti Apache Geode."),
            Line("Before: the whole method is inside one big if, so all the code is pushed to the "
                 "right.",
                 "Më parë: e gjithë metoda është brenda një if-i të madh, ndaj i gjithë kodi "
                 "shtyhet djathtas."),
            Line("After: the condition becomes an early return, and the rest of the code moves one "
                 "level up.",
                 "Më pas: kushti bëhet një dalje e hershme, dhe pjesa tjetër e kodit ngjitet një "
                 "nivel lart."),
            Extra("The tool only does this when the method returns nothing, and its whole body is "
                  "one if, without else.",
                  "Mjeti e bën këtë vetëm kur metoda nuk kthen asgjë, dhe i gjithë trupi i saj "
                  "është një if, pa else."),
            Line("The result compiles, and in my review it was acceptable without any change.",
                 "Rezultati kompilohet, dhe në rishikimin tim ishte i pranueshëm pa asnjë "
                 "ndryshim."),
        ),
        (
            "Conclusions",
            Line("Let me bring it all together. My three questions have three answers.",
                 "Po i bashkoj të gjitha. Tri pyetjet e mia kanë tri përgjigje."),
            Line("One: the published rules are precise, but they miss most smells.",
                 "Një: rregullat e botuara janë të sakta, por i humbasin shumicën e erërave."),
            Line("Two: machine learning on the same numbers does better.",
                 "Dy: mësimi makinerik mbi të njëjtët numra bën më mirë."),
            Line("Three: about one place in four can be fixed automatically. The compiler checks "
                 "every change, but the behavior is not yet proven.",
                 "Tre: rreth një vend në katër mund të ndreqet automatikisht. Kompajleri e "
                 "kontrollon çdo ndryshim, por sjellja ende nuk është vërtetuar."),
            Line("Why does this matter in practice? Published thresholds are a starting point, not "
                 "a final answer. And an automatic fix is a proposal that a developer still "
                 "reviews.",
                 "Pse ka rëndësi kjo në praktikë? Pragjet e botuara janë pikënisje, jo përgjigje "
                 "përfundimtare. Dhe një ndreqje automatike është propozim që programuesi e "
                 "rishikon prapë."),
            Line("And what comes next? I want to run the projects' own tests, support more fixes, "
                 "and ask other developers to judge the changes.",
                 "Dhe çka vjen më pas? Dua t'i ekzekutoj testet e vetë projekteve, të shtoj më "
                 "shumë ndreqje, dhe t'u kërkoj programuesve të tjerë t'i gjykojnë ndryshimet."),
            Extra("The limits: one dataset, four smells, and only Java. And I picked the best "
                  "model using the same results I report, so its numbers are a little "
                  "optimistic.",
                  "Kufijtë: një bazë të dhënash, katër erëra, dhe vetëm Java. Dhe modelin më të "
                  "mirë e zgjodha me të njëjtat rezultate që raportoj, ndaj numrat e tij janë pak "
                  "optimistë."),
            Line("Everything is open source, every number can be recreated with a script, and "
                 "the references are on the next slide.",
                 "Gjithçka është me burim të hapur, çdo numër mund të rikrijohet me një skript, "
                 "dhe referencat janë në sllajdin tjetër."),
            Line("Thank you. I am happy to answer your questions.",
                 "Faleminderit. Me kënaqësi u përgjigjem pyetjeve tuaja."),
        ),
        # Referencat nuk thuhen: sllajdi është aty që salla dhe mentorja t'i gjejnë burimet.
        ("References",),
    ]  # fmt: skip


# ----------------------------------------------------------------------
# Sistemi vizual i prezantimit
# ----------------------------------------------------------------------
# Segoe UI ndodhet në çdo Windows, ndaj kompjuteri i sallës e vizaton sllajdin ashtu siç
# u ndërtua; një font i shkarkuar mund të zëvendësohet atje dhe ta prishë faqosjen.
FONT = "Segoe UI"
FONT_STRONG = "Segoe UI Semibold"

# Ngjyrat janë të projektit: blu e figurave të punimit dhe e ndërfaqes, e kuqja vetëm për
# atë që nuk funksionon. Sfondi i errët vetëm për hapjen dhe mbylljen, që salla ta ndiejë
# fillimin dhe fundin pa e lexuar.
NAVY = RGBColor(0x0E, 0x27, 0x40)
ON_DARK = RGBColor(0xF4, 0xF6, 0xF9)
ON_DARK_SOFT = RGBColor(0xA9, 0xB8, 0xC9)
ACCENT_ON_DARK = RGBColor(0x8D, 0xBB, 0xE6)
DANGER_ON_DARK = RGBColor(0xF0, 0x8A, 0x7E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF2, 0xF4, 0xF7)
TINT = RGBColor(0xE4, 0xED, 0xF6)
REFUSED = RGBColor(0xDD, 0xE2, 0xE8)
FRAME = RGBColor(0xC9, 0xD0, 0xD8)
# Gri e rregullave si tekst: më e errët se shtylla, që të lexohet nga fundi i sallës.
RULE_TEXT = RGBColor(0x5E, 0x68, 0x74)
REMOVED_TINT = RGBColor(0xF7, 0xE1, 0xDC)
ADDED_TINT = RGBColor(0xDF, 0xEF, 0xE3)

# Shkalla e shkronjave: titulli lexohet nga rreshti i fundit, teksti nga mesi i sallës.
SIZE_TITLE = Pt(28)
SIZE_LEAD = Pt(20)
SIZE_CAPTION = Pt(16)
SIZE_SMALL = Pt(14)
SIZE_FOOT = Pt(12)
SIZE_STAT = Pt(40)
SIZE_NUMBER = Pt(56)
RADIUS = Inches(0.08)
STAT_HEIGHT = Inches(1.7)
# Trashësia e një metode ose fushe në vizatimet e erërave.
BAR_HEIGHT = Inches(0.09)
# Kodi: rreshti më i gjatë i shembullit duhet të hyjë në dy të tretat e sllajdit.
CODE_SIZE = Pt(13)
TABLE_ROW = Inches(0.5)
GAP = Inches(0.3)
CONTENT = WIDTH - 2 * MARGIN
# Titulli lë vend djathtas për hartën e pyetjeve, në çdo sllajd në të njëjtin vend.
TITLE_WIDTH = WIDTH - 2 * MARGIN - Inches(2.75)
TOP = Inches(2.0)
# Maja e boshtit MCC, e njëjtë te dy grafikët që salla t'i krahasojë me sy.
MCC_TOP = 0.9

# Pjesa e fotos së ndërfaqes që tregon dy gjetje me kushtet e tyre, në pikselat e origjinalit:
# e plotë, teksti i saj del sa një gjilpërë në projektor.
INTERFACE_CROP = (195, 1090, 1300, 1640)


def _frame(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    """Një kuti teksti pa kufi të brendshëm, që teksti të rreshtohet me format përreth."""
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    return frame


def _run(paragraph, text: str, size, colour, strong: bool = False):
    run = paragraph.add_run()
    run.text = text
    run.font.size = size
    run.font.color.rgb = colour
    run.font.name = FONT_STRONG if strong else FONT
    return run


def _put(frame, text: str, size, colour, *, strong=False, new=False, align=PP_ALIGN.LEFT,
         before=0, after=0):  # fmt: skip
    paragraph = frame.add_paragraph() if new else frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    _run(paragraph, text, size, colour, strong)
    return paragraph


def _shape(slide, left, top, width, height, fill, kind=MSO_SHAPE.ROUNDED_RECTANGLE, line=None):
    """Një formë e mbushur, pa hije: rrumbullakimi është i njëjtë në gjithë prezantimin."""
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = min(RADIUS / min(width, height), 0.5)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.18)
    frame.margin_top = frame.margin_bottom = Inches(0.1)
    return shape


def _background(slide, colour) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = colour


def _heading(slide, title: str, subtitle: str = "", dark: bool = False, kicker: str = ""):
    """Titulli në të njëjtin vend në çdo sllajd; titulli thotë gjetjen, nëntitulli burimin."""
    frame = _frame(slide, MARGIN, Inches(0.45), TITLE_WIDTH, Inches(1.4))
    paragraph = frame.paragraphs[0]
    if kicker:
        _run(paragraph, kicker, SIZE_TITLE, ON_DARK_SOFT if dark else SOFT, strong=True)
    _run(paragraph, title, SIZE_TITLE, ON_DARK if dark else INK, strong=True)
    if subtitle:
        _put(frame, subtitle, SIZE_CAPTION, ON_DARK_SOFT if dark else SOFT, new=True, before=6)
    return frame


def _map(slide, active: set[int], role: str = "") -> None:
    """Tri pyetjet lart djathtas; ndizet ajo për të cilën flet sllajdi.

    Salla nuk e mban mend skeletin e një prezantimi dhjetë-minutësh; harta ia kujton në
    çdo sllajd pa asnjë fjalë më shumë. «Q» e jo «RQ», sepse shkurtesa e kërkimit nuk i
    thotë asgjë një studenti nga një fushë tjetër. `role` thotë çka bën sllajdi me
    pyetjen, p.sh. «answers», që salla ta dijë kur dëgjon një përgjigje.
    """
    width, height, gap = Inches(0.78), Inches(0.36), Inches(0.1)
    left = WIDTH - MARGIN - 3 * width - 2 * gap
    for index in (1, 2, 3):
        on = index in active
        pill = _shape(slide, Emu(int(left + (index - 1) * (width + gap))), Inches(0.5), width,
                      height, ACCENT if on else PANEL)  # fmt: skip
        pill.adjustments[0] = 0.5
        frame = pill.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _put(frame, f"Q{index}", SIZE_SMALL, WHITE if on else SOFT, strong=True,
             align=PP_ALIGN.CENTER)  # fmt: skip
    if role:
        label = f"{role} " + " · ".join(f"Q{index}" for index in sorted(active))
        _put(_frame(slide, left, Inches(0.94), Emu(int(WIDTH - MARGIN - left)), Inches(0.3)),
             label, SIZE_SMALL, ACCENT, strong=True, align=PP_ALIGN.RIGHT)  # fmt: skip


def _mcc_key(slide, left, top, width) -> None:
    """Shkalla e MCC-së: pa të, «0.54» nuk i thotë sallës nëse është mirë apo keq."""
    _shape(slide, left, Emu(int(top + Inches(0.08))), width, Inches(0.04), FRAME,
           kind=MSO_SHAPE.RECTANGLE)  # fmt: skip
    for share, label, align in ((0.0, "0 = guessing", PP_ALIGN.LEFT),
                                (1.0, "1 = perfect", PP_ALIGN.RIGHT)):  # fmt: skip
        x = Emu(int(left + width * share))
        # Shenja e skajit të djathtë hyn brenda vijës, që të mos dalë përtej kolonës.
        tick = x if share == 0 else Emu(int(x - Inches(0.04)))
        _shape(slide, tick, top, Inches(0.04), Inches(0.2), SOFT, kind=MSO_SHAPE.RECTANGLE)
        box_w = Emu(int(width / 2))
        box_left = x if align == PP_ALIGN.LEFT else Emu(int(x - box_w))
        _put(_frame(slide, box_left, Emu(int(top + Inches(0.26))), box_w, Inches(0.3)), label,
             SIZE_SMALL, SOFT, strong=True, align=align)  # fmt: skip


def _footer(slide, index: int, dark: bool = False) -> None:
    """Emri dhe numri i sllajdit, që një pyetje të mund të thotë «te sllajdi 5»."""
    colour = ON_DARK_SOFT if dark else SOFT
    top, height = HEIGHT - Inches(0.55), Inches(0.3)
    _put(_frame(slide, MARGIN, top, Inches(4), height), "Florent Latifi  ·  UBT 2026", SIZE_FOOT,
         colour)  # fmt: skip
    _put(_frame(slide, WIDTH - MARGIN - Inches(1), top, Inches(1), height), str(index), SIZE_FOOT,
         colour, align=PP_ALIGN.RIGHT)  # fmt: skip


def _bullet(paragraph, colour) -> None:
    """Një pikë e vërtetë e PowerPoint-it, me dhëmbëzim të varur, jo një «·» në tekst."""
    properties = paragraph._p.get_or_add_pPr()
    indent = int(Inches(0.3))
    properties.set("marL", str(indent))
    properties.set("indent", str(-indent))
    tint = properties.makeelement(qn("a:buClr"), {})
    tint.append(tint.makeelement(qn("a:srgbClr"), {"val": str(colour)}))
    properties.append(tint)
    properties.append(properties.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    properties.append(properties.makeelement(qn("a:buChar"), {"char": "•"}))


def _bullet_list(slide, left, top, width, height, lines, size=SIZE_LEAD, colour=INK, mark=ACCENT,
                 after=12):  # fmt: skip
    frame = _frame(slide, left, top, width, height)
    for index, line in enumerate(lines):
        _bullet(_put(frame, line, size, colour, new=index > 0, after=after), mark)
    return frame


def _stat(slide, left, top, width, number: str, caption: str, colour=ACCENT, size=SIZE_NUMBER,
          caption_colour=SOFT, height=STAT_HEIGHT):  # fmt: skip
    """Një shifër e madhe me nënshkrimin e saj: syri e gjen shifrën para se ta dëgjojë."""
    frame = _frame(slide, left, top, width, height)
    _put(frame, number, size, colour, strong=True)
    _put(frame, caption, SIZE_CAPTION, caption_colour, new=True, before=2)
    return frame


def _badge(slide, left, top, letter: str, fill) -> None:
    """Shkronja e qasjes (A, B, C) në një rreth: e njëjta shenjë në çdo sllajd ku shfaqet."""
    size = Inches(0.36)
    circle = _shape(slide, left, top, size, size, fill, kind=MSO_SHAPE.OVAL)
    frame = circle.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _put(frame, letter, SIZE_SMALL, WHITE, strong=True, align=PP_ALIGN.CENTER)


def _node(slide, left, top, width, height, text: str, fill=PANEL, letter="", letter_fill=ACCENT):
    """Një hap i rrjedhës, me shkronjën e qasjes kur hapi i përket njërës prej tyre."""
    _shape(slide, left, top, width, height, fill)
    inset = Inches(0.16)
    if letter:
        _badge(slide, Emu(int(left + inset)), Emu(int(top + (height - Inches(0.36)) / 2)), letter,
               letter_fill)  # fmt: skip
        inset = Inches(0.66)
    frame = _frame(slide, Emu(int(left + inset)), top, Emu(int(width - inset - Inches(0.14))),
                   height, MSO_ANCHOR.MIDDLE)  # fmt: skip
    _put(frame, text, SIZE_CAPTION, INK)


def _arrow(slide, centre, top, height) -> None:
    arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Emu(int(centre - Inches(0.1))), top,
                                   Inches(0.2), height)  # fmt: skip
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = FRAME
    arrow.line.fill.background()
    arrow.shadow.inherit = False


def _diff(slide, label, lines, marked, tint, colour, left, top, width, size=CODE_SIZE) -> Emu:
    """Kod me rreshtat e ndryshuar në një brez ngjyre, si në një pamje ndryshimesh.

    Lartësia e rreshtit është e saktë, ndaj brezi bie mbi rreshtin e vet pa u matur me sy.
    """
    line = Pt(round(size.pt * 1.15, 1))
    _put(_frame(slide, left, top, width, Inches(0.3)), label, SIZE_SMALL, SOFT, strong=True)
    top = Emu(int(top + Inches(0.34)))
    pad_x, pad_y = Inches(0.16), Inches(0.1)
    height = Emu(int(line * len(lines) + 2 * pad_y))
    _shape(slide, left, top, width, height, PANEL)
    for index in sorted(marked):
        _shape(slide, left, Emu(int(top + pad_y + line * index)), width, line, tint,
               kind=MSO_SHAPE.RECTANGLE)  # fmt: skip
    frame = _frame(slide, Emu(int(left + pad_x)), Emu(int(top + pad_y)),
                   Emu(int(width - 2 * pad_x)), Emu(int(height - 2 * pad_y)))  # fmt: skip
    frame.word_wrap = False
    for index, text in enumerate(lines):
        paragraph = frame.add_paragraph() if index else frame.paragraphs[0]
        paragraph.line_spacing = line
        paragraph.space_before = paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = text or " "
        run.font.name = MONO
        run.font.size = size
        run.font.bold = index in marked
        run.font.color.rgb = colour if index in marked else INK
    return Emu(int(top + height))


def _data_table(slide, headers, rows, top=TOP, size=SIZE_CAPTION, row=TABLE_ROW):
    """Tabelat rezervë: koka me blu, rreshtat me një nuancë të lehtë në çdo të dytin."""
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), MARGIN, top, CONTENT,
                                   Emu(int(row * (len(rows) + 1))))  # fmt: skip
    table = shape.table
    # Kolonat sipas tekstit më të gjatë të secilës: me gjerësi të barabarta, «Feature Envy»
    # dhe emrat e modeleve thyhen në dy rreshta ndërsa kolonat e numrave mbeten gjysmë bosh.
    longest = [max(max(len(str(row[column])) for row in rows), len(headers[column]), 6)
               for column in range(len(headers))]  # fmt: skip
    for column, length in zip(table.columns, longest, strict=True):
        column.width = Emu(int(CONTENT * length / sum(longest)))
    for line, values in enumerate([headers, *rows]):
        for column, value in enumerate(values):
            cell = table.cell(line, column)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if line == 0 else (PANEL if line % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame = cell.text_frame
            frame.paragraphs[0].text = ""
            _put(frame, str(value), size if line else SIZE_SMALL, WHITE if line == 0 else INK,
                 strong=line == 0 or column == 0,
                 align=PP_ALIGN.LEFT if column == 0 else PP_ALIGN.RIGHT)  # fmt: skip
    return table


# ----------------------------------------------------------------------
# Sllajdet
# ----------------------------------------------------------------------
def _headings(f: dict) -> list[str]:
    """Titulli i çdo sllajdi kryesor: gjetja në një fjali, që salla ta lexojë para se ta dëgjojë.

    Prezantimi dhe udhëzuesi i fjalimit i marrin nga këtu, që të mos thonë dy tituj.
    """
    better = sum(f["best"][s]["mcc"] > f["strategy"][s]["mcc"] for s in f["smells"])
    count = "all four" if better == len(f["smells"]) == 4 else f"{better} of {len(f['smells'])}"
    chance = _and([SMELL_SQ[s] for s in f["chance"]])
    return [
        TITLE,
        "Code smells make every change expensive, and today's tools only point at them",
        "What this work adds to earlier studies",
        "Four code smells, each a different design problem",
        f"JavaSmell measures {f['metrics']} metrics once and answers with three approaches",
        "Professional developers label the code; no model is tested on a project it saw",
        f"Learned models beat the rules on {count} smells",
        f"For {chance}, even two experts agree no better than chance",
        "The engine rewrites only what it can prove, and refuses the rest with a reason",
        "Each rewrite is checked before and after, but compiling ≠ the same behavior",
        "Real code: the outer if becomes an early return",
        "Key findings and what comes next",
        "References",
    ]


def _count_word(count: int, total: int) -> str:
    """«three», «all four»: një numër i vogël lexohet më shpejt si fjalë në një fjali."""
    words = {1: "one", 2: "two", 3: "three", 4: "four"}
    return f"all {words.get(total, str(total))}" if count == total else words.get(count, str(count))


def _slide_title(slide, f, heading) -> None:
    _background(slide, NAVY)
    _put(_frame(slide, MARGIN, Inches(0.6), CONTENT, Inches(0.4)), EVENT, SIZE_CAPTION,
         ON_DARK_SOFT)  # fmt: skip
    _put(_frame(slide, MARGIN, Inches(1.75), Inches(11.0), Inches(2.4)), heading, Pt(40), ON_DARK,
         strong=True)  # fmt: skip
    frame = _frame(slide, MARGIN, Inches(4.45), CONTENT, Inches(1.0))
    _put(frame, AUTHORS, SIZE_LEAD, ON_DARK)
    _put(frame, AFFILIATION, SIZE_CAPTION, ON_DARK_SOFT, new=True, before=4)
    _put(_frame(slide, MARGIN, Inches(6.5), CONTENT, Inches(0.4)), SESSION, SIZE_CAPTION,
         ON_DARK_SOFT)  # fmt: skip


def _slide_problem(slide, f, heading) -> None:
    _heading(slide, heading, f"Code smells are design problems, not bugs {_cite('Fowler 2018')}")
    _bullet_list(slide, MARGIN, TOP, CONTENT, Inches(1.5), [
        "The program works, but every change costs more than it should.",
        f"Existing tools use fixed thresholds, calibrated on other code "
        f"{_cite('Lanza & Marinescu 2006')}.",
        "They point at the problem, but rarely fix it.",
    ], after=8)  # fmt: skip
    questions = [
        ("Question 1", "How accurate are published detection rules?", "A", "Rules", MUTED),
        ("Question 2", "Does machine learning on the same metrics do better?", "B",
         "Machine learning", ACCENT),
        ("Question 3", "Can fixes be applied automatically without breaking compilation?", "C",
         "Refactoring engine", ACCENT),
    ]  # fmt: skip
    width = Emu(int((CONTENT - 2 * GAP) / 3))
    top, height = Inches(3.85), Inches(2.6)
    for index, (label, question, letter, approach, fill) in enumerate(questions):
        left = Emu(int(MARGIN + index * (width + GAP)))
        _shape(slide, left, top, width, height, PANEL)
        frame = _frame(slide, Emu(int(left + Inches(0.25))), Emu(int(top + Inches(0.25))),
                       Emu(int(width - Inches(0.5))), Inches(1.6))  # fmt: skip
        _put(frame, f"Q{index + 1}  ·  {label}", SIZE_CAPTION, ACCENT, strong=True)
        _put(frame, question, SIZE_LEAD, INK, strong=True, new=True, before=6)
        badge_top = Emu(int(top + height - Inches(0.6)))
        _badge(slide, Emu(int(left + Inches(0.25))), badge_top, letter, fill)
        _put(_frame(slide, Emu(int(left + Inches(0.72))), badge_top, Inches(2.6), Inches(0.36),
                    MSO_ANCHOR.MIDDLE), approach, SIZE_CAPTION, SOFT)  # fmt: skip


def _slide_contribution(slide, f, heading) -> None:
    """Çka e dallon punimin: secila mangësi e literaturës krah përgjigjes së punimit.

    Rreshtat janë ata të «Boshllëku në literaturë» te punimi (2.7), me të njëjtat burime,
    që mentorja dhe salla ta lidhin kontributin me studimet që ai i plotëson.
    """
    _heading(slide, heading, "Three gaps in earlier studies, and how this work addresses each")
    earlier_w = Inches(5.0)
    arrow_w = Inches(0.55)
    ours_left = Emu(int(MARGIN + earlier_w + arrow_w))
    ours_w = Emu(int(WIDTH - MARGIN - ours_left))
    top = Inches(1.95)
    _put(_frame(slide, MARGIN, top, earlier_w, Inches(0.3)), "Earlier studies", SIZE_SMALL, SOFT,
         strong=True)  # fmt: skip
    _put(_frame(slide, ours_left, top, ours_w, Inches(0.3)), "This work", SIZE_SMALL, ACCENT,
         strong=True)  # fmt: skip
    rows = [
        (f"Published rules are taken as correct, with thresholds set on other systems "
         f"{_cite('Lanza & Marinescu 2006')}.",
         f"The rules are measured against 26 professional developers in MLCQ "
         f"{_cite('Madeyski & Lewowski 2020')}."),
        (f"Rules and machine learning are rarely compared on the same data, and scores depend on "
         f"how the data are built {_cite('Arcelli Fontana et al. 2016', 'Di Nucci et al. 2018')}.",
         f"Same {f['samples']:,} samples, same score (MCC), split by project, with 95% "
         "intervals."),
        (f"Finding and fixing are studied apart: a fix is proposed, and the decision is left to "
         f"the developer {_cite('Tsantalis & Chatzigeorgiou 2009')}.",
         f"One chain: find → rewrite → compile → measure again, with the preconditions proven "
         f"first {_cite('Opdyke 1992')}."),
    ]  # fmt: skip
    top, height, gap = Inches(2.35), Inches(1.05), Inches(0.15)
    for earlier, ours in rows:
        for left, width, fill, text, colour in ((MARGIN, earlier_w, PANEL, earlier, RULE_TEXT),
                                                (ours_left, ours_w, TINT, ours, INK)):  # fmt: skip
            frame = _shape(slide, left, top, width, height, fill).text_frame
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame.margin_left = frame.margin_right = Inches(0.22)
            _put(frame, text, SIZE_CAPTION, colour, strong=colour == INK)
        _put(_frame(slide, Emu(int(MARGIN + earlier_w)), top, arrow_w, height, MSO_ANCHOR.MIDDLE),
             "→", SIZE_LEAD, ACCENT, strong=True, align=PP_ALIGN.CENTER)  # fmt: skip
        top = Emu(int(top + height + gap))
    _put(_frame(slide, ours_left, Emu(int(top + Inches(0.05))), ours_w, Inches(0.6)),
         "And an open tool: a plain-language interface in Albanian, and every number "
         "regenerated by one script.", SIZE_CAPTION, ACCENT, strong=True)  # fmt: skip


def _box(slide, left, top, width, height, fill=WHITE):
    return _shape(slide, left, top, width, height, fill, line=FRAME)


def _bar(slide, left, top, width, colour, height=BAR_HEIGHT) -> None:
    _shape(slide, left, top, width, height, colour, kind=MSO_SHAPE.RECTANGLE)


def _field(slide, left, top, width) -> None:
    """Një fushë: pikë gri dhe një vijë. Metodat janë vija blu, që dallohen pa legjendë."""
    _shape(slide, left, Emu(int(top - Inches(0.02))), Inches(0.13), Inches(0.13), MUTED,
           kind=MSO_SHAPE.OVAL)  # fmt: skip
    _bar(slide, Emu(int(left + Inches(0.2))), top, Emu(int(width - Inches(0.2))), REFUSED)


def _link(slide, x1, y1, x2, y2, colour) -> None:
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = colour
    line.line.width = Pt(1.5)


def _caption(slide, left, top, width, text) -> None:
    """Një etiketë e qendërzuar mbi një kuti, më e gjerë se kutia që të mos thyhet."""
    spare = Inches(0.2)
    _put(_frame(slide, Emu(int(left - spare)), top, Emu(int(width + 2 * spare)), Inches(0.3)), text,
         SIZE_SMALL, SOFT, align=PP_ALIGN.CENTER)  # fmt: skip


def _draw_blob(slide, x, y) -> None:
    """Një klasë e madhe me shumë metoda, që merr edhe të dhëna nga dy klasa të vogla."""
    _box(slide, x, y, Inches(1.4), Inches(2.0))
    for row in range(9):
        _bar(slide, Emu(int(x + Inches(0.18))), Emu(int(y + Inches(0.2) + row * Inches(0.18))),
             Inches(1.04), ACCENT)  # fmt: skip
    for row in (0, 1):
        top = Emu(int(y + Inches(0.3) + row * Inches(1.05)))
        _box(slide, Emu(int(x + Inches(1.7))), top, Inches(0.42), Inches(0.42))
        _link(slide, Emu(int(x + Inches(1.4))), Emu(int(top + Inches(0.21))),
              Emu(int(x + Inches(1.7))), Emu(int(top + Inches(0.21))), FRAME)  # fmt: skip


def _draw_data_class(slide, x, y) -> None:
    """Një klasë vetëm me fusha, dhe një klasë tjetër me metodën që punon me to."""
    _box(slide, x, Emu(int(y + Inches(0.1))), Inches(1.3), Inches(1.5))
    for row in range(5):
        _field(slide, Emu(int(x + Inches(0.18))), Emu(int(y + Inches(0.35) + row * Inches(0.25))),
               Inches(0.95))  # fmt: skip
    other_x, other_y = Emu(int(x + Inches(1.55))), Emu(int(y + Inches(1.2)))
    _box(slide, other_x, other_y, Inches(0.6), Inches(0.8))
    _bar(slide, Emu(int(other_x + Inches(0.1))), Emu(int(other_y + Inches(0.35))), Inches(0.4),
         ACCENT)  # fmt: skip
    _link(slide, Emu(int(other_x + Inches(0.1))), Emu(int(other_y + Inches(0.39))),
          Emu(int(x + Inches(1.3))), Emu(int(y + Inches(0.9))), ACCENT)  # fmt: skip


def _draw_long_method(slide, x, y) -> None:
    """Një metodë e vetme si një kolonë e gjatë rreshtash kodi."""
    _box(slide, Emu(int(x + Inches(0.35))), y, Inches(1.45), Inches(2.1))
    lengths = [0.9, 0.7, 1.0, 0.55, 0.85, 0.95, 0.6, 1.0, 0.75, 0.5, 0.9, 0.8, 0.65, 0.95, 0.7]
    for row, length in enumerate(lengths):
        _bar(slide, Emu(int(x + Inches(0.5) + Inches(0.12 if row % 4 else 0))),
             Emu(int(y + Inches(0.15) + row * Inches(0.12))), Inches(length), RULE_TEXT,
             Inches(0.05))  # fmt: skip


def _draw_feature_envy(slide, x, y) -> None:
    """Një metodë në klasën e vet që zgjatet te fushat e një klase tjetër."""
    own_x, other_x, top = x, Emu(int(x + Inches(1.3))), Emu(int(y + Inches(0.4)))
    _caption(slide, own_x, y, Inches(0.85), "own class")
    _caption(slide, other_x, y, Inches(0.85), "other class")
    _box(slide, own_x, top, Inches(0.85), Inches(1.4))
    _box(slide, other_x, top, Inches(0.85), Inches(1.4))
    method_y = Emu(int(top + Inches(0.65)))
    _bar(slide, Emu(int(own_x + Inches(0.12))), method_y, Inches(0.6), ACCENT)
    for row in range(3):
        field_y = Emu(int(top + Inches(0.3) + row * Inches(0.38)))
        _field(slide, Emu(int(other_x + Inches(0.12))), field_y, Inches(0.6))
        _link(slide, Emu(int(own_x + Inches(0.72))), Emu(int(method_y + Inches(0.045))),
              Emu(int(other_x + Inches(0.12))), Emu(int(field_y + Inches(0.045))), ACCENT)  # fmt: skip


def _slide_smells(slide, f, heading) -> None:
    _heading(slide, heading, f"Definitions after Fowler {_cite('Fowler 2018')}; all four are labeled "
                             f"by professional developers in MLCQ "
                             f"{_cite('Madeyski & Lewowski 2020')}")  # fmt: skip
    cards = [
        ("Blob", "A class that does too much.", _draw_blob),
        ("Data Class", "Only holds data; almost no behavior.", _draw_data_class),
        ("Long Method", "A method that has grown too long.", _draw_long_method),
        ("Feature Envy", "A method that uses another class more than its own.",
         _draw_feature_envy),
    ]  # fmt: skip
    width = Emu(int((CONTENT - 3 * GAP) / 4))
    height = Inches(3.85)
    for index, (name, definition, draw) in enumerate(cards):
        left = Emu(int(MARGIN + index * (width + GAP)))
        _shape(slide, left, TOP, width, height, PANEL)
        draw(slide, Emu(int(left + Inches(0.25))), Emu(int(TOP + Inches(0.3))))
        frame = _frame(slide, Emu(int(left + Inches(0.25))), Emu(int(TOP + Inches(2.65))),
                       Emu(int(width - Inches(0.5))), Inches(1.5))  # fmt: skip
        _put(frame, name, SIZE_LEAD, INK, strong=True)
        _put(frame, definition, SIZE_CAPTION, INK, new=True, before=4)
    legend = _frame(slide, MARGIN, Inches(6.05), CONTENT, Inches(0.3))
    paragraph = legend.paragraphs[0]
    _run(paragraph, "Blue line", SIZE_SMALL, ACCENT, strong=True)
    _run(paragraph, " = method   ·   ", SIZE_SMALL, SOFT)
    _run(paragraph, "gray dot", SIZE_SMALL, RULE_TEXT, strong=True)
    _run(paragraph, " = field (a piece of data)   ·   box = class", SIZE_SMALL, SOFT)


def _slide_system(slide, f, heading) -> None:
    _heading(slide, heading, "One analyzer: A rules · B machine learning · C refactoring engine")
    left, width, height, gap = MARGIN, Inches(6.2), Inches(0.52), Inches(0.2)
    half = Emu(int((width - Inches(0.2)) / 2))
    centre = Emu(int(left + width / 2))
    rows = [
        ("Java source code", None),
        (f"{f['metrics']} metrics: lines, complexity, coupling, … "
         f"{_cite('Lanza & Marinescu 2006', 'Chidamber & Kemerer 1994')}", None),
        ("split", None),
        ("Smells, each with the conditions that fired", None),
        ("Refactoring engine: rewrites only what it can prove", "C"),
        ("javac check, then a patch — never written in place", None),
    ]
    # Rreshti i ndarë (A dhe B) merr dy shigjeta, një për secilën kuti, që rrjedha të mos
    # duket sikur kalon mes tyre.
    halves = [Emu(int(left + half / 2)), Emu(int(left + half + Inches(0.2) + half / 2))]
    top = TOP
    for index, (text, letter) in enumerate(rows):
        if index:
            split = "split" in (text, rows[index - 1][0])
            for centre_x in halves if split else [centre]:
                _arrow(slide, centre_x, Emu(int(top - gap + Inches(0.03))),
                       Emu(int(gap - Inches(0.06))))  # fmt: skip
        if text == "split":
            _node(slide, left, top, half, height, f"{f['strategies']} published rules", PANEL, "A",
                  MUTED)  # fmt: skip
            _node(slide, Emu(int(left + half + Inches(0.2))), top, half, height,
                  "Machine learning", TINT, "B")  # fmt: skip
        else:
            _node(slide, left, top, width, height, text, TINT if letter else PANEL, letter or "")
        top = Emu(int(top + height + gap))
    picture_left = Emu(int(MARGIN + width + Inches(0.45)))
    picture_width = Emu(int(WIDTH - MARGIN - picture_left))
    path = os.path.join(FIGURES, "nderfaqja_paneli.png")
    x0, y0, x1, y1 = INTERFACE_CROP
    picture = slide.shapes.add_picture(path, picture_left, Inches(2.45), width=picture_width,
                                       height=Emu(int(picture_width * (y1 - y0) / (x1 - x0))))  # fmt: skip
    pixels_w, pixels_h = picture.image.size
    picture.crop_left, picture.crop_right = x0 / pixels_w, 1 - x1 / pixels_w
    picture.crop_top, picture.crop_bottom = y0 / pixels_h, 1 - y1 / pixels_h
    picture.line.color.rgb = FRAME
    picture.line.width = Pt(0.75)
    _put(_frame(slide, picture_left, Emu(int(picture.top + picture.height + Inches(0.15))),
                picture_width, Inches(0.9)),
         "Web interface in plain language (Albanian): every finding shows the conditions that "
         "fired", SIZE_CAPTION, SOFT)  # fmt: skip


def _slide_data(slide, f, heading) -> None:
    _heading(slide, heading, f"MLCQ {_cite('Madeyski & Lewowski 2020')}: reviewed by 26 "
                             "professional developers")  # fmt: skip
    _map(slide, {1, 2}, "testing")
    _stat(slide, MARGIN, TOP, Inches(6.0), f"{f['samples']:,}",
          "pieces of code (classes or methods) judged by professional developers")  # fmt: skip
    _stat(slide, Emu(int(MARGIN + Inches(6.4))), Emu(int(TOP + Inches(0.12))), Inches(5.0),
          f"{f['repositories']}", "open-source projects they come from", ACCENT, SIZE_STAT)  # fmt: skip
    # Pesë grupet e projekteve: modeli mëson nga katër dhe provohet te i pesti.
    folds, top, fold_w, fold_h = 5, Inches(3.95), Inches(1.22), Inches(1.6)
    gap, tile, tile_gap = Inches(0.17), Inches(0.24), Inches(0.1)
    for index in range(folds):
        left = Emu(int(MARGIN + index * (fold_w + gap)))
        test = index == folds - 1
        _shape(slide, left, top, fold_w, fold_h, TINT if test else PANEL)
        grid = 3 * tile + 2 * tile_gap
        x0 = Emu(int(left + (fold_w - grid) / 2))
        for cell in range(9):
            row, column = divmod(cell, 3)
            _shape(slide, Emu(int(x0 + column * (tile + tile_gap))),
                   Emu(int(top + Inches(0.2) + row * (tile + tile_gap))), tile, tile,
                   ACCENT if test else MUTED)  # fmt: skip
        _put(_frame(slide, left, Emu(int(top + fold_h - Inches(0.44))), fold_w, Inches(0.34)),
             "test" if test else "learn", SIZE_CAPTION, ACCENT if test else SOFT, strong=test,
             align=PP_ALIGN.CENTER)  # fmt: skip
    diagram_w = Emu(int(folds * fold_w + (folds - 1) * gap))
    _put(_frame(slide, MARGIN, Emu(int(top + fold_h + Inches(0.15))), diagram_w, Inches(0.8)),
         f"Projects split into {f['ml']['folds']} groups: learn from four, test on the fifth, "
         "which the model has never seen. Five times.", SIZE_CAPTION, SOFT)  # fmt: skip
    majority = f["ml"]["per_smell"]["feature envy"]["models"]["majority"]["accuracy"]
    panel_left = Emu(int(MARGIN + diagram_w + Inches(0.45)))
    panel_w = Emu(int(WIDTH - MARGIN - panel_left))
    _shape(slide, panel_left, top, panel_w, Inches(2.45), PANEL)
    frame = _frame(slide, Emu(int(panel_left + Inches(0.25))), Emu(int(top + Inches(0.2))),
                   Emu(int(panel_w - Inches(0.5))), Inches(2.1))  # fmt: skip
    _put(frame, f"Why MCC, not accuracy {_cite('Chicco & Jurman 2020')}", SIZE_CAPTION, INK,
         strong=True)  # fmt: skip
    _put(frame, f"{majority:.1%}", SIZE_STAT, CRITICAL, strong=True, new=True, before=2)
    _put(frame, "accuracy of a detector that always says “clean” on Feature Envy. Its MCC is 0, "
         "the same as guessing.", SIZE_CAPTION, INK, new=True, before=2)  # fmt: skip


def _side(slide):
    """Kolona djathtas e grafikëve: e njëjta gjerësi dhe e njëjta shkallë në të dy sllajdet."""
    left = Emu(int(MARGIN + Inches(8.6)))
    width = Emu(int(WIDTH - MARGIN - left))
    _put(_frame(slide, left, Inches(1.95), width, Inches(0.3)), "How to read MCC", SIZE_SMALL, INK,
         strong=True)  # fmt: skip
    _mcc_key(slide, left, Inches(2.3), width)
    return left, width


def _slide_results(slide, f, heading) -> None:
    _heading(slide, heading, f"Same {f['samples']:,} samples, same scoring · higher is better")
    _map(slide, {1, 2}, "answers")
    # Vendimi statistikor nën emrin e erës: grafiku e thotë vetë cila fitore është e sigurt.
    names = [f"{SMELL_SQ[s]}\n({'significant' if SMELL_SQ[s] in f['sure'] else 'could be chance'})"
             for s in f["smells"]]  # fmt: skip
    _bar_chart(slide, names, [("Published rules (A)", f["rule_mcc"], MUTED),
                              ("Best ML model (B)", f["model_mcc"], ACCENT)],
               top=TOP, height=Inches(4.7), width=Inches(8.2), maximum=MCC_TOP)  # fmt: skip
    left, width = _side(slide)
    found = [f["best"][s]["recall"] for s in f["smells"]]
    _put(_frame(slide, left, Inches(3.2), width, Inches(0.3)), "Out of 100 real smells", SIZE_CAPTION,
         INK, strong=True)  # fmt: skip
    _stat(slide, left, Inches(3.55), width,
          f"{min(f['recall']) * 100:.0f}–{max(f['recall']) * 100:.0f}", "found by the rules",
          RULE_TEXT, SIZE_STAT, height=Inches(1.1))  # fmt: skip
    _stat(slide, left, Inches(4.7), width, f"{min(found) * 100:.0f}–{max(found) * 100:.0f}",
          "found by the models", ACCENT, SIZE_STAT, height=Inches(1.1))  # fmt: skip
    _put(_frame(slide, left, Inches(5.85), width, Inches(0.6)),
         "The price: usually more false alarms.", SIZE_CAPTION, SOFT)  # fmt: skip


def _slide_agreement(slide, f, heading) -> None:
    _heading(slide, heading, "Two reviewers judged the same code: how often did they agree?")
    _map(slide, {2}, "explains")
    smells, chance = f["smells"], f["chance"]
    chart = _bar_chart(slide, f["names"], [("Two reviewers", [f["reviewers"][s] for s in smells],
                                            ACCENT)],
                       top=TOP, height=Inches(4.7), width=Inches(8.2), minimum=-0.2,
                       maximum=MCC_TOP)  # fmt: skip
    # Erët ku pajtimi është në nivelin e rastësisë dalin me ngjyrë tjetër: për to flet sllajdi.
    for index, smell in enumerate(smells):
        if smell in chance:
            point = chart.plots[0].series[0].points[index]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = CRITICAL
            # Një pikë me format të vetin nuk e trashëgon `invertIfNegative` të serisë, dhe
            # PowerPoint-i e vizaton shtyllën negative bosh nëse pika nuk e thotë vetë.
            dpt = point.format._element
            flag = dpt.makeelement(qn("c:invertIfNegative"), {"val": "0"})
            dpt.find(qn("c:idx")).addnext(flag)
    left, width = _side(slide)
    for s in chance[:1]:
        _stat(slide, left, Inches(3.2), width, f"{f['reviewers'][s]:.2f}".replace("-", "−"),
              f"{SMELL_SQ[s]}: agreement of a coin flip", CRITICAL)  # fmt: skip
    _put(_frame(slide, left, Inches(4.95), width, Inches(1.7)),
         "If experts do not agree on what a smell is, no detector can learn it well.",
         SIZE_LEAD, INK, strong=True)  # fmt: skip


def _slide_engine(slide, f, heading) -> None:
    _heading(slide, heading, "Three transformations: Extract Method · Guard Clauses · Introduce "
                             "Parameter Object")  # fmt: skip
    _map(slide, {3}, "answers")
    refactoring = f["refactoring"]
    applied, detected = f["applied"], f["detected"]
    refused, missing = refactoring["refused"], refactoring["unlocatable"]
    frame = _frame(slide, MARGIN, TOP, CONTENT, Inches(0.7), MSO_ANCHOR.BOTTOM)
    paragraph = frame.paragraphs[0]
    _run(paragraph, f"{detected:,}", SIZE_STAT, INK, strong=True)
    _run(paragraph, "   places in the code where a fix could apply", SIZE_CAPTION, SOFT)
    top, height = Inches(2.85), Inches(0.8)
    rewritten_w = Emu(int(CONTENT * applied / detected))
    refused_w = Emu(int(CONTENT * refused / detected) - Inches(0.05))
    # Gjerësia e çdo pjese është pjesa e vendeve: shiriti e tregon raportin pa e lexuar.
    for segment_left, segment_w, fill, colour, label, share in (
        (MARGIN, rewritten_w, ACCENT, WHITE, f"{applied:,} rewritten", applied / detected),
        (Emu(int(MARGIN + rewritten_w + Inches(0.05))), refused_w, REFUSED, INK,
         f"{refused:,} refused: not provably safe, each with a reason", refused / detected),
    ):  # fmt: skip
        frame = _shape(slide, segment_left, top, segment_w, height, fill).text_frame
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_top = frame.margin_bottom = Inches(0.04)
        _put(frame, label, SIZE_CAPTION, colour, strong=True)
        _put(frame, _pct(share), SIZE_CAPTION, colour, new=True)
    # Pjesa e tretë është shumë e hollë për shiritin, ndaj shuma shkruhet e plotë: pa të, dy
    # numrat e shiritit nuk mblidhen te i pari (vërejtja e mentores, 23 tetor).
    frame = _frame(slide, MARGIN, Inches(3.72), CONTENT, Inches(0.3))
    paragraph = frame.paragraphs[0]
    _run(paragraph, f"+ {missing} ({_pct(missing / detected)}) could not be found again in the "
         "file, so no fix was tried", SIZE_SMALL, SOFT)  # fmt: skip
    _run(paragraph, f"      {applied:,} + {refused:,} + {missing} = {detected:,}", SIZE_SMALL, INK,
         strong=True)  # fmt: skip
    column_w = Emu(int((CONTENT - Inches(0.6)) / 2))
    reasons = sorted(refactoring["refused_by_reason"].items(), key=lambda pair: -pair[1])
    reason_rows = [(REASON_EN[name], count) for name, count in reasons[:3]]
    rest = reasons[3:]
    reason_rows.append((f"{_count_word(len(rest), len(reasons)).capitalize()} other reasons",
                        sum(count for _, count in rest)))  # fmt: skip
    by_fix = sorted(refactoring["applied_by_refactoring"].items(), key=lambda pair: -pair[1])
    for left, title, rows, total, colour in (
        (MARGIN, f"The {applied:,} rewrites, by fix",
         [(_words(name), count) for name, count in by_fix], applied, ACCENT),
        (Emu(int(MARGIN + column_w + Inches(0.6))), f"Why the {refused:,} were refused", reason_rows,
         refused, MUTED),
    ):  # fmt: skip
        _put(_frame(slide, left, Inches(4.3), column_w, Inches(0.3)), title, SIZE_CAPTION, SOFT,
             strong=True)  # fmt: skip
        _share_rows(slide, left, Inches(4.75), column_w, rows, total, colour)


# Arsyet e refuzimit me fjalë sallë; kodet janë ato të motorit (`refactor/base.py`).
REASON_EN = {
    "shape_not_matched": "Code shape does not match the fix",
    "control_flow_escapes": "A return, break or continue leaves the block",
    "not_definitely_assigned": "An input value is not definitely assigned",
    "multiple_outputs": "More than one value comes out",
    "ambiguous_overload": "An overload cannot be told apart",
    "unresolved_name": "A name cannot be resolved",
}
# Të njëjtat arsye me fjalët e Tabelës 11 të punimit, për udhëzuesin shqip.
REASON_SQ_SHORT = {
    "shape_not_matched": "forma e kodit nuk përputhet",
    "control_flow_escapes": "rrjedha e kontrollit del nga blloku",
    "not_definitely_assigned": "vlerë e pacaktuar me siguri",
    "multiple_outputs": "më shumë se një vlerë dalëse",
    "ambiguous_overload": "mbingarkesë e paqartë",
    "unresolved_name": "emër i pazgjidhur",
}


def _per_fix(f: dict) -> dict[str, dict[str, int]]:
    """Vendet, rishkrimet dhe verdiktet sipas ndreqjes, nga rreshti i çdo vendi.

    Llogaritet si tabela «sipas transformimit» e punimit (`chapters._by_transformation`),
    nga i njëjti skedar, që shumat të përkojnë me totalet e sllajdit të motorit. Skedari nuk
    i mban vendet që nuk u gjetën dot sërish, ndaj vendet e tij janë totali pa to.
    """
    rows = _rows_if_present("refactoring_sites.csv") or []
    out: dict[str, dict[str, int]] = {}
    for row in rows:
        entry = out.setdefault(row["refactoring"], dict.fromkeys(
            ("sites", "applied", "passed", "new_errors", "resolved"), 0))  # fmt: skip
        entry["sites"] += 1
        if row["applied"] != "1":
            continue
        entry["applied"] += 1
        entry["passed"] += row["verdict"] in ("compiles", "no_new_errors")
        entry["new_errors"] += row["verdict"] == "new_errors"
        entry["resolved"] += row["resolution"] == "resolved"
    totals = {key: sum(entry[key] for entry in out.values()) for key in ("sites", "applied")}
    expected = {"sites": f["detected"] - f["refactoring"]["unlocatable"], "applied": f["applied"]}
    if totals != expected:
        raise ValueError(f"refactoring_sites.csv nuk përkon me totalet: {totals}")
    return out


def _share_rows(slide, left, top, width, rows, total, colour) -> None:
    """Emri, numri dhe një shirit sa pjesa e tij: renditja lexohet pa e lexuar numrin."""
    pitch, bar_h = Inches(0.42), Inches(0.08)
    count_w = Inches(1.0)
    for label, count in rows:
        _put(_frame(slide, left, top, Emu(int(width - count_w)), Inches(0.28)), label, SIZE_SMALL,
             INK)  # fmt: skip
        _put(_frame(slide, Emu(int(left + width - count_w)), top, count_w, Inches(0.28)),
             f"{count:,}", SIZE_SMALL, INK, strong=True, align=PP_ALIGN.RIGHT)  # fmt: skip
        _shape(slide, left, Emu(int(top + Inches(0.29))), width, bar_h, PANEL,
               kind=MSO_SHAPE.RECTANGLE)  # fmt: skip
        _shape(slide, left, Emu(int(top + Inches(0.29))),
               Emu(max(int(width * count / total), int(Inches(0.04)))), bar_h, colour,
               kind=MSO_SHAPE.RECTANGLE)  # fmt: skip
        top = Emu(int(top + pitch))


def _slide_validation(slide, f, heading) -> None:
    """Si u kontrollua çdo rishkrim, hap pas hapi, dhe çka nuk e kontrollon asnjë hap.

    Mentorja kërkoi që procesi i validimit të shpjegohej dhe që të thuhej qartë se
    kompilimi nuk e garanton sjelljen; paneli djathtas e thotë këtë me shkronja të mëdha.
    """
    _heading(slide, heading, "How the rewrites were validated, from the syntax tree to a manual "
                             "review")  # fmt: skip
    _map(slide, {3}, "validates")
    refactoring, quality = f["refactoring"], f["quality"]
    applied = f["applied"]
    verdicts, resolution = refactoring["verdicts"], refactoring["resolution"]
    passed = verdicts.get("compiles", 0) + verdicts.get("no_new_errors", 0)
    verify = _load("verify_with_project.json")
    in_project = verify["compiled_in_project"]
    changed = sum(e["behaviour"]["changed"] for e in quality["by_refactoring"].values())
    steps = [
        (f"Before: prove the preconditions {_cite('Opdyke 1992')}",
         "on the syntax tree, or refuse with a reason",
         f"{applied:,} rewritten", f"{refactoring['refused']:,} refused"),
        ("After: compile with javac",
         "the rewritten file may add no new kind of error",
         f"{passed:,} pass ({_pct(passed / applied)})",
         f"{verdicts.get('new_errors', 0)} fail, never written · {verdicts.get('parses', 0)} no "
         "verdict"),
        ("Measure the smell again",
         "does the detector still fire on the new code?",
         f"{resolution['resolved']:,} removed ({_pct(f['resolved'])})",
         f"{resolution['persists']:,} remain · {resolution['unknown']} not identifiable"),
        ("Sample: compile inside the project",
         f"{verify['rewrites']} rewrites in {verify['files_checked']} files, with their neighbours",
         f"{in_project['compiles']} compile fully",
         f"{in_project['new_errors']} new errors · {in_project['not_checked']} no verdict"),
        ("Read by hand",
         f"{quality['reviewed']} rewrites: behavior, benefit, acceptance",
         f"{_pct(quality['acceptable_reweighted'])} acceptable",
         f"{changed} changed the behavior"),
    ]  # fmt: skip
    width, height, pitch = Inches(8.0), Inches(0.8), Inches(0.9)
    text_w, outcome_w = Inches(4.15), Inches(3.05)
    top = Inches(1.95)
    for index, (title, detail, outcome, cost) in enumerate(steps, 1):
        _shape(slide, MARGIN, top, width, height, PANEL)
        _badge(slide, Emu(int(MARGIN + Inches(0.16))), Emu(int(top + (height - Inches(0.36)) / 2)),
               str(index), ACCENT)  # fmt: skip
        frame = _frame(slide, Emu(int(MARGIN + Inches(0.66))), top, text_w, height,
                       MSO_ANCHOR.MIDDLE)  # fmt: skip
        _put(frame, title, SIZE_CAPTION, INK, strong=True)
        _put(frame, detail, SIZE_SMALL, SOFT, new=True)
        frame = _frame(slide, Emu(int(MARGIN + width - outcome_w - Inches(0.18))), top, outcome_w,
                       height, MSO_ANCHOR.MIDDLE)  # fmt: skip
        _put(frame, outcome, SIZE_CAPTION, ACCENT, strong=True, align=PP_ALIGN.RIGHT)
        _put(frame, cost, SIZE_SMALL, CRITICAL if index > 1 else SOFT, new=True,
             align=PP_ALIGN.RIGHT)  # fmt: skip
        top = Emu(int(top + pitch))
    left = Emu(int(MARGIN + width + Inches(0.35)))
    panel_w = Emu(int(WIDTH - MARGIN - left))
    panel_h = Emu(int(4 * pitch + height))
    _shape(slide, left, Inches(1.95), panel_w, panel_h, REMOVED_TINT)
    frame = _frame(slide, Emu(int(left + Inches(0.25))), Inches(2.15),
                   Emu(int(panel_w - Inches(0.5))), Emu(int(panel_h - Inches(0.4))))  # fmt: skip
    _put(frame, "Not verified", SIZE_CAPTION, CRITICAL, strong=True)
    _put(frame, "Compiling ≠ the same behavior", SIZE_LEAD, INK, strong=True, new=True, before=4)
    _put(frame, "No step runs the program. The projects' own tests were not run: the data keeps "
         "only .java files, without build files or libraries.", SIZE_CAPTION, INK, new=True,
         before=10)  # fmt: skip
    _put(frame, "So the claim is: the rewrite compiles, or adds no new kind of error.",
         SIZE_CAPTION, INK, strong=True, new=True, before=10)  # fmt: skip


def _slide_example(slide, f, heading) -> None:
    _heading(slide, heading, "Apache Geode · DistributedRegion.checkForLimitedOrNoAccess · Guard "
                             "Clauses")  # fmt: skip
    _map(slide, {3}, "example for")
    code = _example()
    if code is None:
        _put(_frame(slide, MARGIN, TOP, CONTENT, Inches(0.6)),
             "[Corpus missing: run build_code_example.py]", SIZE_LEAD, CRITICAL)  # fmt: skip
        return
    before, removed, after, added = code
    width = Inches(8.3)
    bottom = _diff(slide, "Before", before, removed, REMOVED_TINT, CRITICAL, MARGIN,
                   Inches(1.6), width)  # fmt: skip
    _diff(slide, "After", after, added, ADDED_TINT, ADDED, MARGIN, Emu(int(bottom + Inches(0.1))),
          width)  # fmt: skip
    left = Emu(int(MARGIN + width + Inches(0.4)))
    column = Emu(int(WIDTH - MARGIN - left))
    _stat(slide, left, Inches(1.75), column, "−1 level",
          "of nesting: everything after the guard moves up", ADDED, SIZE_STAT, INK)  # fmt: skip
    steps = [
        "The outer condition becomes an early return.",
        "Passed the javac check.",
        "Judged acceptable as-is in the manual review.",
    ]
    top = Inches(3.55)
    for index, text in enumerate(steps, 1):
        _badge(slide, left, top, str(index), ADDED if index == 1 else ACCENT)
        _put(_frame(slide, Emu(int(left + Inches(0.5))), top, Emu(int(column - Inches(0.5))),
                    Inches(1.0)), text, SIZE_CAPTION, INK)  # fmt: skip
        top = Emu(int(top + Inches(1.0)))


def _slide_closing(slide, f, heading) -> None:
    """Gjetjet, pse kanë rëndësi, dhe çka vjen më pas: tri pjesët që kërkoi mentorja.

    Pikat «Why it matters» dhe «Next» janë ato të punimit (6.5 dhe 6.7), jo të reja.
    """
    _background(slide, NAVY)
    _heading(slide, heading, dark=True)
    applied, detected = f["applied"], f["detected"]
    lowest, highest = min(f["recall"]) * 100, max(f["recall"]) * 100
    better = sum(f["best"][s]["mcc"] > f["strategy"][s]["mcc"] for s in f["smells"])
    verdicts = f["refactoring"]["verdicts"]
    passed = (verdicts.get("compiles", 0) + verdicts.get("no_new_errors", 0)) / applied
    findings_w = Inches(5.6)
    _put(_frame(slide, MARGIN, Inches(1.4), findings_w, Inches(0.45)), "Key findings", SIZE_LEAD,
         ACCENT_ON_DARK, strong=True)  # fmt: skip
    answers = [
        f"Rules are precise, but find only {lowest:.0f}–{highest:.0f} of every 100 smells.",
        f"Learned models do better on {_count_word(better, len(f['smells']))}, "
        f"clearly on {_count_word(len(f['sure']), len(f['smells']))}.",
        f"Yes, for {applied / detected:.0%} of places; {passed:.0%} compile or add no new error, "
        "but behavior is not yet proven.",
    ]
    top = Inches(1.95)
    pill_w, pill_h = Inches(0.62), Inches(0.36)
    for index, text in enumerate(answers, 1):
        pill = _shape(slide, MARGIN, top, pill_w, pill_h, ACCENT_ON_DARK)
        pill.adjustments[0] = 0.5
        frame = pill.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _put(frame, f"Q{index}", SIZE_SMALL, NAVY, strong=True, align=PP_ALIGN.CENTER)
        _put(_frame(slide, Emu(int(MARGIN + pill_w + Inches(0.2))), Emu(int(top - Inches(0.02))),
                    Emu(int(findings_w - pill_w - Inches(0.2))), Inches(0.9)), text, Pt(18),
             ON_DARK)  # fmt: skip
        top = Emu(int(top + Inches(0.95)))
    frame = _frame(slide, MARGIN, Emu(int(top + Inches(0.05))), findings_w, Inches(0.6))
    paragraph = frame.paragraphs[0]
    _run(paragraph, "Limits: ", SIZE_SMALL, DANGER_ON_DARK, strong=True)
    _run(paragraph, "one dataset, four smells, Java only; the best model was chosen on the same "
         "predictions, so it is slightly optimistic.", SIZE_SMALL, ON_DARK_SOFT)  # fmt: skip
    left = Emu(int(MARGIN + findings_w + Inches(0.55)))
    width = Emu(int(WIDTH - MARGIN - left))
    for top, title, lines in (
        (Inches(1.4), "Why it matters", [
            "Published thresholds are a starting point: calibrate or learn them for your own "
            "code.",
            "An automatic fix is a compiler-checked proposal that a developer reviews and tests.",
        ]),
        (Inches(3.3), "Future research", [
            "Run the projects' own tests, to go from “compiles” to “behaves the same”.",
            "Resolve names inside a project, to add Move Method and Encapsulate Field.",
            "More smells beyond MLCQ, and reviewers other than the author.",
        ]),
    ):  # fmt: skip
        _put(_frame(slide, left, top, width, Inches(0.45)), title, SIZE_LEAD, ACCENT_ON_DARK,
             strong=True)  # fmt: skip
        _bullet_list(slide, left, Emu(int(top + Inches(0.55))), width, Inches(1.5), lines,
                     SIZE_CAPTION, ON_DARK, ACCENT_ON_DARK, after=6)  # fmt: skip
    frame = _frame(slide, MARGIN, Inches(5.65), CONTENT, Inches(1.1))
    _put(frame, "Thank you  ·  Questions?", Pt(36), ON_DARK, strong=True)
    _put(frame, REPOSITORY.replace("https://", ""), SIZE_CAPTION, ON_DARK_SOFT, new=True, before=2)


def _slide_references(slide, f, heading) -> None:
    """Burimet e cituara në sllajde, në IEEE, në dy kolona që të lexohen nga salla."""
    # Pa nëntitull: autori e hoqi vetë në PowerPoint (23 tetor), dhe ndërtuesi e ndjek.
    _heading(slide, heading)
    half = (len(DECK_SOURCES) + 1) // 2
    column_w = Emu(int((CONTENT - Inches(0.5)) / 2))
    hang = int(Inches(0.42))
    for column, entries in enumerate((DECK_SOURCES[:half], DECK_SOURCES[half:])):
        left = Emu(int(MARGIN + column * (column_w + Inches(0.5))))
        frame = _frame(slide, left, Inches(1.75), column_w, Inches(4.95))
        for offset, (_, text) in enumerate(entries):
            number = column * half + offset + 1
            paragraph = frame.add_paragraph() if offset else frame.paragraphs[0]
            paragraph.space_after = Pt(6)
            properties = paragraph._p.get_or_add_pPr()
            properties.set("marL", str(hang))
            properties.set("indent", str(-hang))
            _run(paragraph, f"[{number}]  ", SIZE_FOOT, ACCENT, strong=True)
            # IEEE shkruan me të pjerrëta titullin e librit dhe emrin e revistës: `*...*`.
            for part, piece in enumerate(text.split("*")):
                if piece:
                    _run(paragraph, piece, SIZE_FOOT, INK).font.italic = part % 2 == 1


LAYOUTS = [_slide_title, _slide_problem, _slide_contribution, _slide_smells, _slide_system,
           _slide_data, _slide_results, _slide_agreement, _slide_engine, _slide_validation,
           _slide_example, _slide_closing, _slide_references]  # fmt: skip
# Hapja dhe mbyllja me sfond të errët; referencat vijnë pas mbylljes, me sfond të çelët.
DARK = {0, LAYOUTS.index(_slide_closing)}


def build_deck(f: dict, long: bool) -> tuple[str, list[tuple[str, list[Line]]], list[str]]:
    """Prezantimi; me `long` shënimet mbajnë edhe fjalitë `Extra`.

    Kthen tekstin e plotë të çdo sllajdi kryesor, me shtesat të shënuara, që fleta e
    përgatitjes të llogarisë kohën e të dy versioneve nga i njëjti burim, dhe titujt e
    sllajdeve rezervë me radhë, që udhëzuesi t'i numërojë siç janë në prezantim.
    """
    deck = Presentation()
    deck.slide_width = WIDTH
    deck.slide_height = HEIGHT
    script = [(title, list(parts)) for title, *parts in _script(f)]
    for index, ((_, parts), layout, heading) in enumerate(
        zip(script, LAYOUTS, _headings(f), strict=True)
    ):
        slide = _blank(deck)
        _speech(slide, _spoken(parts, long))
        layout(slide, f, heading)
        if index:
            _footer(slide, index + 1, dark=index in DARK)
    main = len(deck.slides)
    refactoring, smells, detected, pmd = f["refactoring"], f["smells"], f["detected"], f["pmd"]
    quality = f["quality"]

    # ------------------------------------------------------------------
    # Rezervë: vetëm për pyetjet, nuk thuhen në dhjetë minutat
    # ------------------------------------------------------------------
    backups: list[str] = []

    def backup(title: str, subtitle: str, notes: str):
        slide = _blank(deck)
        backups.append(title)
        _heading(slide, title, subtitle, kicker="Backup · ")
        _note(slide, notes)
        _footer(slide, len(deck.slides))
        return slide

    ml = f["ml"]
    slide = backup("Full detection results", "Precision, recall and MCC for both approaches",
                   "Use when asked for precision or recall per smell.")  # fmt: skip
    _data_table(
        slide,
        ["Smell", "Rules P", "Rules R", "Rules MCC", "Model", "Model P", "Model R", "Model MCC"],
        [
            [
                SMELL_SQ[s],
                f"{f['strategy'][s]['precision']:.3f}",
                f"{f['strategy'][s]['recall']:.3f}",
                f"{f['strategy'][s]['mcc']:.3f}",
                MODEL_EN[ml["per_smell"][s]["best_model"]],
                f"{f['best'][s]['precision']:.3f}",
                f"{f['best'][s]['recall']:.3f}",
                f"{f['best'][s]['mcc']:.3f}",
            ]
            for s in smells
        ],
    )

    intervals = _load("bootstrap_intervals.json")["per_smell"]
    slide = backup("Confidence intervals", "95% bootstrap intervals, resampling whole repositories",
                   "Above zero means the model's advantage is not explained by which projects "
                   "happen to be in the sample.")  # fmt: skip
    _data_table(
        slide,
        ["Smell", "Rules (A)", "Model (B)", "B − A", "Above zero"],
        [
            [
                SMELL_SQ[s],
                "[{low:.3f}, {high:.3f}]".format(**intervals[s]["intervals"]["rules"]),
                "[{low:.3f}, {high:.3f}]".format(**intervals[s]["intervals"][intervals[s]["best_model"]]),
                "[{low:.3f}, {high:.3f}]".format(**intervals[s]["intervals"]["difference_model_minus_rules"]),
                "yes" if intervals[s]["intervals"]["difference_model_minus_rules"]["low"] > 0 else "no",
            ]
            for s in smells
        ],
    )  # fmt: skip

    slide = backup("Every algorithm", "MCC of each classifier, same folds",
                   "If asked whether choosing the best model inflates the result: yes, slightly. "
                   "Even the weakest learning algorithm beats the rules on three of four smells.")  # fmt: skip
    _data_table(
        slide,
        ["Smell", *[MODEL_EN[m] for m in ("logistic", "random_forest", "gradient_boosting")], "Rules (A)"],
        [
            [SMELL_SQ[s]]
            + [f"{ml['per_smell'][s]['models'][m]['mcc']:.3f}"
               for m in ("logistic", "random_forest", "gradient_boosting")]
            + [f"{f['strategy'][s]['mcc']:.3f}"]
            for s in smells
        ],
    )  # fmt: skip

    names = PMD_NAMES
    slide = backup("Comparison with PMD", f"PMD {pmd['pmd_version']} {_cite('PMD Team 2026')} with its default thresholds, same samples",
                   "PMD is never ahead once the confidence interval is considered. PMD has no "
                   "Feature Envy rule.")  # fmt: skip
    _data_table(
        slide,
        ["Smell", "PMD MCC", "Our MCC", "Difference, 95% CI"],
        [
            [
                names.get(key, key),
                f"{entry['pmd']['mcc']:.3f}",
                "—" if entry.get("ours") is None else f"{entry['ours']['mcc']:.3f}",
                "—" if entry.get("difference") is None
                else f"{entry['difference']['low']:+.3f} to {entry['difference']['high']:+.3f}",
            ]
            for key, entry in pmd["by_smell"].items()
        ],
    )  # fmt: skip

    verify = _load("verify_with_project.json")
    slide = backup("Verification in detail", "Which rewrites succeeded and which had problems, per fix and per check",
                   "Isolated files often cannot compile because their neighbours are missing; "
                   "inside the project many more compile fully, and 2 show an error the file "
                   "alone hid.")  # fmt: skip
    _put(_frame(slide, MARGIN, Inches(1.62), CONTENT, Inches(0.3)),
         f"All {detected:,} places, by fix: “pass” = compiles, or adds no new kind of error",
         SIZE_SMALL, SOFT, strong=True)  # fmt: skip
    per_fix = _per_fix(f)
    totals = {key: sum(entry[key] for entry in per_fix.values())
              for key in ("sites", "applied", "passed", "new_errors", "resolved")}  # fmt: skip
    _data_table(
        slide,
        ["Fix", "Places", "Rewritten", "Pass javac", "New error", "Smell removed", "Acceptable*"],
        [
            [_words(name), f"{e['sites']:,}", f"{e['applied']:,}",
             f"{e['passed']:,} ({_pct(e['passed'] / e['applied'], 0)})", f"{e['new_errors']}",
             f"{e['resolved']:,} ({_pct(e['resolved'] / e['applied'], 0)})",
             f"{quality['by_refactoring'][name]['acceptable']:.0%}"]
            for name, e in sorted(per_fix.items(), key=lambda pair: -pair[1]["applied"])
        ]
        + [["Not located again", f"{refactoring['unlocatable']}", "—", "—", "—", "—", "—"]]
        + [["All", f"{totals['sites'] + refactoring['unlocatable']:,}", f"{totals['applied']:,}",
            f"{totals['passed']:,} ({_pct(totals['passed'] / totals['applied'], 0)})",
            f"{totals['new_errors']}",
            f"{totals['resolved']:,} ({_pct(totals['resolved'] / totals['applied'], 0)})",
            _pct(quality["acceptable_reweighted"])]],
        top=Inches(1.95), size=SIZE_SMALL, row=Inches(0.4),
    )  # fmt: skip
    alone, inside = verify["compiled_alone"], verify["compiled_in_project"]
    _put(_frame(slide, MARGIN, Inches(4.55), CONTENT, Inches(0.3)),
         f"The same {verify['rewrites']} rewrites ({verify['files_checked']} files), compiled two ways",
         SIZE_SMALL, SOFT, strong=True)  # fmt: skip
    _data_table(
        slide,
        ["Compiled", "Compile fully", "No new kind of error", "New error", "No verdict"],
        [[label, str(counts.get("compiles", 0)), str(counts.get("no_new_errors", 0)),
          str(counts.get("new_errors", 0)), str(counts.get("not_checked", 0))]
         for label, counts in (("The file alone", alone), ("Inside its project", inside))],
        top=Inches(4.88), size=SIZE_SMALL, row=Inches(0.4),
    )  # fmt: skip
    introduced = refactoring["introduced_smells"]
    newest = max(introduced, key=introduced.get)
    _put(_frame(slide, MARGIN, Inches(6.2), CONTENT, Inches(0.6)),
         f"* Manual review of {quality['per_refactoring']} rewrites per fix; “All” is reweighted by "
         f"how often each fix occurs. {refactoring['verdicts'].get('parses', 0)} rewrites got no javac "
         f"verdict. After "
         f"rewriting, {sum(introduced.values()):,} new smells appeared, most often {_words(newest)} "
         f"({introduced[newest]}).", SIZE_SMALL, SOFT)  # fmt: skip

    slide = backup("Extract Method in the interface", "jsoup · NodeTraversor.traverse, preview before any file is touched",
                   "The extracted method is named 'extracted'; choosing a meaningful name is "
                   "left to the developer.")  # fmt: skip
    path = os.path.join(FIGURES, "nderfaqja_rishkrimi.png")
    picture = slide.shapes.add_picture(path, MARGIN, TOP, width=CONTENT)
    picture.line.color.rgb = FRAME
    picture.line.width = Pt(0.75)

    slide = backup("Comparison with the literature", f"Madeyski & Lewowski {_cite('Madeyski & Lewowski 2023')}, same MLCQ dataset, DS1",
                   "They used random 5-fold cross-validation and tuned hyperparameters; we "
                   "grouped by repository and did not tune. Our results are in the same range.")  # fmt: skip
    _data_table(
        slide,
        ["Smell", "Rules (A)", "Model (B)", "Madeyski & Lewowski 2023"],
        [[SMELL_SQ[s], f"{f['strategy'][s]['mcc']:.3f}", f"{f['best'][s]['mcc']:.3f}",
          MADEYSKI_2023_DS1[s]] for s in smells],
    )  # fmt: skip

    os.makedirs(OUT_DIR, exist_ok=True)
    path = DECK_LONG if long else DECK
    deck.save(path)
    return path, script[:main], backups


# ----------------------------------------------------------------------
# Fleta e përgatitjes
# ----------------------------------------------------------------------
PHRASES = [
    ("Thank you, that is a good question.", "Fillim i sigurt për çdo pyetje; të jep dy sekonda kohë."),
    ("If I understand correctly, you are asking about …", "Kur pyetja është e gjatë ose e paqartë."),
    ("Could you repeat the question, please?", "Kur nuk e dëgjove ose nuk e kuptove."),
    ("We did not measure that directly, but …", "Kur pyetja del jashtë asaj që mate."),
    ("That is a limitation of this study, and it is part of our future work.",
     "Kur të kapin te një kufizim: pranoje, mos u mbro."),
    ("Let me show you the backup slide.", "Kur ke sllajd rezervë për atë pyetje."),
]

PRONUNCIATION = [
    ("code smell", "kod smel"),
    ("refactoring", "ri-FÉK-tëring"),
    ("precondition", "pri-kën-DÍ-shën"),
    ("Matthews correlation coefficient", "MÉ-thjuz ko-rë-LÉJ-shën ko-i-FÍ-shënt"),
    ("repository", "ri-PÓ-zi-to-ri"),
    ("threshold", "THRÉSH-hold (th si te «thikë»)"),
    ("Feature Envy", "FÍ-çër ÉN-vi"),
    ("compiler / compilation", "kom-PÁJ-lër / kom-pi-LÉJ-shën"),
    ("behavior", "bi-HÉJ-vjër"),
    ("precision / recall", "pri-SÍ-zhën / ri-KÓLL"),
    ("confidence interval", "KÓN-fi-dëns ÍN-tër-vël"),
    ("bootstrap", "BÚT-strep"),
    ("syntax tree", "SÍN-teks tri"),
    ("Guard Clauses", "gard KLLÓ-zëz"),
    ("Apache Geode", "ë-PÉ-çi XHÍ-oud"),
]

CONFIDENCE = [
    "Mëso përmendësh vetëm dy fjalitë e para dhe dy të fundit. Fillimi i sigurt e qetëson "
    "zërin, dhe fundi i sigurt lë përshtypjen e fundit; mesin e mban sllajdi.",
    "Provoje me zë të lartë, me kronometër, të paktën pesë herë, dhe të dy versionet. Nëse "
    "i gjati, me pauzat, të del mbi 9½ minuta, hiqi disa fjali të gjelbra që në provë: në "
    "sallë flitet më shpejt nga nervozizmi, por kryesuesi mund të ta ndërpresë te minuta "
    "e dhjetë.",
    "Regjistrohu një herë me telefon dhe dëgjoje. Do të vëresh vetë ku nxiton dhe ku "
    "ngec; një provë para një shoku a familjari vlen sa tri vetëm.",
    "Mos e lexo sllajdin. Shiko titullin, thuaj mesazhin me fjalët e tua, pastaj trego "
    "shifrën. Shënimet folëse janë rrjetë sigurie, jo tekst për t'u lexuar.",
    "Ndalu pas çdo shifre të rëndësishme. Një pauzë prej një sekonde tingëllon si "
    "vetëbesim, jo si harresë.",
    "Para se të dalësh: frymë brenda katër sekonda, jashtë gjashtë, tri herë. Qëndro "
    "drejt, me këmbët të palëvizura, dhe shiko tri fytyra miqësore në pjesë të ndryshme "
    "të sallës.",
    "Nëse humbet fillin, shiko titullin e sllajdit dhe thuaje me zë: «So, the point here "
    "is …». Askush në sallë nuk e di çfarë kishe ndër mend të thoshe.",
    "Nëse të mungon një fjalë anglisht, thuaje më thjeshtë. «Let me say it in a simpler "
    "way» është plotësisht normale në një konferencë ndërkombëtare.",
    "Te pyetjet: dëgjoje deri në fund, ndalu dy sekonda, përgjigju me dy-tri fjali dhe "
    "ndalu. Përgjigja e shkurtër dhe e saktë duket më e sigurt se ajo e gjata.",
    "Kufizimet i ke vetë te sllajdet 10 dhe 12. Kur dikush të pyet për to, nuk të ka kapur: ti i "
    "ke thënë i pari. Pranoje, thuaj pse, dhe çka do të bëje më pas.",
    "Kujto: e ndërtove vetë, dhe askush në sallë nuk e njeh këtë sistem më mirë se ti. "
    "Shumica e prezantuesve janë studentë, dhe salla dëshiron që të dalësh mirë.",
]

CHECKLIST = [
    "Të mërkurën, 21 tetor: shiko agjendën përfundimtare dhe kontrollo radhën, orën dhe "
    "sallën.",
    "Merr të dy versionet në USB, si .pptx dhe si PDF, dhe dërgoji edhe me email vetes.",
    "Arrij te salla I-302 rreth 08:30, para se të fillojë seanca. Hape prezantimin në "
    "kompjuterin e sallës dhe kontrollo që kodi në sllajdin 11 duket njësoj (shkronja "
    "Consolas).",
    "Prezantohu te kryesuesi i seancës (Blerim Zylifu). Pyete sa minuta ka çdo prezantim "
    "dhe si ta sinjalizon kohën. Me 10 minuta merr versionin e gjatë; me më "
    "pak, të shkurtrin.",
    "Ujë pranë, telefoni pa zë, fleta e pyetjeve në xhep.",
    "Me versionin e shkurtër, kur vendi është 7 minuta: nëse në minutën e pestë je ende te "
    "sllajdi 7, kalo sllajdin 11 (shembullin) me një fjali dhe mbyll me sllajdin 12.",
    "Me versionin e gjatë: nëse në minutën e shtatë je ende te sllajdi 8, lëri të gjitha "
    "fjalitë e gjelbra që mbeten dhe vazhdo me tekstin e shkurtër.",
]

# Nga agjenda paraprake që dërguan organizatorët; agjenda përfundimtare vjen më 21 tetor.
CONFERENCE = [
    "15th UBT Annual International Conference on Computer Science and Engineering, 23–25 "
    "tetor 2026. Ceremonia e hapjes është të premten, 23 tetor, në UBT Innovation Campus "
    "– Lipjan. Prezantimi mund të mbahet fizikisht ose online.",
    "Seanca jote: «Session 1 – Software Engineering», e shtunë 24 tetor, 09:00–11:00, "
    "salla I-302 në UBT Campus. Kryesues Blerim Zylifu, bashkëkryesues Lavdim Menxhiqi.",
    "Je i 14-ti nga 17 prezantime në seancë. Email-i thotë 10 minuta për prezantim, por "
    "17 prezantime në 120 minuta lënë rreth 7 për secilin. Prandaj ke dy versione, dhe "
    "zgjedh njërin sipas kohës që jep kryesuesi.",
    "Në të njëjtën seancë ka tri punime për kodin dhe modelet gjuhësore: optimizimi i kodit "
    "Java me LLM, besueshmëria e kodit të gjeneruar nga AI, dhe punimi i mentores «Trust "
    "but Verify». Pyetja «pse jo një LLM?» ka gjasa të dalë; përgjigjen e ke më poshtë.",
    "Agjenda përfundimtare vjen të mërkurën, 21 tetor: kontrollo radhën, orën dhe sallën.",
    "Abstraktet botohen në «Abstract Book» me ISBN dhe DOI në UBT Knowledge Center. "
    "Pyete mentoren nëse kërkohet më vonë edhe punim i plotë (6 faqe, formati Springer).",
]


def build_sheet(f: dict, script: list[tuple[str, list[str]]]) -> str:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = DocPt(11)

    doc.add_heading("Konferenca e UBT-së 2026: përgatitja", level=0)
    doc.add_paragraph(f"{TITLE}\n{AUTHORS}")

    doc.add_heading("Çka dimë për konferencën", level=1)
    for line in CONFERENCE:
        doc.add_paragraph(line, style="List Bullet")

    words = {
        long: sum(_word_count(_spoken(parts, long)) for _, parts in script) for long in (False, True)
    }

    def minutes(long: bool, rate: int) -> str:
        return _clock(round(words[long] / rate * 60))

    doc.add_heading("Dy versione", level=1)
    for line in [
        f"{os.path.basename(DECK)}: versioni i shkurtër, {words[False]} fjalë, rreth "
        f"{minutes(False, WORDS_PER_MINUTE)} me ritëm të qetë. Për rastin kur kryesuesi jep "
        "8 deri 9 minuta.",
        f"{os.path.basename(DECK_LONG)}: versioni i gjatë, {words[True]} fjalë, rreth "
        f"{minutes(True, WORDS_PER_MINUTE)} me ritëm të qetë. Për rastin kur jep vërtet 10 "
        "minuta, siç thotë email-i.",
        "Sllajdet janë të njëjta; ndryshon vetëm teksti folës. I gjati është i shkurtri plus "
        "disa fjali shtesë, që te shënimet e PowerPoint-it janë me shkronja të pjerrëta e "
        "këtu më poshtë me të gjelbër: kur koha nuk mjafton, i lë dhe je te versioni i "
        "shkurtër pa humbur fillin.",
        f"Me ritëm të shpejtë ({FAST_WORDS_PER_MINUTE} fjalë në minutë) i shkurtri zgjat "
        f"{minutes(False, FAST_WORDS_PER_MINUTE)} dhe i gjati "
        f"{minutes(True, FAST_WORDS_PER_MINUTE)}, pa pauzat. Ky ritëm është supozim. Për ta matur tëndin, lexo me zë tekstin e "
        "sllajdit 2 me kronometër: fjalët e tij (te tabela) pjesëtuar me minutat janë "
        "ritmi yt.",
    ]:  # fmt: skip
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Koha për sllajd", level=1)
    doc.add_paragraph(
        f"E llogaritur nga fjalët e shënimeve folëse, me {WORDS_PER_MINUTE} fjalë në minutë, "
        "pa pauzat dhe kalimet mes sllajdeve (zakonisht edhe rreth një minutë gjithsej)."
    )
    headers = ["#", "Sllajdi", "Fjalë",
               "Shkurt: sek.", "Shkurt: gjithsej", "9 min: sek.", "9 min: gjithsej"]  # fmt: skip
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
    # Gjithsej nga fjalët e grumbulluara, jo nga sekondat e rrumbullakuara të çdo sllajdi,
    # që rreshti i fundit të thotë të njëjtën kohë si përmbledhja më lart.
    spoken_so_far = {False: 0, True: 0}
    for index, (title, parts) in enumerate(script, 1):
        row = table.add_row().cells
        row[0].text, row[1].text = str(index), title
        counts = [_word_count(_spoken(parts, long)) for long in (False, True)]
        row[2].text = str(counts[0]) if counts[0] == counts[1] else f"{counts[0]} / {counts[1]}"
        for column, long in ((3, False), (5, True)):
            spoken_so_far[long] += counts[long]
            row[column].text = str(round(counts[long] / WORDS_PER_MINUTE * 60))
            row[column + 1].text = _clock(round(spoken_so_far[long] / WORDS_PER_MINUTE * 60))

    doc.add_heading("Si të marrësh vetëbesim", level=1)
    for line in CONFIDENCE:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Pyetjet e mundshme (përgjigju në anglisht, shkurt)", level=1)
    doc.add_paragraph(
        "Lista e plotë, me përgjigjen shqip dhe sllajdin rezervë për secilën, është te "
        f"{os.path.basename(SPEECH)}."
    )
    for level, (name, _) in LEVELS.items():
        doc.add_heading(f"Niveli {level}: {name}", level=2)
        _qa(doc, [(q.question, q.answer) for q in _question_bank(f) if q.level == level])

    doc.add_heading("Fraza të gatshme për pyetjet", level=1)
    for phrase, when in PHRASES:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(phrase).bold = True
        paragraph.add_run(f"  ({when})")

    doc.add_heading("Shqiptimi i termave", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Termi", "Si thuhet (përafërsisht)"
    for term, sound in PRONUNCIATION:
        row = table.add_row().cells
        row[0].text, row[1].text = term, sound

    doc.add_heading("Dita e konferencës", level=1)
    for line in CHECKLIST:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Teksti i plotë që thua (shënimet e sllajdeve)", level=1)
    doc.add_paragraph(
        "Fjalitë me të gjelbër dhe të pjerrëta thuhen vetëm në versionin 9-minutësh. Pa to, "
        "mbetet saktësisht teksti i versionit të shkurtër."
    )
    for index, (title, parts) in enumerate(script, 1):
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{index}. {title}").bold = True
        paragraph = doc.add_paragraph()
        for position, part in enumerate(parts):
            run = paragraph.add_run(part if position == 0 else " " + part)
            if isinstance(part, Extra):
                run.italic = True
                run.font.color.rgb = DocRGB(*ADDED)

    doc.save(SHEET)
    return SHEET


# ----------------------------------------------------------------------
# Ndihmës për dokumentet Word
# ----------------------------------------------------------------------
def _document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = DocPt(11)
    return doc


def _list(doc, lines: list[str]) -> None:
    for line in lines:
        doc.add_paragraph(line, style="List Bullet")


# Gjerësia e shkrimit në një faqe Letter me kufijtë e parazgjedhur të Word-it, dhe sa zë
# përafërsisht një shkronjë Calibri 11 dhe kufiri i brendshëm i një qelize.
TEXT_WIDTH_CM = 16.4
CHAR_CM = 0.21
CELL_PADDING_CM = 0.5


def _grid(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    """Një tabelë, me gjerësi kolonash në centimetra.

    Word-i u jep kolonave gjerësi të barabarta, dhe një kolonë me fjali të gjata ngushtohet
    sa një kolonë me një numër. Pa `widths`, secila kolonë merr së paku vendin e fjalës së
    saj më të gjatë, që asnjë fjalë të mos thyhet në mes, dhe pjesa që mbetet ndahet sipas
    gjatësisë mesatare të tekstit.
    """
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
    for values in rows:
        for cell, text in zip(table.add_row().cells, values, strict=True):
            cell.text = text
    if widths is None:
        columns = list(zip(headers, *rows, strict=True))
        floor = [max(len(word) for text in column for word in text.split() or [""]) * CHAR_CM
                 + CELL_PADDING_CM for column in columns]  # fmt: skip
        weights = [sum(map(len, column)) / len(column) for column in columns]
        spare = max(TEXT_WIDTH_CM - sum(floor), 0.0)
        widths = [low + spare * weight / sum(weights) for low, weight in zip(floor, weights, strict=True)]
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = DocCm(width)
    doc.add_paragraph()


def _code(doc, lines: list[str]) -> None:
    """Kod me shkronja me gjerësi të njëtrajtshme, që dhëmbëzimi të mbetet i lexueshëm."""
    for line in lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = DocPt(0)
        run = paragraph.add_run(line or " ")
        run.font.name = MONO
        run.font.size = DocPt(9.5)
    doc.add_paragraph()


def _qa(doc, pairs: list[tuple[str, str]]) -> None:
    for question, answer in pairs:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = DocPt(0)
        paragraph.add_run(question).bold = True
        doc.add_paragraph(answer)


# ----------------------------------------------------------------------
# Pyetjet e mundshme nga salla
# ----------------------------------------------------------------------
@dataclass
class Question:
    """Një pyetje e mundshme, me përgjigjen angleze që thuhet dhe atë shqipe.

    Salla e konferencës është kryesisht shqipfolëse, ndaj pyetja mund të vijë në cilëndo
    gjuhë: përgjigjja shqipe nuk është përkthim fjalë për fjalë, por ajo që thuhet po të
    vijë pyetja shqip.
    """

    level: int
    question: str
    pyetja: str
    answer: str
    pergjigjja: str
    slide: str = ""
    backup: str = ""


LEVELS = {
    1: ("Dikush që e dëgjon temën për herë të parë",
        "Nuk e di çka është një erë kodi, as mësimi makinerik. Pyet për kuptimin dhe dobinë."),
    2: ("Student i informatikës",
        "E di çka është një klasë dhe një model, por jo detajet. Pyet si funksionon."),
    3: ("Pedagog",
        "Pyet për metodën: si e provove, pse kështu dhe jo ndryshe, me çka krahasohet."),
    4: ("Recensent i rreptë",
        "Kërkon dobësitë: çka mund të jetë e gabuar, e fryrë ose e pavërtetuar."),
}  # fmt: skip


def _question_bank(f: dict) -> list[Question]:
    """Rreth dyzet pyetje, nga më e thjeshta te më e vështira, me shifrat e punimit."""
    smells, ml = f["smells"], f["ml"]
    refactoring, quality = f["refactoring"], f["quality"]
    applied, detected = f["applied"], f["detected"]
    reasons = refactoring["refused_by_reason"]
    majority = ml["per_smell"]["feature envy"]["models"]["majority"]["accuracy"]
    limit = f["reference"]["thresholds"]["long_method_loc"]
    winners = {s: MODEL_EN[ml["per_smell"][s]["best_model"]].lower() for s in smells}
    by_model: dict[str, list[str]] = {}
    for s in smells:
        by_model.setdefault(winners[s], []).append(SMELL_SQ[s])
    winners_en = "; ".join(f"{name} for {_and(names)}" for name, names in by_model.items())
    winners_sq = "; ".join(f"{name} për {_dhe(names)}" for name, names in by_model.items())
    pricier = sum(f["best"][s]["precision"] < f["strategy"][s]["precision"] for s in smells)
    introduced = f["introduced"]
    newest = _words(max(introduced, key=introduced.get))
    envy = f["reviewers"]["feature envy"]
    agree = [f["reviewers"][s] for s in f["moderate"]]
    changed = sum(e["behaviour"]["changed"] for e in quality["by_refactoring"].values())
    top = [_metric_name(x).split(" (")[0] for x in ml["per_smell"]["blob"]["top_features"][:2]]
    n = len(smells)
    shape, escape = reasons["shape_not_matched"] / detected, reasons["control_flow_escapes"] / detected
    return [
        # --- 1: dikush që e dëgjon për herë të parë ---
        Question(1, "What is a code smell?", "Çka është një erë e kodit?",
                 "A sign that code is badly organized: it works, but it is hard to change. Like a "
                 "smell in the fridge, it does not prove something is rotten, but it tells you to "
                 "look.",
                 "Shenjë që kodi është i organizuar keq: punon, por ndryshohet me vështirësi. Si "
                 "era në frigorifer: nuk vërteton që diçka është prishur, por të thotë të shikosh.",
                 "The problem"),
        Question(1, "If the program works, why does it matter?",
                 "Nëse programi punon, pse ka rëndësi?",
                 "Because most of the cost of software comes after it is written: fixing it, "
                 "changing it, adding features. Smelly code makes every change slower and riskier.",
                 "Sepse shumica e kostos së një programi vjen pasi shkruhet: ndreqjet, ndryshimet, "
                 "veçoritë e reja. Kodi me erë e bën çdo ndryshim më të ngadaltë dhe më të "
                 "rrezikshëm.", "The problem"),
        Question(1, "So is it like a spell checker for code?",
                 "Pra është si një kontrollues drejtshkrimi për kodin?",
                 "A little, yes. It points at the places that are hard to maintain, explains why, "
                 "and for some of them it prepares a fix.",
                 "Pak a shumë, po. Tregon vendet që mirëmbahen me vështirësi, shpjegon pse, dhe për "
                 "disa prej tyre përgatit një ndreqje.", "What we built"),
        Question(1, "Did you use ChatGPT or another AI to fix the code?",
                 "A përdore ChatGPT ose ndonjë AI tjetër për ta ndrequr kodin?",
                 "No. The fixes are fixed program transformations, not generated text. An AI model "
                 "can write code that looks right, but it cannot prove that the change is safe, "
                 "and it can give a different answer each time.",
                 "Jo. Ndreqjet janë transformime të caktuara të programit, jo tekst i gjeneruar. "
                 "Një model AI mund të shkruajë kod që duket i saktë, por nuk e vërteton dot që "
                 "ndryshimi është i sigurt, dhe mund të japë përgjigje tjetër çdo herë.",
                 "What we built"),
        Question(1, "Who decided which code is smelly?",
                 "Kush vendosi cili kod ka erë?",
                 "Twenty-six professional developers, in a public dataset called MLCQ. I did not "
                 "label the data myself, so my own opinion does not shape the results.",
                 "Njëzet e gjashtë programues profesionistë, në një bazë publike të dhënash që "
                 "quhet MLCQ. Të dhënat nuk i etiketova vetë, ndaj mendimi im nuk i ndikon "
                 "rezultatet.", "Data and evaluation"),
        Question(1, "What does machine learning mean here, in simple words?",
                 "Çka do të thotë mësimi makinerik këtu, me fjalë të thjeshta?",
                 "Instead of fixed rules, the computer learns from thousands of examples that "
                 "developers judged, and finds by itself the line between smelly and clean code.",
                 "Në vend të rregullave fikse, kompjuteri mëson nga mijëra shembuj që i gjykuan "
                 "programuesit, dhe e gjen vetë kufirin mes kodit me erë dhe kodit të pastër.",
                 "Rules vs. machine learning"),
        Question(1, "Does it change my code without asking me?",
                 "A ma ndryshon kodin pa më pyetur?",
                 "No. It first shows the change, and it writes it as a separate patch file. "
                 "Changing the files directly needs an explicit option.",
                 "Jo. Fillimisht e tregon ndryshimin, dhe e shkruan si një skedar më vete "
                 "ndryshimesh. Për t'i ndryshuar skedarët drejtpërdrejt duhet një zgjedhje e "
                 "qartë.", "What we built", "Extract Method in the interface"),
        Question(1, "Does it work for Python or other languages?",
                 "A punon për Python ose gjuhë të tjera?",
                 "Not yet. The parser I use supports many languages, but the measurements and the "
                 "fixes are written for Java.",
                 "Ende jo. Parseri që përdor i mbështet shumë gjuhë, por matjet dhe ndreqjet janë "
                 "shkruar për Java."),
        Question(1, "Can I use it? Does it cost anything?",
                 "A mund ta përdor? A kushton diçka?",
                 "Yes, and it is free. It is open source on GitHub under the MIT license, and it "
                 "runs on a normal laptop.",
                 "Po, dhe është falas. Është me burim të hapur në GitHub, me licencën MIT, dhe "
                 "punon në një laptop të zakonshëm.", "Conclusions"),
        Question(1, "What surprised you the most?",
                 "Çka të befasoi më shumë?",
                 "That for Feature Envy even two experts agree no better than chance. So part of "
                 "the difficulty is in the definition of the smell, not only in the tools.",
                 "Që për Feature Envy edhe dy ekspertë pajtohen jo më shumë se me hamendje. Pra një "
                 "pjesë e vështirësisë është te përkufizimi i erës, jo vetëm te mjetet.",
                 "Human agreement"),
        # --- 2: student i informatikës ---
        Question(2, "What is MCC, and why not accuracy?",
                 "Çka është MCC-ja, dhe pse jo saktësia?",
                 f"MCC is a score from minus one to one: zero means guessing, one means perfect. "
                 f"Most code is clean, so on Feature Envy a detector that always says “clean” is "
                 f"{majority:.0%} accurate, but its MCC is zero.",
                 f"MCC-ja është një notë nga minus një deri në një: zero do të thotë hamendje, një "
                 f"përsosmëri. Shumica e kodit është i pastër, ndaj te Feature Envy një detektor që "
                 f"thotë gjithmonë «i pastër» është {majority:.0%} i saktë, por MCC-ja e tij është "
                 "zero.", "Data and evaluation"),
        Question(2, "What are the 26 metrics?", "Cilat janë 26 metrikat?",
                 f"Simple numbers about the code: {len(f['reference']['metrics']['class'])} for "
                 f"classes and {len(f['reference']['metrics']['method'])} for methods. For example "
                 "lines of code, number of methods, how many paths the code has, how deep the "
                 "blocks are nested, and how much a method uses data from other classes.",
                 f"Numra të thjeshtë për kodin: {len(f['reference']['metrics']['class'])} për "
                 f"klasat dhe {len(f['reference']['metrics']['method'])} për metodat. Për shembull "
                 "rreshtat e kodit, numri i metodave, sa rrugë ka kodi, sa thellë futen blloqet, "
                 "dhe sa përdor një metodë të dhëna nga klasa të tjera.", "What we built"),
        Question(2, "Which rules did you use?", "Cilat rregulla përdore?",
                 f"{f['strategies']} detection strategies, mostly from Lanza and Marinescu's book "
                 "Object-Oriented Metrics in Practice, and from Fowler for Long Method and Long "
                 "Parameter List.",
                 f"{f['strategies']} strategji gjetjeje, kryesisht nga libri i Lanza-s dhe "
                 "Marinescu-t, «Object-Oriented Metrics in Practice», dhe nga Fowler-i për Long "
                 "Method dhe Long Parameter List.", "What we built"),
        Question(2, "Which machine learning model was the best?",
                 "Cili model i mësimit makinerik doli më i miri?",
                 f"It depends on the smell: {winners_en}. The backup slide shows all of them.",
                 f"Varet nga era: {winners_sq}. Sllajdi rezervë i tregon të gjithë.",
                 "Rules vs. machine learning", "Every algorithm"),
        Question(2, "What are precision and recall?", "Çka janë precision dhe recall?",
                 "Precision: when the tool says “smell”, how often it is right. Recall: of all the "
                 "real smells, how many it finds. The rules have high precision but low recall.",
                 "Precision: kur mjeti thotë «erë», sa shpesh ka të drejtë. Recall: nga të gjitha "
                 "erërat e vërteta, sa i gjen. Rregullat kanë precision të lartë, por recall të "
                 "ulët.", "Rules vs. machine learning", "Full detection results"),
        Question(2, "Why do the models give more false alarms?",
                 "Pse modelet japin më shumë alarme të rreme?",
                 f"Because they find many more real smells, and the price is more alarms that "
                 f"turn out wrong. On {pricier} of the {n} smells their precision is lower than "
                 "the rules'.",
                 f"Sepse gjejnë shumë më tepër erëra të vërteta, dhe çmimi është më shumë alarme "
                 f"që dalin të gabuara. Te {pricier} nga {n} erërat precision-i i tyre është më i "
                 "ulët se i rregullave.", "Rules vs. machine learning", "Full detection results"),
        Question(2, f"Why were only {applied / detected:.0%} of the places changed?",
                 f"Pse u ndryshuan vetëm {applied / detected:.0%} e vendeve?",
                 f"Because the engine refuses whenever it cannot prove the change is safe. The "
                 f"main reasons: the code does not have the right shape ({shape:.0%}), or a "
                 f"return or break would leave the moved block ({escape:.0%}).",
                 f"Sepse motori refuzon sa herë nuk e vërteton dot që ndryshimi është i sigurt. "
                 f"Arsyet kryesore: kodi nuk ka formën e duhur ({shape:.0%}), ose një return a "
                 f"break do të dilte nga blloku i zhvendosur ({escape:.0%}).",
                 "Automatic refactoring", "Verification in detail"),
        Question(2, "What does “compiles” mean, if the projects cannot be built?",
                 "Çka do të thotë «kompilohet», nëse projektet nuk ndërtohen dot?",
                 "The projects come without their build files, so many files do not compile alone "
                 "even before the change. So I compile before and after, and a change counts only "
                 "if it adds no new kind of error.",
                 "Projektet vijnë pa skedarët e ndërtimit, ndaj shumë skedarë nuk kompilohen të "
                 "vetëm as para ndryshimit. Prandaj kompiloj para dhe pas, dhe një ndryshim pranohet "
                 "vetëm nëse nuk shton asnjë lloj të ri gabimi.", "Automatic refactoring",
                 "Verification in detail"),
        Question(2, "Which three fixes, exactly?", "Cilat janë saktësisht tri ndreqjet?",
                 "Extract Method moves part of a long method into a new method. Guard Clauses "
                 "turns a big if around the whole method into an early return. Introduce "
                 "Parameter Object groups many parameters into one object.",
                 "Extract Method e zhvendos një pjesë të një metode të gjatë në një metodë të re. "
                 "Guard Clauses e kthen një if të madh rreth gjithë metodës në një dalje të "
                 "hershme. Introduce Parameter Object i mbledh shumë parametra në një objekt.",
                 "Automatic refactoring"),
        Question(2, "What is MLCQ?", "Çka është MLCQ?",
                 "A public dataset by Madeyski and Lewowski, 2020: about fifteen thousand reviews "
                 "of real open-source Java code by twenty-six professional developers, for four "
                 "smells.",
                 "Një bazë publike të dhënash nga Madeyski dhe Lewowski, 2020: rreth pesëmbëdhjetë "
                 "mijë vlerësime të kodit të vërtetë Java me burim të hapur, nga njëzet e gjashtë "
                 "programues profesionistë, për katër erëra.", "Data and evaluation"),
        # --- 3: pedagog ---
        Question(3, "How did you avoid data leakage?", "Si e shmange rrjedhjen e të dhënave?",
                 f"I grouped the data by project, with {ml['folds']}-fold grouped cross-validation. "
                 "A project is never in training and testing at the same time, because code from "
                 "one project is very similar.",
                 f"I grupova të dhënat sipas projektit, me provë të kryqëzuar në {ml['folds']} "
                 "grupe. Një projekt nuk është kurrë njëkohësisht në mësim dhe në provë, sepse "
                 "kodi i një projekti ngjan shumë.", "Data and evaluation"),
        Question(3, "How do you know the improvement is not luck?",
                 "Si e di që përmirësimi nuk është fat?",
                 f"I resampled whole projects {f['resamples']:,} times and recomputed the "
                 f"difference each time. For {_and(f['sure'])} the 95 percent interval stays above "
                 f"zero; for {_and(f['unsure'])} it includes zero, so I do not claim it.",
                 f"I rizgjodha projektet e plota {f['resamples']:,} herë dhe e rillogarita "
                 f"ndryshimin çdo herë. Për {_dhe(f['sure'])} intervali 95 për qind mbetet mbi "
                 f"zero; për {_dhe(f['unsure'])} e përfshin zeron, ndaj nuk e pretendoj.",
                 "Rules vs. machine learning", "Confidence intervals"),
        Question(3, "Why not deep learning?", "Pse jo deep learning?",
                 "The data is small and it is a table of numbers, the models must be explainable, "
                 "and everything had to run on a laptop at zero cost. Simple models were already "
                 "enough to beat the rules.",
                 "Të dhënat janë të pakta dhe janë një tabelë numrash, modelet duhet të "
                 "shpjegohen, dhe gjithçka duhej të punonte në një laptop pa kosto. Modelet e "
                 "thjeshta mjaftuan për t'i mundur rregullat.", "Rules vs. machine learning"),
        Question(3, "How does it compare with tools like PMD or SonarQube?",
                 "Si krahasohet me mjete si PMD ose SonarQube?",
                 "I compared with PMD on the same samples: once the confidence interval is "
                 "considered it is never ahead, and for Long Method my rules are clearly better. "
                 "I did not compare with SonarQube; that is a good next step.",
                 "E krahasova me PMD-në mbi të njëjtat mostra: kur merret parasysh intervali i "
                 "besimit nuk del kurrë përpara, dhe për Long Method rregullat e mia janë qartë më "
                 "të mira. Me SonarQube nuk e krahasova; ai është një hap i mirë për më pas.",
                 "Rules vs. machine learning", "Comparison with PMD"),
        Question(3, "Where do the thresholds come from?", "Nga vijnë pragjet?",
                 f"Mostly from Lanza and Marinescu. Four values have no published source, for "
                 f"example {limit:g} lines for Long Method. I say so in the thesis, and I tested "
                 "how sensitive the results are to that one.",
                 f"Kryesisht nga Lanza dhe Marinescu. Katër vlera nuk kanë burim të botuar, për "
                 f"shembull {limit:g} rreshta për Long Method. Këtë e them hapur në punim, dhe "
                 "provova sa ndikon ajo vlerë te rezultatet.", "The problem"),
        Question(3, "How does the tool decide how serious a smell is?",
                 "Si vendos mjeti sa e rëndë është një erë?",
                 "It measures how far each value is past its limit, as a ratio, and takes the "
                 "average. Below 1.5 it is minor, below 2.5 major, above that critical: the same "
                 "scale as MLCQ.",
                 "Mat sa larg kufirit është secila vlerë, si raport, dhe merr mesataren. Nën 1.5 "
                 "është e lehtë, nën 2.5 e mesme, mbi të e rëndë: e njëjta shkallë si në MLCQ."),
        Question(3, "Do the new smells after a fix cancel the benefit?",
                 "A e anulojnë dobinë erërat e reja pas ndreqjes?",
                 f"They are a real cost, and I report them: {sum(introduced.values()):,} new smells "
                 f"appeared, most often {newest}, mostly from Extract Method. But the targeted "
                 f"smell disappeared in {f['resolved']:.0%} of the cases.",
                 f"Janë një kosto e vërtetë, dhe i raportoj: u shfaqën {sum(introduced.values()):,} "
                 f"erëra të reja, më shpesh {newest}, kryesisht nga Extract Method. Por era që "
                 f"synohej u zhduk në {f['resolved']:.0%} të rasteve.", "Automatic refactoring"),
        Question(3, "How did you review the changes by hand?",
                 "Si i rishikove ndryshimet me dorë?",
                 f"I sampled {quality['per_refactoring']} changes for each of the three fixes, "
                 f"{quality['reviewed']} in total, and judged each as acceptable, acceptable after a "
                 "small edit, or rejected. Extract Method is most of the changes, so the total is "
                 "weighted by how common each fix is.",
                 f"Zgjodha {quality['per_refactoring']} ndryshime për secilën nga tri ndreqjet, "
                 f"{quality['reviewed']} gjithsej, dhe secilin e gjykova si të pranueshëm, të "
                 "pranueshëm pas një ndreqjeje të vogël, ose të refuzuar. Extract Method përbën "
                 "shumicën e ndryshimeve, ndaj totali peshohet sipas sa i shpeshtë është secili.",
                 "Automatic refactoring"),
        Question(3, "How fast is it? Can it handle a real project?",
                 "Sa i shpejtë është? A e përballon një projekt të vërtetë?",
                 "On jsoup, about 35 thousand lines in 199 files, the rules run in about two "
                 "seconds on a laptop. The whole corpus of hundreds of projects was analyzed the "
                 "same way.",
                 "Te jsoup-i, rreth 35 mijë rreshta në 199 skedarë, rregullat punojnë rreth dy "
                 "sekonda në një laptop. I gjithë korpusi me qindra projekte u analizua njësoj.",
                 "What we built"),
        Question(3, "What does a developer actually see in the tool?",
                 "Çka sheh në të vërtetë një programues te mjeti?",
                 "For each finding: the smell, how serious it is, the conditions that fired with "
                 "the measured values, and where possible a preview of the fix before any file "
                 "changes. There is also a plain-language mode in Albanian.",
                 "Për çdo gjetje: erën, sa e rëndë është, kushtet që e shkaktuan me vlerat e "
                 "matura, dhe kur mundet një pamje paraprake të ndreqjes para se të preket ndonjë "
                 "skedar. Ka edhe një mënyrë me gjuhë të thjeshtë, shqip.", "What we built",
                 "Extract Method in the interface"),
        # --- 4: recensent i rreptë ---
        Question(4, "Isn't choosing the best model on the same predictions optimistic?",
                 "A nuk është optimiste të zgjedhësh modelin më të mirë mbi të njëjtat parashikime?",
                 f"Yes, slightly, and I say so. That is why I report every algorithm: even the "
                 f"weakest one beats the rules on {f['weakest']} of {n} smells, so the conclusion "
                 "does not depend on the choice.",
                 f"Po, pak, dhe e them hapur. Prandaj i raportoj të gjithë algoritmet: edhe më i "
                 f"dobëti i mund rregullat te {f['weakest']} nga {n} erërat, ndaj përfundimi nuk "
                 "varet nga zgjedhja.", "Conclusions", "Every algorithm"),
        Question(4, "If reviewers disagree, are your labels reliable at all?",
                 "Nëse rishikuesit nuk pajtohen, a janë fare të besueshme etiketat?",
                 f"For Feature Envy, not really: the reviewers' MCC is {envy:.2f}, and I report that "
                 f"as a limit. For the other three it is {min(agree):.2f} to {max(agree):.2f}, so the "
                 "labels carry real signal, and each label is the average of its reviewers.",
                 f"Për Feature Envy, jo vërtet: MCC-ja mes rishikuesve është {envy:.2f}, dhe këtë e "
                 f"raportoj si kufizim. Për tri të tjerat është {min(agree):.2f} deri "
                 f"{max(agree):.2f}, ndaj etiketat kanë sinjal të vërtetë, dhe çdo etiketë është "
                 "mesatarja e rishikuesve të saj.", "Human agreement"),
        Question(4, "How can you call a change safe without running tests?",
                 "Si mund ta quash një ndryshim të sigurt pa ekzekutuar teste?",
                 f"I do not claim that behavior is preserved. I claim two things: the "
                 f"preconditions are proven on the syntax tree, and the result adds no new compile "
                 f"error. In my manual review {changed} Extract Method change still seemed to "
                 "change behavior, which is exactly why I do not claim more.",
                 f"Nuk pretendoj që sjellja ruhet. Pretendoj dy gjëra: kushtet paraprake "
                 f"vërtetohen në pemën e sintaksës, dhe rezultati nuk shton gabim të ri kompilimi. "
                 f"Në rishikimin tim {changed} ndryshim Extract Method dukej se e ndryshonte "
                 "sjelljen, dhe pikërisht prandaj nuk pretendoj më shumë.",
                 "Conclusions"),
        Question(4, "Why tree-sitter, and not a compiler front end that resolves names?",
                 "Pse tree-sitter, dhe jo një pjesë kompajleri që i zgjidh emrat?",
                 "Tree-sitter is fast, it works on code that does not compile, and it needs no "
                 "build setup. The price is that names are not resolved, so the engine refuses "
                 "when a name is ambiguous instead of guessing.",
                 "Tree-sitter është i shpejtë, punon edhe me kod që nuk kompilohet, dhe nuk kërkon "
                 "konfigurim ndërtimi. Çmimi është që emrat nuk zgjidhen, ndaj motori refuzon kur "
                 "një emër është i paqartë, në vend që të hamendësojë.", "What we built"),
        Question(4, "Could the model just learn size or project style instead of smells?",
                 "A mund të mësojë modeli thjesht madhësinë ose stilin e projektit, jo erërat?",
                 f"Grouping by project stops it from memorizing projects. Size does matter for some "
                 f"smells: for Blob the strongest features are {top[0]} and {top[1]}, which fits "
                 "the definition of a class that does too much.",
                 f"Grupimi sipas projektit e pengon të mësojë përmendësh projektet. Madhësia ka "
                 f"rëndësi për disa erëra: për Blob tiparet më të forta janë {top[0]} dhe "
                 f"{top[1]}, që përputhet me përkufizimin e një klase që bën shumë.",
                 "Rules vs. machine learning"),
        Question(4, "How did you handle the imbalance, with so few smelly samples?",
                 "Si e trajtove pabarazinë, me kaq pak mostra me erë?",
                 "Every model that supports it uses balanced class weights. I did not oversample, "
                 "because copying rows could cross the project groups and undo the grouping. And "
                 "the score is MCC, which is not fooled by imbalance.",
                 "Çdo model që e lejon përdor pesha të balancuara për klasat. Nuk shumëfishova "
                 "mostra, sepse kopjimi i rreshtave mund t'i kalonte grupet e projekteve dhe ta "
                 "prishte ndarjen. Dhe nota është MCC, që nuk mashtrohet nga pabarazia.",
                 "Data and evaluation"),
        Question(4, "How do your numbers compare with the literature?",
                 "Si krahasohen numrat e tu me literaturën?",
                 f"They are in the same range as Madeyski and Lewowski 2023 on the same dataset: "
                 f"for Feature Envy they report {MADEYSKI_2023_DS1['feature envy']} and I get "
                 f"{f['best']['feature envy']['mcc']:.2f}. They used a random split and tuned "
                 "models, which usually gives higher numbers, and I grouped by project.",
                 f"Janë në të njëjtin rang me Madeyski dhe Lewowski 2023, mbi të njëjtën bazë: për "
                 f"Feature Envy ata raportojnë {MADEYSKI_2023_DS1['feature envy']} dhe unë marr "
                 f"{f['best']['feature envy']['mcc']:.2f}. Ata përdorën ndarje rastësore dhe "
                 "modele të akorduara, që zakonisht japin numra më të lartë, ndërsa unë ndava "
                 "sipas projektit.", "Rules vs. machine learning", "Comparison with the literature"),
        Question(4, "Why keep Introduce Parameter Object if none of its changes were acceptable?",
                 "Pse e mban Introduce Parameter Object, nëse asnjë ndryshim i tij nuk ishte i "
                 "pranueshëm?",
                 "To report a negative result honestly. It does remove the Long Parameter List "
                 "smell, but a developer gains nothing, because the long list only moves into a "
                 "new class.",
                 "Për ta raportuar me ndershmëri një rezultat negativ. E heq erën Long Parameter "
                 "List, por programuesi nuk fiton asgjë, sepse lista e gjatë thjesht kalon në një "
                 "klasë të re.", "Automatic refactoring"),
        Question(4, "Can someone else reproduce your numbers?",
                 "A mund t'i riprodhojë dikush tjetër numrat e tu?",
                 "Yes. Every number comes from a script and is stored with fixed random seeds and "
                 "pinned versions, so a clean copy of the project can regenerate every table.",
                 "Po. Çdo numër vjen nga një skript dhe ruhet me fara të fiksuara të rastit dhe me "
                 "versione të fiksuara, ndaj një kopje e pastër e projektit mund ta rikrijojë çdo "
                 "tabelë.", "Conclusions"),
        Question(4, "Would this work on industrial code, not open source?",
                 "A do të punonte kjo për kod industrial, jo me burim të hapur?",
                 "Probably with recalibration: MLCQ is open-source code, so the thresholds and the "
                 "models may need new data. That is a limit I state, and the pipeline is scripted "
                 "so it can be rerun on another dataset.",
                 "Ndoshta, por me rikalibrim: MLCQ është kod me burim të hapur, ndaj pragjet dhe "
                 "modelet mund të kenë nevojë për të dhëna të reja. Këtë e them si kufizim, dhe "
                 "rrjedha është me skripte, ndaj mund të riekzekutohet mbi një bazë tjetër.",
                 "Conclusions"),
    ]  # fmt: skip


# ----------------------------------------------------------------------
# Udhëzuesi i fjalimit
# ----------------------------------------------------------------------
@dataclass
class SlideGuide:
    """Çka i duhet folësit për një sllajd, përtej fjalive që thotë."""

    heading: str
    message: str
    shows: list[str]
    point: list[str]
    keywords: str


def _slide_guides(f: dict) -> list[SlideGuide]:
    """Një udhëzim për secilin sllajd kryesor, me radhën e prezantimit.

    Titujt vijnë nga `_headings`, të njëjtët që vë prezantimi, dhe «Çka shihet» përshkruan
    pamjen që vizatojnë funksionet `_slide_*`: kur ndryshon një sllajd, ndryshon edhe këtu.
    """
    smells, ml = f["smells"], f["ml"]
    titles = _headings(f)
    lowest, highest = min(f["recall"]) * 100, max(f["recall"]) * 100
    pairs = "; ".join(
        f"{SMELL_SQ[s]} {f['strategy'][s]['mcc']:.2f} → {f['best'][s]['mcc']:.2f}" for s in smells
    )
    agree = "; ".join(f"{SMELL_SQ[s]} {f['reviewers'][s]:.2f}" for s in smells)
    majority = ml["per_smell"]["feature envy"]["models"]["majority"]["accuracy"]
    applied, detected = f["applied"], f["detected"]
    refused = sum(f["refactoring"]["refused_by_reason"].values())
    chance = _dhe([SMELL_SQ[s] for s in f["chance"]])
    moderate = _dhe([SMELL_SQ[s] for s in f["moderate"]])
    quality = f["quality"]
    found_low, found_high = (bound(f["best"][s]["recall"] for s in smells) * 100
                             for bound in (min, max))  # fmt: skip
    refactoring = f["refactoring"]
    missing = refactoring["unlocatable"]
    verdicts, resolution = refactoring["verdicts"], refactoring["resolution"]
    passed = verdicts.get("compiles", 0) + verdicts.get("no_new_errors", 0)
    verify = _load("verify_with_project.json")
    in_project = verify["compiled_in_project"]
    changed = sum(e["behaviour"]["changed"] for e in quality["by_refactoring"].values())
    by_fix = sorted(refactoring["applied_by_refactoring"].items(), key=lambda pair: -pair[1])
    reasons = sorted(refactoring["refused_by_reason"].items(), key=lambda pair: -pair[1])
    references = len(DECK_SOURCES)
    return [
        SlideGuide(
            heading=titles[0],
            message="Kush je, nga vjen puna, dhe tri gjërat që do të dëgjojnë: si gjenden "
            "erërat, sa saktë, dhe sa mund të ndreqen.",
            shows=[
                "Sfond blu i errët: kështu duken vetëm sllajdi i parë dhe i fundit, që salla ta "
                "ndiejë fillimin dhe mbylljen.",
                "Lart, me shkronja të vogla: emri i konferencës. Në mes: titulli i punimit.",
                f"Poshtë titullit: emrat ({AUTHORS}) dhe UBT. Në fund: seanca dhe data.",
            ],
            point=[
                "Mos e lexo titullin: salla e lexon vetë. Shiko sallën dhe thuaj hyrjen "
                "përmendësh.",
                "Kalo te sllajdi 2 pas fjalisë «…and whether it can also fix them».",
            ],
            keywords="emri · teza në UBT · gjej / sa saktë / ndreq",
        ),
        SlideGuide(
            heading=titles[1],
            message="Erërat e kodit nuk janë gabime, por e bëjnë çdo ndryshim të shtrenjtë. "
            "Mjetet e sotme përdorin numra fiksë dhe nuk ndreqin, prandaj tri pyetje.",
            shows=[
                "Titulli e thotë problemin në një fjali: erërat e bëjnë çdo ndryshim të "
                "shtrenjtë, dhe mjetet e sotme vetëm i tregojnë. Nëntitulli: janë probleme "
                "ndërtimi, jo gabime.",
                "Tri pika: programi punon, por çdo ndryshim kushton më shumë se ç'duhet; mjetet "
                "përdorin pragje fikse, të përcaktuara mbi kod tjetër; e tregojnë problemin, "
                "por rrallë e ndreqin.",
                "Poshtë, tri karta: tri pyetjet e kërkimit, Q1, Q2 dhe Q3. Q1: sa të sakta janë "
                "rregullat e botuara? Q2: a bën më mirë mësimi makinerik me të njëjtat metrika? "
                "Q3: a mund të ndreqen automatikisht pa e prishur kompilimin?",
                "Në fund të çdo karte, një rreth me shkronjë: A (gri) rregullat, B (blu) mësimi "
                "makinerik, C (blu) motori. Janë të njëjtat shkronja që shfaqen te sllajdi 5.",
                "Numrat në kllapa katrore, si [1] te nëntitulli, janë referencat IEEE; lista është "
                f"te sllajdi {len(titles)}. Nuk i lexon me zë.",
                "Nga sllajdi 6 deri te 11, lart djathtas ka tri etiketa Q1 · Q2 · Q3. Ajo me blu "
                "është pyetja për të cilën flet sllajdi, dhe poshtë tyre shkruan çka bën sllajdi "
                "me të, p.sh. «answers Q1 · Q2»: kjo është harta e prezantimit.",
            ],
            point=[
                "Kur thua fjalitë e para, trego tri pikat me radhë.",
                "Te «So I asked three questions», trego tri kartat nga e majta në të djathtë "
                "dhe ndalu një sekondë te secila.",
            ],
            keywords="punon, por kushton · pragje fikse · nuk ndreqin · tri pyetje",
        ),
        SlideGuide(
            heading=titles[2],
            message="Çka e dallon punimin: tri mangësi të studimeve të mëparshme dhe si e trajton "
            "secilën, plus një mjet i hapur.",
            shows=[
                "Dy kolona: majtas, me gri, «Earlier studies»; djathtas, me blu të çelët dhe "
                "shkronja të trasha, «This work». Shigjeta mes tyre e lidh secilën mangësi me "
                "përgjigjen e punimit.",
                f"Rreshti 1: rregullat e botuara merren si të sakta, me pragje të vendosura mbi "
                f"sisteme të tjera {_cite('Lanza & Marinescu 2006')} → këtu maten kundrejt 26 "
                f"programuesve profesionistë të MLCQ-së {_cite('Madeyski & Lewowski 2020')}.",
                "Rreshti 2: rregullat dhe mësimi makinerik rrallë krahasohen mbi të njëjtat të "
                "dhëna, dhe notat varen nga mënyra si ndërtohen të dhënat "
                f"{_cite('Arcelli Fontana et al. 2016', 'Di Nucci et al. 2018')} → këtu të "
                f"njëjtat {f['samples']:,} mostra, e njëjta notë (MCC), ndarje sipas projektit, "
                "me intervale 95%.",
                "Rreshti 3: gjetja dhe ndreqja studiohen veç e veç; ndreqja propozohet dhe "
                f"vendimi i lihet programuesit {_cite('Tsantalis & Chatzigeorgiou 2009')} → këtu "
                "një zinxhir: gjej → rishkruaj → kompilo → mat sërish, me parakushtet të "
                f"vërtetuara më parë {_cite('Opdyke 1992')}.",
                "Poshtë djathtas, me blu: mjeti është i hapur, me ndërfaqe në shqip të thjeshtë, "
                "dhe çdo numër rikrijohet me një skript.",
            ],
            point=[
                "Te «First…», «Second…» dhe «Third…», trego rreshtin përkatës: fillo te kutia gri "
                "majtas, pastaj kalo me dorë te kutia blu djathtas.",
            ],
            keywords="tri mangësi · rregullat maten · e njëjta provë · një zinxhir · mjet i hapur",
        ),
        SlideGuide(
            heading=titles[3],
            message="Katër erërat që mat studimi, të vizatuara: salla e sheh çka është secila "
            "para se të dëgjojë një numër.",
            shows=[
                "Katër karta, një për çdo erë. Në secilën: një vizatim, emri dhe një fjali që e "
                "përkufizon.",
                "Si lexohen vizatimet: kutia është një klasë, vija blu një metodë, pika gri një "
                "fushë (një e dhënë). Kjo shkruhet edhe poshtë kartave.",
                "Blob: një kuti e madhe plot metoda, që merr të dhëna edhe nga dy klasa të vogla. "
                "Data Class: një kuti vetëm me fusha, dhe një klasë tjetër me metodën që punon me "
                "to. Long Method: një kolonë e gjatë rreshtash kodi. Feature Envy: një metodë në "
                "klasën e vet («own class») që zgjatet me tri vija te fushat e një klase tjetër.",
                f"Nëntitulli: përkufizimet sipas Fowler-it {_cite('Fowler 2018')}, dhe të katërat "
                "janë të etiketuara nga programues profesionistë në MLCQ "
                f"{_cite('Madeyski & Lewowski 2020')}.",
            ],
            point=[
                "Trego kartat me radhë, nga e majta në të djathtë, kur thua emrin e secilës.",
                "Te Feature Envy, ndiq me gisht vijat blu nga «own class» te «other class».",
            ],
            keywords="Blob bën shumë · Data Class vetëm të dhëna · Long Method e gjatë · "
            "Feature Envy te klasa tjetër",
        ),
        SlideGuide(
            heading=titles[4],
            message=f"Një mjet, {f['metrics']} metrika, tri qasje: A rregulla, B modele, C një "
            "motor që ndreq dhe çdo ndreqje e kontrollon me javac.",
            shows=[
                "Majtas, rrjedha e mjetit nga lart poshtë: kodi Java → "
                f"{f['metrics']} metrika për çdo klasë e metodë → dy kuti krah për krah, "
                f"A ({f['strategies']} rregulla të botuara) dhe B (mësimi makinerik) → erërat, "
                "secila me kushtet që e shkaktuan → C, motori që rishkruan vetëm atë që "
                "vërteton → kontrolli me javac dhe një patch, pa e prekur origjinalin.",
                "AST është pema e sintaksës: kodi i kthyer në pemë, si analiza gramatikore e një "
                "fjalie. javac është kompajleri i Java-s: programi që e kthen kodin në program "
                "të ekzekutueshëm dhe tregon gabimet.",
                "Djathtas, një copë e zmadhuar e ndërfaqes web gjatë analizës së jsoup-it (një "
                "librari e njohur Java për faqet HTML), shqip: dy gjetje, secila me kushtet e "
                "veta, p.sh. «Rreshtat e kodit të metodës: 53 (problem kur është mbi 30)».",
            ],
            point=[
                "Ndiqe rrjedhën me dorë nga lart poshtë: te «Part A…» kutia A, te «Part B…» "
                "kutia B, te «Part C…» kutia C.",
                "Te «On the right you see the web interface», trego foton djathtas.",
            ],
            keywords=f"{f['metrics']} metrika · A rregulla · B modele · C ndreqje + javac · "
            "ndërfaqja",
        ),
        SlideGuide(
            heading=titles[5],
            message="Etiketat vijnë nga programues profesionistë (MLCQ), dhe prova bëhet e "
            "drejtë duke i ndarë të dhënat sipas projektit dhe duke përdorur MCC.",
            shows=[
                "Lart djathtas ndizen Q1 dhe Q2, me «testing» poshtë tyre: ky sllajd tregon si u "
                "provuan rregullat dhe modelet.",
                f"Dy numra: {f['samples']:,} copa kodi (klasa ose metoda) të gjykuara nga "
                f"programues profesionistë, dhe {f['repositories']} projektet nga vijnë.",
                f"Poshtë majtas, {ml['folds']} kuti me katrorë: çdo kuti është një grup "
                "projektesh. Katër gri janë «learn», ku modeli mëson; e pesta, blu, është «test», "
                "ku provohet mbi projekte që nuk i ka parë. Një projekt ndodhet vetëm në një kuti.",
                f"Poshtë djathtas, kutia «Why the score is MCC, not accuracy»: {majority:.1%} me "
                "të kuqe është saktësia e një detektori që thotë gjithmonë «i pastër» te Feature "
                "Envy. MCC-ja e tij është 0, sa hamendja.",
            ],
            point=[
                "Kur thua numrat, trego numrat e mëdhenj, njërin pas tjetrit.",
                "Te «it learns from four groups», trego kutitë: «mëson nga këto katër, "
                "provohet te kjo».",
                "Te «a detector that always says “clean”», trego numrin e kuq.",
            ],
            keywords=f"MLCQ · 26 programues · {f['samples']:,} mostra · ndarje sipas projektit "
            "· MCC",
        ),
        SlideGuide(
            heading=titles[6],
            message="Rregullat kanë të drejtë kur alarmojnë, por humbasin shumicën e rasteve; "
            f"modelet janë më të mira te të katër erërat, dhe qartë te {_dhe(f['sure'])}.",
            shows=[
                "Grafik me shtylla. Për çdo erë ka dy shtylla krah për krah: gri janë rregullat "
                "e botuara (A), blu është modeli më i mirë (B).",
                "Numri mbi çdo shtyllë është MCC: 0 do të thotë jo më mirë se hamendja, 1 do të "
                "thotë i përsosur. Sa më e lartë shtylla, aq më mirë.",
                f"Numrat, rregullat → modeli: {pairs}.",
                "Nën emrin e çdo ere: «(significant)» do të thotë që fitorja e modelit është e "
                "sigurt statistikisht; «(could be chance)» do të thotë që mund të jetë rastësi.",
                "Djathtas lart, «How to read MCC»: një vijë nga «0 = guessing» te «1 = perfect». "
                "E njëjta shkallë del edhe te sllajdi tjetër.",
                f"Poshtë saj, «Out of 100 real smells»: rregullat gjejnë {lowest:.0f}–"
                f"{highest:.0f} (gri), modelet {found_low:.0f}–{found_high:.0f} (blu), dhe "
                "çmimi: zakonisht më shumë alarme të rreme.",
            ],
            point=[
                "Te «Gray is the published rules, blue is machine learning», trego një shtyllë gri dhe një blu, që "
                "salla ta lidhë ngjyrën me kuptimin.",
                f"Te «they find only {lowest:.0f} to {highest:.0f}», trego numrin gri djathtas; te "
                f"«It finds {found_low:.0f} to {found_high:.0f} out of 100», numrin blu poshtë tij.",
                f"Te «For {_and(f['unsure'])}, it could be luck», trego dy "
                "shtyllat e saj: janë më të ulëtat dhe më afër njëra-tjetrës.",
            ],
            keywords=f"gri = rregullat · blu = modeli · kapin {lowest:.0f}–{highest:.0f}% · "
            f"më mirë te të katërta · {_dhe(f['unsure'])} ndoshta rastësi",
        ),
        SlideGuide(
            heading=titles[7],
            message=f"Për {chance} edhe ekspertët nuk pajtohen me njëri-tjetrin; kufiri nuk është "
            "vetëm algoritmi, por vetë përkufizimi i erës.",
            shows=[
                "Grafik me një shtyllë për çdo erë: sa pajtohen dy rishikues njerëzorë që panë "
                "të njëjtin kod, i matur me MCC. Vija e hollë horizontale është zero, pra "
                "hamendja.",
                f"Numrat: {agree}.",
                f"Shtylla e {chance} është e kuqe dhe nën zero: pajtimi është sa me hamendje. "
                f"Për {moderate} pajtimi është mesatar.",
                "Djathtas: e njëjta shkallë e MCC-së si te sllajdi para tij, numri i madh me të "
                "kuqe («sa një monedhë e hedhur»), dhe fjalia kryesore me shkronja të trasha.",
            ],
            point=[
                f"Te «For {_and([SMELL_SQ[s] for s in f['chance']])}, they agree no more than a "
                "coin flip», trego shtyllën e kuqe.",
                "Fjalinë «If experts do not agree…» thuaje ngadalë, duke parë sallën: është "
                "mesazhi kryesor i pjesës së parë, dhe e ke të shkruar djathtas.",
            ],
            keywords="sa pajtohen njerëzit · Feature Envy sa hamendja · problemi është "
            "përkufizimi",
        ),
        SlideGuide(
            heading=titles[8],
            message=f"Motori ndreq vetëm kur mund ta vërtetojë: rishkruan {applied / detected:.0%} "
            f"të vendeve, refuzon {refused / detected:.0%} me arsye, dhe {missing} vende nuk u "
            f"gjetën dot sërish; të trija bashkë janë {detected:,}.",
            shows=[
                f"Lart: {detected:,} vende në kod ku mund të zbatohej një ndreqje («places in the "
                "code»).",
                f"Shiriti poshtë tij: pjesa blu janë {applied:,} të rishkruarat "
                f"({_pct(applied / detected)}), pjesa gri janë {refused:,} të refuzuarat "
                f"({_pct(refused / detected)}), sepse nuk u vërtetuan të sigurta. Gjatësia e "
                "secilës pjese është sa pjesa e saj.",
                f"Nën shirit, me shkronja të vogla: {missing} vende ({_pct(missing / detected)}) "
                "nuk u gjetën dot sërish në skedar, ndaj aty nuk u provua asnjë ndreqje; djathtas, "
                f"shuma e plotë {applied:,} + {refused:,} + {missing} = {detected:,}, që salla ta "
                "shohë se numrat mblidhen.",
                f"Poshtë majtas, «The {applied:,} rewrites, by fix»: "
                + "; ".join(f"{_words(name)} {count:,}" for name, count in by_fix)
                + ", secila me një shirit sa pjesa e saj.",
                f"Poshtë djathtas, «Why the {refused:,} were refused»: "
                + "; ".join(f"{REASON_SQ_SHORT[name]} {count:,}" for name, count in reasons[:3])
                + f"; tri arsye të tjera {sum(count for _, count in reasons[3:]):,}.",
                "Tri ndreqjet, në nëntitull. Extract Method: një pjesë e metodës së gjatë kalon "
                "në një metodë të re. Guard Clauses: if-i i madh që mbështjell gjithë metodën "
                "bëhet dalje e hershme. Introduce Parameter Object: shumë parametra mblidhen në "
                "një objekt.",
            ],
            point=[
                "Te «it found … places», trego numrin e madh, pastaj pjesën blu të shiritit.",
                "Te «It refused …», trego pjesën gri, pastaj arsyet poshtë djathtas.",
                f"Te «The last {missing}…», trego rreshtin e vogël nën shirit, me shumën.",
                "Te «For me, a refusal is a good result», mos lëviz: thuaje duke parë sallën.",
            ],
            keywords=f"tri ndreqje · {applied:,} + {refused:,} + {missing} = {detected:,} · "
            "refuzimi është i saktë · arsyet",
        ),
        SlideGuide(
            heading=titles[9],
            message="Si u kontrollua çdo rishkrim, në pesë hapa, dhe çka nuk e kontrollon asnjë "
            "hap: sjelljen e programit.",
            shows=[
                "Majtas, pesë hapa me numër, nga lart poshtë. Djathtas në secilin, me blu, sa e "
                "kaluan hapin; poshtë tyre, me të kuqe, sa patën problem.",
                f"1. Para: kushtet vërtetohen mbi pemën e sintaksës {_cite('Opdyke 1992')}: "
                f"{applied:,} u rishkruan, {refused:,} u refuzuan.",
                f"2. Pas: javac mbi skedarin e rishkruar, që nuk lejohet të shtojë lloj të ri "
                f"gabimi: {passed:,} kaluan ({_pct(passed / applied)}); "
                f"{verdicts.get('new_errors', 0)} dështuan dhe nuk shkruhen kurrë; "
                f"{verdicts.get('parses', 0)} mbetën pa verdikt.",
                f"3. Era matet sërish: u hoq te {resolution['resolved']:,} "
                f"({_pct(f['resolved'])}); mbeti te {resolution['persists']:,}; te "
                f"{resolution['unknown']} entiteti nuk u identifikua dot pa mëdyshje.",
                f"4. Mostër: {verify['rewrites']} rishkrime në {verify['files_checked']} skedarë, "
                f"të kompiluara brenda projektit të vet: {in_project['compiles']} kompilohen "
                f"plotësisht, {in_project['new_errors']} kanë gabim të ri, "
                f"{in_project['not_checked']} mbetën pa verdikt.",
                f"5. Lexim me dorë i {quality['reviewed']} rishkrimeve: "
                f"{_pct(quality['acceptable_reweighted'])} të pranueshme (të ripeshuara sipas "
                f"ndreqjes); {changed} e ndryshonte sjelljen.",
                "Djathtas, kutia rozë «Not verified»: kompilimi ≠ sjellja e njëjtë. Asnjë hap nuk "
                "e ekzekuton programin; testet e projekteve nuk u ekzekutuan, sepse të dhënat kanë "
                "vetëm skedarët .java. Prandaj pretendimi është: kompilohet, ose nuk shton lloj të "
                "ri gabimi.",
            ],
            point=[
                "Ndiq hapat me dorë nga 1 te 5, një hap për çdo fjali.",
                "Te «So, to be clear…», trego kutinë rozë djathtas dhe thuaje ngadalë, duke parë "
                "sallën.",
            ],
            keywords="provë para · javac pas · era rimatet · brenda projektit · me dorë · "
            "kompilimi ≠ sjellja",
        ),
        SlideGuide(
            heading=titles[10],
            message="Një shembull i vërtetë: if-i i madh bëhet dalje e hershme, kodi ngjitet një "
            "nivel lart, kompilohet dhe është i pranueshëm.",
            shows=[
                "Lart, «Before»: kodi origjinal nga Apache Geode, metoda "
                "checkForLimitedOrNoAccess. Rreshtat në brez të kuq janë ata që ndryshojnë: if-i "
                "i jashtëm dhe kllapa e tij e fundit.",
                "Poshtë, «After»: kodi pas rishkrimit. Rreshtat në brez të gjelbër janë të rinj: "
                "kushti me «!( … )» dhe «return;».",
                "Komenti «/* … unchanged lines */» do të thotë që një pjesë e gjatë që nuk "
                "ndryshon është mbledhur në një rresht, që kodi të lexohet nga salla.",
                "Djathtas lart, «−1 level» me të gjelbër: thellësia e blloqeve bie me një nivel, "
                "sepse gjithçka pas «rojës» ngjitet lart.",
                "Poshtë saj, tri hapa me numër: kushti bëhet dalje e hershme; e kaloi javac; u "
                "gjykua i pranueshëm ashtu siç është.",
            ],
            point=[
                "Te «Before…», trego brezin e kuq lart.",
                "Te «After…», trego brezin e gjelbër poshtë. Te «moves one level up», trego "
                "«−1 level» djathtas.",
            ],
            keywords="Apache Geode · if i madh → dalje e hershme · një nivel lart · kompilohet",
        ),
        SlideGuide(
            heading=titles[11],
            message="Tri përgjigjet, pse kanë rëndësi në praktikë, dhe çka vjen më pas; kufijtë "
            "me një fjali.",
            shows=[
                "Sfond blu i errët, si sllajdi i parë: salla e kupton që po mbyll.",
                "Majtas, «Key findings»: përgjigjja e secilës pyetje në një fjali, me etiketën Q1, "
                "Q2, Q3. Poshtë tyre, me të kuqe, «Limits»: një bazë të dhënash, katër erëra, "
                "vetëm Java; modeli më i mirë u zgjodh mbi të njëjtat parashikime, ndaj është pak "
                "optimist.",
                "Djathtas lart, «Why it matters»: pragjet e botuara janë pikënisje dhe duhen "
                "kalibruar ose mësuar për kodin tënd; një ndreqje automatike është propozim i "
                "kontrolluar nga kompajleri, që programuesi e rishikon dhe e teston.",
                "Djathtas poshtë, «Future research»: testet e vetë projekteve, që pretendimi të "
                "ngrihet nga «kompilohet» te «sillet njësoj»; zgjidhja e emrave brenda projektit, "
                "për Move Method dhe Encapsulate Field; më shumë erëra se ato të MLCQ-së, dhe "
                "rishikues të tjerë përveç autorit. Të trija janë nga «Puna e ardhshme» e "
                "punimit.",
                "Poshtë: «Thank you · Questions?» dhe adresa e projektit në GitHub.",
            ],
            point=[
                "Te «My three questions have three answers», trego Q1, Q2 dhe Q3 me radhë.",
                "Te «Why does this matter in practice?», trego «Why it matters»; te «And what "
                "comes next?», «Future research».",
                "Pas «I am happy to answer your questions», ndalu, buzëqesh dhe prit. Mos kalo te "
                "referencat: ky sllajd, me adresën e GitHub-it, mbetet në ekran gjatë pyetjeve.",
            ],
            keywords="tri përgjigjet · pragjet janë pikënisje · propozim i kontrolluar · testet "
            "e projekteve · faleminderit",
        ),
        SlideGuide(
            heading=titles[12],
            message="Burimet e cituara në sllajde, në stilin IEEE. Nuk thuhet; e hap vetëm nëse "
            "dikush pyet për një burim.",
            shows=[
                f"{references} referenca në dy kolona, me numrat [1]–[{references}] që dalin në "
                "sllajde, sipas radhës së citimit të parë.",
                "Teksti i secilës është ai i listës IEEE të punimit; punimi ka gjithsej 45 "
                "burime.",
            ],
            point=[
                "Nëse të pyesin nga vjen një rregull ose baza e të dhënave, kalo këtu: "
                f"{', '.join(str(len(titles)))}, Enter.",
            ],
            keywords=f"IEEE · {_cite('Fowler 2018')} Fowler · {_cite('Lanza & Marinescu 2006')} "
            f"Lanza & Marinescu · {_cite('Madeyski & Lewowski 2020')} MLCQ · "
            f"{_cite('Opdyke 1992')} Opdyke",
        ),
    ]  # fmt: skip


def _backup_guides(f: dict) -> dict[str, tuple[str, str]]:
    """Për çdo sllajd rezervë: çka tregon dhe kur hapet."""
    return {
        "Full detection results": (
            "Tabelë: për çdo erë precision (P), recall (R) dhe MCC, për rregullat dhe për "
            "modelin më të mirë.",
            "Kur pyesin për precision ose recall.",
        ),
        "Confidence intervals": (
            "Intervalet 95% për rregullat, për modelin dhe për ndryshimin B − A. Kolona «Above "
            "zero» thotë nëse ndryshimi është i sigurt.",
            "Kur pyesin «si e dini që nuk është fat?».",
        ),
        "Every algorithm": (
            "MCC i çdo algoritmi (logistic regression, random forest, gradient boosting) "
            "krahas rregullave.",
            "Kur pyesin nëse zgjedhja e modelit më të mirë e fryn rezultatin: edhe algoritmi më "
            f"i dobët i mund rregullat te {f['weakest']} nga {len(f['smells'])} erërat.",
        ),
        "Comparison with PMD": (
            f"PMD {f['pmd']['pmd_version']}, me pragjet e veta, mbi të njëjtat mostra: MCC i "
            "PMD-së, MCC ynë dhe intervali i ndryshimit.",
            "Kur pyesin si krahasohet me mjetet ekzistuese.",
        ),
        "Verification in detail": (
            "Dy tabela. E para, sipas ndreqjes: vende, rishkrime, sa e kaluan javac-un, sa "
            "shtuan gabim të ri, te sa u hoq era, dhe sa ishin të pranueshme në rishikim; me "
            "rreshtin e vendeve që nuk u gjetën dot sërish dhe totalin. E dyta: të njëjtat "
            "rishkrime të mostrës, të kompiluara një herë vetëm skedari dhe një herë brenda "
            "projektit. Poshtë: erërat e reja pas rishkrimit.",
            "Kur pyesin cilat rishkrime dolën mirë e cilat patën probleme, ose çka do të thotë "
            "«kompilohet».",
        ),
        "Extract Method in the interface": (
            "Foto e ndërfaqes që tregon një Extract Method te jsoup, para se të preket ndonjë "
            "skedar.",
            "Kur pyesin si duket mjeti ose si e përdor një programues.",
        ),
        "Comparison with the literature": (
            "Rezultatet tuaja krahas Madeyski & Lewowski (2023), mbi të njëjtën bazë MLCQ.",
            "Kur pyesin si krahasohet me punë të tjera.",
        ),
    }


def _say(doc, line: Line) -> None:
    """Fjalia angleze që thuhet, dhe poshtë saj kuptimi shqip që nuk thuhet."""
    extra = isinstance(line, Extra)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = DocPt(0)
    if extra:
        mark = paragraph.add_run("+ ")
        mark.bold = True
        mark.font.color.rgb = DocRGB(*ADDED)
    run = paragraph.add_run(str(line))
    run.font.size = DocPt(12)
    if extra:
        run.italic = True
        run.font.color.rgb = DocRGB(*ADDED)
    meaning = doc.add_paragraph()
    meaning.paragraph_format.left_indent = DocPt(18)
    meaning.paragraph_format.space_after = DocPt(8)
    run = meaning.add_run(line.sq)
    run.italic = True
    run.font.size = DocPt(10)
    run.font.color.rgb = DocRGB(*SOFT)


def _answer(doc, question: Question, backups: list[str], main: int) -> None:
    """Pyetja dhe përgjigjja në të dy gjuhët, me numrin e sllajdit rezervë kur ka."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = DocPt(0)
    paragraph.add_run(question.question).bold = True
    run = paragraph.add_run(f"   ({question.pyetja})")
    run.italic = True
    run.font.color.rgb = DocRGB(*SOFT)
    say = doc.add_paragraph()
    say.paragraph_format.space_after = DocPt(0)
    say.add_run(question.answer).font.size = DocPt(12)
    meaning = doc.add_paragraph()
    meaning.paragraph_format.left_indent = DocPt(18)
    meaning.paragraph_format.space_after = DocPt(4 if question.backup else 10)
    run = meaning.add_run(question.pergjigjja)
    run.italic = True
    run.font.size = DocPt(10)
    run.font.color.rgb = DocRGB(*SOFT)
    if question.backup:
        number = main + 1 + backups.index(question.backup)
        hint = doc.add_paragraph()
        hint.paragraph_format.left_indent = DocPt(18)
        hint.paragraph_format.space_after = DocPt(10)
        run = hint.add_run(f"Sllajdi rezervë {number} · {question.backup}: "
                           f"{', '.join(str(number))}, Enter.")  # fmt: skip
        run.font.size = DocPt(10)
        run.font.color.rgb = DocRGB(*ADDED)


def _bank_section(doc, bank: list[Question], backups: list[str], main: int) -> None:
    """Të gjitha pyetjet, nga më e thjeshta te më e vështira, me si t'u përgjigjesh."""
    doc.add_page_break()
    doc.add_heading("Pyetjet e mundshme: nga më e thjeshta te më e vështira", level=1)
    _list(doc, [
        "Pyetja mund të vijë anglisht ose shqip. Përgjigju në gjuhën e pyetjes: me të zeza "
        "është përgjigjja angleze, me gri ajo shqipe.",
        "Dëgjoje pyetjen deri në fund, ndalu dy sekonda, dhe përgjigju me një-dy fjali. Një "
        "përgjigje e shkurtër dhe e saktë duket më e sigurt se një e gjatë.",
        "Nëse pyetja ka një sllajd rezervë, e ke numrin poshtë përgjigjes: shtyp numrin dhe "
        "Enter, pastaj kthehu te sllajdi i fundit.",
        "Nëse nuk e di përgjigjen, thuaje: «I did not measure that, but…» dhe thuaj çka di. "
        "Askush nuk pret që të dish gjithçka.",
    ])  # fmt: skip
    for level, (name, who) in LEVELS.items():
        doc.add_heading(f"Niveli {level}: {name}", level=2)
        paragraph = doc.add_paragraph()
        paragraph.add_run(who).italic = True
        for question in bank:
            if question.level == level:
                _answer(doc, question, backups, main)


def build_speech_guide(f: dict, script: list[tuple[str, list[Line]]], backups: list[str]) -> str:
    """Çka të thuhet dhe ku të tregohet, sllajd për sllajd.

    Sllajdet janë të rralla me tekst me qëllim, që salla të dëgjojë e jo të lexojë; ky
    dokument është gjysma tjetër: ai i shpjegon folësit çdo element dhe i jep kuptimin e
    çdo fjalie që thotë.
    """
    guides = _slide_guides(f)
    doc = _document()
    doc.add_heading("Çka të flas: prezantimi sllajd për sllajd", level=0)
    doc.add_paragraph(f"{TITLE}\n{AUTHORS}")
    doc.add_paragraph(
        "E shtunë, 24 tetor 2026 · Session 1 – Software Engineering · salla I-302 · "
        "prezantimi i 14-të nga 17."
    )

    doc.add_heading("Si ta përdorësh këtë dokument", level=1)
    _list(doc, [
        "Për çdo sllajd ke: mesazhin në një fjali, çka shihet në sllajd, ku të tregosh, "
        "çka thua, fjalët kyçe dhe pyetjet e mundshme.",
        "Te «Çka thua», fjalia me të zeza është ajo që e thua në anglisht. Poshtë saj, me gri, "
        "është kuptimi në shqip: atë nuk e thua, është që ta dish çka po thua.",
        "Fjalitë me të gjelbër dhe me «+» para thuhen vetëm në versionin 9-minutësh. Pa to "
        "mbetet saktësisht versioni i shkurtër.",
        "Mos e mëso përmendësh çdo fjali. Mëso mesazhin e çdo sllajdi dhe fjalët kyçe, dhe "
        "vetëm hyrjen e mbylljen fjalë për fjalë (në fund të dokumentit).",
        "Lexoje një herë të gjithin me zë, me sllajdet hapur pranë, dhe trego me gisht atë që "
        "thotë «Ku të tregosh».",
        "Lart djathtas, nga sllajdi 6 te 11, etiketat Q1 · Q2 · Q3 tregojnë për cilën pyetje "
        "flet sllajdi; ajo me blu është pyetja e tanishme. Nuk ke nevojë ta thuash: salla e "
        "sheh vetë.",
    ])  # fmt: skip

    doc.add_heading("Historia e prezantimit në një faqe", level=1)
    doc.add_paragraph(
        "Prezantimi është një histori me katër pjesë: problemi, kontributi dhe erërat (1–4), "
        "çka ndërtove dhe si e provove (5–6), çka gjete dhe si e verifikove (7–11), dhe "
        "përfundimet me referencat (12–13). Nëse e kupton këtë tabelë, e di prezantimin."
    )
    rows, totals = [], {False: 0, True: 0}
    for index, ((_, parts), guide) in enumerate(zip(script, guides, strict=True), 1):
        cells = [str(index), guide.heading if index > 1 else "Titulli", guide.message]
        for long in (False, True):
            totals[long] += _word_count(_spoken(parts, long))
            cells.append(_clock(round(totals[long] / WORDS_PER_MINUTE * 60)))
        rows.append(cells)
    _grid(doc, ["#", "Sllajdi", "Mesazhi", "Shkurt: deri këtu", "9 min: deri këtu"], rows,
          widths=[0.8, 5.0, 7.0, 1.8, 1.8])  # fmt: skip

    doc.add_heading("Para se të fillosh", level=1)
    verification = len(script) + 1 + backups.index("Verification in detail")
    keys = [key for key, _ in script]
    closing = keys.index("Conclusions") + 1
    _list(doc, [
        "Nise prezantimin me pamjen e folësit: Alt+F5, ose Slide Show → Presenter View. Salla "
        "sheh sllajdin; ti sheh edhe shënimet dhe sllajdin e radhës.",
        "Përpara: shigjeta djathtas, hapësira ose klikuesi. Prapa: shigjeta majtas.",
        f"Për të hapur një sllajd rezervë, shtyp numrin e tij dhe Enter (p.sh. "
        f"{', '.join(str(verification))}, Enter për sllajdin {verification}). Për t'u kthyer "
        f"te mbyllja: {', '.join(str(closing))}, Enter.",
        "Qëndro anash ekranit, jo para tij, me trup nga salla. Trego me dorë ose me laser, "
        "pastaj kthehu nga salla para se të flasësh.",
        "Shënimet në PowerPoint janë rrjetë sigurie: shiko fjalët kyçe, jo gjithë tekstin.",
    ])  # fmt: skip

    bank = _question_bank(f)
    for index, ((key, parts), guide) in enumerate(zip(script, guides, strict=True), 1):
        doc.add_page_break()
        doc.add_heading(f"Sllajdi {index} · {guide.heading}", level=1)
        paragraph = doc.add_paragraph()
        paragraph.add_run("Mesazhi: ").bold = True
        paragraph.add_run(guide.message)
        short, full = (round(_word_count(_spoken(parts, long)) / WORDS_PER_MINUTE * 60)
                       for long in (False, True))  # fmt: skip
        doc.add_paragraph(f"Koha: rreth {short} sekonda në versionin e shkurtër, {full} në atë "
                          "9-minutësh.")  # fmt: skip
        doc.add_heading("Çka shihet në sllajd", level=2)
        _list(doc, guide.shows)
        doc.add_heading("Ku të tregosh", level=2)
        _list(doc, guide.point)
        doc.add_heading("Çka thua", level=2)
        for line in parts:
            _say(doc, line)
        doc.add_heading("Nëse të ngec mendja", level=2)
        doc.add_paragraph(guide.keywords)
        related = [question for question in bank if question.slide == key]
        if related:
            doc.add_heading("Nëse të pyesin për këtë sllajd", level=2)
            for question in related:
                _answer(doc, question, backups, len(script))

    doc.add_page_break()
    doc.add_heading("Sllajdet rezervë: vetëm për pyetjet", level=1)
    doc.add_paragraph(
        "Këto nuk i tregon gjatë prezantimit. I hap vetëm kur një pyetje i kërkon: shtyp "
        "numrin dhe Enter, përgjigju, pastaj kthehu te sllajdi i fundit."
    )
    explained = _backup_guides(f)
    for offset, title in enumerate(backups, len(script) + 1):
        shows, when = explained[title]
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = DocPt(0)
        paragraph.add_run(f"Sllajdi {offset} · {title}").bold = True
        doc.add_paragraph(shows)
        paragraph = doc.add_paragraph()
        paragraph.add_run("Kur e hap: ").italic = True
        paragraph.add_run(when)

    _bank_section(doc, bank, backups, len(script))

    doc.add_heading("Hyrja dhe mbyllja, përmendësh", level=1)
    doc.add_paragraph(
        "Këto janë fjalitë e vetme që ia vlen t'i mësosh fjalë për fjalë: fillimi i sigurt e "
        "qetëson zërin, dhe mbyllja e sigurt lë përshtypjen e fundit."
    )
    for line in script[0][1][:2] + script[closing - 1][1][-2:]:
        _say(doc, line)

    doc.save(SPEECH)
    return SPEECH


# ----------------------------------------------------------------------
# Libri: punimi nga e para
# ----------------------------------------------------------------------
REASON_SQ = {
    "shape_not_matched": "Kodi nuk ka formën që i duhet ndreqjes, p.sh. metoda nuk është e "
    "gjitha një if.",
    "control_flow_escapes": "Një return, break ose continue del jashtë pjesës që do të "
    "nxirrej: në metodën e re do të kishte kuptim tjetër.",
    "not_definitely_assigned": "Pjesa e nxjerrë lexon një variabël që mund të mos ketë marrë "
    "vlerë në çdo rrugë.",
    "multiple_outputs": "Pjesa ndryshon më shumë se një variabël që përdoret më pas, ndërsa "
    "një metodë Java kthen vetëm një vlerë.",
    "ambiguous_overload": "Ka disa metoda me të njëjtin emër, dhe nga pema nuk dihet cila "
    "thirret.",
    "unresolved_name": "Një emër në kod nuk zgjidhet dot pa e njohur gjithë projektin.",
    "possible_side_effect": "Një shprehje mund të ndryshojë diçka jashtë saj, dhe zhvendosja "
    "do ta ndryshonte radhën.",
}

VERDICT_SQ = {
    "compiles": "kompilohet plotësisht",
    "no_new_errors": "nuk shton asnjë lloj të ri gabimi; skedari nuk kompilohej as më parë, "
    "sepse i mungojnë fqinjët",
    "new_errors": "shton një lloj të ri gabimi, ndaj nuk shkruhet kurrë",
    "parses": "javac nuk u arrit dot, ndaj u kontrollua vetëm sintaksa",
}

TERMS = [
    ("code smell", "erë e kodit", "Shenjë se kodi është ndërtuar në një mënyrë që e bën të "
     "vështirë për t'u ndryshuar. Nuk është gabim."),
    ("refactoring", "refaktorim, ndreqje", "Ndryshim i strukturës së kodit pa ia ndryshuar "
     "sjelljen."),
    ("metric", "metrikë", "Një numër që përshkruan kodin: rreshta, kompleksitet, lidhje."),
    ("threshold", "prag", "Vlera mbi ose nën të cilën një metrikë konsiderohet problem."),
    ("detection strategy", "rregull gjetjeje", "Disa kushte mbi metrika, të lidhura me DHE/OSE."),
    ("label", "etiketë", "Përgjigjja e njeriut: kjo mostër ka erë apo jo."),
    ("sample", "mostër", "Një klasë ose metodë e vlerësuar nga një njeri."),
    ("repository", "projekt, depo", "Dosja e një projekti, zakonisht në GitHub."),
    ("corpus", "korpus", "Të gjitha projektet e shkarkuara mbi të cilat punon mjeti."),
    ("parser", "parser", "Programi që e lexon tekstin e kodit dhe e kthen në pemë."),
    ("AST, syntax tree", "pema e sintaksës", "Kodi i paraqitur si pemë, si analiza "
     "gramatikore e një fjalie."),
    ("compiler, javac", "kompajler", "Programi që e kthen kodin Java në program dhe tregon "
     "gabimet."),
    ("precondition", "kusht paraprak", "Kushti që duhet vërtetuar para se një ndreqje të jetë "
     "e sigurt."),
    ("refusal", "refuzim", "Motori nuk ndreq, dhe thotë pse. Rezultat i saktë, jo dështim."),
    ("patch", "skedari i ndryshimeve", "Skedar që tregon çdo rresht që hiqet dhe shtohet, pa e "
     "prekur kodin."),
    ("classifier, model", "model", "Program që mëson nga shembuj dhe jep përgjigje për raste "
     "të reja."),
    ("baseline", "pikë krahasimi", "Zgjidhja më e thjeshtë që çdo model duhet ta kalojë."),
    ("confusion matrix", "matrica e konfuzionit", "Tabela 2 × 2 e përgjigjeve të sakta dhe të "
     "gabuara."),
    ("precision", "sa shpesh ka të drejtë", "Kur thotë «ka erë», sa herë ka vërtet."),
    ("recall", "sa nga problemet gjen", "Nga erërat e vërteta, sa i gjen."),
    ("F1", "nota e përbashkët", "Një notë që bashkon precision dhe recall."),
    ("MCC", "nota e saktësisë", "Nga −1 deri në 1; 0 = sa hamendja. Nuk mashtrohet kur "
     "shumica e kodit është e pastër."),
    ("data leakage", "rrjedhje e të dhënave", "Kur modeli provohet mbi diçka që e ka parë gjatë "
     "mësimit."),
    ("fold, cross-validation", "grup, provë e kryqëzuar", "Të dhënat ndahen në grupe; mëson nga "
     "të tjerat, provohet te njëri, me radhë."),
    ("bootstrap", "rizgjedhje", "Rizgjedhja e rastësishme me kthim, shumë herë, për të parë sa "
     "lëviz rezultati."),
    ("confidence interval", "interval besimi", "Kufijtë brenda të cilëve ndodhet vlera e vërtetë "
     "me 95% siguri."),
    ("inter-rater agreement", "pajtimi i rishikuesve", "Sa pajtohen dy njerëz që vlerësojnë të "
     "njëjtën gjë."),
    ("seed", "farë", "Numri nga i cili nisin numrat e rastit; i njëjti numër jep të njëjtin "
     "rezultat."),
    ("severity", "rëndësia", "Sa e rëndë është era: minor, major, critical."),
    ("open source", "burim i hapur", "Kodi është publik dhe mund ta përdorë kushdo."),
]  # fmt: skip

SOURCES = [
    "Fowler, M. (1999; 2nd ed. 2018). Refactoring: Improving the Design of Existing Code. "
    "Addison-Wesley. Erërat e kodit dhe refaktorimet.",
    "Lanza, M. & Marinescu, R. (2006). Object-Oriented Metrics in Practice. Springer. "
    "Rregullat e gjetjes dhe pragjet.",
    "Chidamber, S. R. & Kemerer, C. F. (1994). A Metrics Suite for Object Oriented Design. "
    "IEEE TSE. WMC, CBO, DIT, NOC, RFC, LCOM.",
    "McCabe, T. J. (1976). A Complexity Measure. IEEE TSE. Kompleksiteti ciklomatik (CC).",
    "Bieman, J. M. & Kang, B.-K. (1995). Cohesion and Reuse in an Object-Oriented System. "
    "TCC.",
    "Madeyski, L. & Lewowski, T. (2020). MLCQ: Industry-Relevant Code Smell Data Set. Baza e "
    "të dhënave.",
    "Madeyski, L. & Lewowski, T. (2023), mbi të njëjtën bazë MLCQ. Krahasimi me literaturën.",
]


def _metric_name(feature: str) -> str:
    """`c_CLOC` → `CLOC (Rreshtat e kodit të klasës)`: prefiksi thotë vetëm nivelin."""
    name = feature.split("_", 1)[1] if feature[:2] in ("c_", "m_") else feature
    entry = glossary.metric(name)
    return f"{name} ({entry['name'].lower()})" if entry else name


# Shembujt e vegjël të librit. Nuk janë kod i korpusit: janë shkruar që të llogariten me
# dorë, dhe ndërtuesi i mat me vetë mjetin, që numrat në tekst të jenë ata që jep mjeti.
EXAMPLE_SHOP = """class Shop {
    private int discount;

    int total(int[] prices, boolean member) {
        int sum = 0;
        for (int p : prices) {
            if (p > 0) {
                sum += p;
            }
        }
        if (member && sum > 100) {
            sum = sum - discount;
        }
        return sum;
    }
}"""

EXAMPLE_ACCOUNT = """public class Account {
    private int balance;
    private String owner;
    private String email;

    public void deposit(int amount) {
        balance = balance + amount;
    }

    public int report() {
        return balance * 2;
    }

    public void notify(String text) {
        System.out.println(email + text);
    }
}"""

EXAMPLE_INVOICE = """public class Customer {
    private String name;
    private String street;
    private String city;
    private String zip;
    private String country;
    private String phone;

    public String getName() { return name; }
    public String getStreet() { return street; }
    public String getCity() { return city; }
    public String getZip() { return zip; }
    public String getCountry() { return country; }
    public String getPhone() { return phone; }
}

class Invoice {
    private Customer customer;
    private int amount;

    String label() {
        return customer.getName() + " " + customer.getStreet() + " " + customer.getCity()
            + " " + customer.getZip() + " " + customer.getCountry() + " "
            + customer.getPhone() + " " + amount;
    }
}"""


def _measured(source: str) -> tuple[dict[str, dict], dict[str, dict], list]:
    """Metrikat dhe erërat që jep vetë mjeti për një shembull të vogël."""
    project = analyze_source(source)
    classes: dict[str, dict] = {}
    methods: dict[str, dict] = {}
    for unit in project.units:
        for cls in unit.classes:
            classes[cls.name] = cls.metrics
            for method in cls.methods:
                methods[f"{cls.name}.{method.name}"] = method.metrics
    return classes, methods, detect_all(project)


def _chapter(doc, number: int, title: str, intro: str) -> None:
    doc.add_page_break()
    doc.add_heading(f"{number}. {title}", level=1)
    paragraph = doc.add_paragraph()
    paragraph.add_run(intro).italic = True


def _idea(doc, text: str) -> None:
    """Ideja kryesore e një pjese, që lexuesi ta gjejë kur e rilexon."""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("Ideja kryesore: ")
    run.bold = True
    run.font.color.rgb = DocRGB(*ACCENT)
    paragraph.add_run(text)


def _analogy(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("Si në jetën e përditshme: ")
    run.bold = True
    run.font.color.rgb = DocRGB(*ADDED)
    paragraph.add_run(text).italic = True


def _check(doc, pairs: list[tuple[str, str]]) -> None:
    """Pyetjet e vetëkontrollit në fund të kapitullit, me përgjigjen poshtë secilës."""
    doc.add_heading("Provo veten", level=2)
    for question, answer in pairs:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = DocPt(0)
        paragraph.add_run(question).bold = True
        reply = doc.add_paragraph()
        reply.paragraph_format.left_indent = DocPt(18)
        run = reply.add_run(answer)
        run.italic = True
        run.font.color.rgb = DocRGB(*SOFT)


def build_book(f: dict) -> str:
    """Gjithë punimi i shpjeguar nga e para, si një libër, për dikë që nuk di asgjë.

    Çdo shifër lexohet nga `data/results/` si te sllajdet; shembujt e vegjël maten me vetë
    mjetin gjatë ndërtimit; dhe çdo term merr emrin nga fjalori i thjeshtë i sistemit, që
    libri, ndërfaqja dhe CLI-ja të mos e shpjegojnë të njëjtën gjë në tri mënyra.
    """
    smells, ml = f["smells"], f["ml"]
    reference, dataset, pmd = f["reference"], f["dataset"], f["pmd"]
    refactoring, quality = f["refactoring"], f["quality"]
    applied, detected = f["applied"], f["detected"]
    vocab = glossary.vocabulary()
    intervals = _load("bootstrap_intervals.json")["per_smell"]
    verify = _load("verify_with_project.json")
    blob, envy = f["strategy"]["blob"], f["strategy"]["feature envy"]
    envy_total = envy["tp"] + envy["fp"] + envy["tn"] + envy["fn"]
    majority = ml["per_smell"]["feature envy"]["models"]["majority"]["accuracy"]
    lowest, highest = min(f["recall"]) * 100, max(f["recall"]) * 100
    found = [f["best"][s]["recall"] * 100 for s in smells]
    precision = [f["strategy"][s]["precision"] for s in smells]
    q, t = reference["quantifiers"], reference["thresholds"]
    by_type = refactoring["applied_by_refactoring"]
    refused = sum(refactoring["refused_by_reason"].values())
    winners = {s: ml["per_smell"][s]["best_model"] for s in smells}
    shop_classes, shop_methods, _ = _measured(EXAMPLE_SHOP)
    account_classes, _, _ = _measured(EXAMPLE_ACCOUNT)
    invoice_classes, invoice_methods, invoice_smells = _measured(EXAMPLE_INVOICE)
    total, label = shop_methods["Shop.total"], invoice_methods["Invoice.label"]

    doc = _document()
    doc.add_heading("Punimi nga e para", level=0)
    doc.add_paragraph(f"{TITLE}\n{AUTHORS}")
    doc.add_paragraph(
        "Ky dokument e shpjegon gjithë punimin për dikë që nuk di asgjë për programimin, as "
        "për mësimin makinerik. Nis nga «çka është një program» dhe arrin deri te çdo numër i "
        "prezantimit. Lexoje me radhë: çdo kapitull mbështetet te ai para tij."
    )
    doc.add_heading("Si ta lexosh", level=2)
    _list(doc, [
        "Çdo kapitull nis me një fjali që thotë çka do të mësosh, dhe mbaron me «Provo veten»: "
        "pyetje me përgjigjen poshtë. Mbuloje përgjigjen dhe provo ta thuash vetë.",
        "«Ideja kryesore» është fjalia që duhet ta mbash mend nga ajo pjesë. «Si në jetën e "
        "përditshme» është një krahasim që e bën idenë më të lehtë.",
        "Çdo numër në këtë dokument vjen nga skedarët e rezultateve të punimit, si te sllajdet. "
        "Shembujt e vegjël me kod janë matur me vetë mjetin gjatë ndërtimit të dokumentit.",
        "Fjalët teknike shpjegohen kur dalin për herë të parë, dhe i gjen të gjitha te fjalori "
        "në fund.",
    ])  # fmt: skip
    doc.add_heading("Harta e dokumentit", level=2)
    _grid(doc, ["Pjesa", "Kapitujt", "Çka mëson"], [
        ["Bazat", "1–4", "Çka është kodi, çka është një erë, çka është refaktorimi, si e lexon "
                          "mjeti kodin"],
        ["Gjetja e erërave", "5–11", "Metrikat, rregullat, mësimi makinerik, të dhënat, si matet "
                                     "saktësia, rezultatet, pajtimi i njerëzve"],
        ["Ndreqja e erërave", "12–13", "Si punon motori i refaktorimit dhe çka arriti"],
        ["Mjeti dhe shkenca", "14–17", "Si është ndërtuar sistemi, pse rezultatet janë të "
                                       "besueshme, si u bë punimi hap pas hapi, kufizimet"],
        ["Për konferencën", "18–21", "Lidhja me sllajdet, pyetjet e mundshme, fjalori, burimet"],
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 1, "Çka është një program, dhe pse kodi duhet të jetë i rregullt",
             "Në këtë kapitull mëson fjalët bazë: kod, Java, klasë, metodë, fushë, projekt.")
    doc.add_paragraph(
        "Një program është një varg udhëzimesh për kompjuterin. Njeriu i shkruan këto udhëzime "
        "në një gjuhë programimi, si tekst. Ky tekst quhet kod burimor (source code). Java është "
        "një nga gjuhët më të përdorura në botë: me të shkruhen sisteme bankash, servera, "
        "aplikacione Android dhe shumë mjete për kompani."
    )
    doc.add_heading("Pesë fjalë Java", level=2)
    doc.add_paragraph("Ky është një program i vogël Java, me shpjegimin në çdo rresht:")
    _code(doc, [
        "class BankAccount {                   // KLASË: një «dosje» me të dhëna dhe veprime",
        "    private double balance;           // FUSHË: një e dhënë brenda klasës",
        "",
        "    void deposit(double amount) {     // METODË: një veprim; amount është PARAMETËR",
        "        balance += amount;            // veprimi: bilanci rritet me shumën",
        "    }",
        "",
        "    double getBalance() {             // një metodë që KTHEN një vlerë",
        "        return balance;",
        "    }",
        "}",
    ])  # fmt: skip
    _list(doc, [
        "Klasa është si një formular: ka rubrikat e veta (fushat) dhe butonat e vet (metodat). "
        "Nga një klasë mund të krijohen shumë objekte, si shumë llogari bankare nga i njëjti "
        "formular.",
        "Fusha është një e dhënë që klasa e mban, p.sh. bilanci.",
        "Metoda është një veprim. Merr hyrje (parametrat), bën diçka, dhe ndonjëherë kthen një "
        "vlerë me return. Një metodë «void» nuk kthen asgjë.",
        "Një projekt (repository) është një dosje me qindra ose mijëra klasa, zakonisht në "
        "GitHub. «Burim i hapur» (open source) do të thotë që kodi është publik.",
    ])  # fmt: skip
    doc.add_heading("Pse ka rëndësi që kodi të jetë i rregullt", level=2)
    doc.add_paragraph(
        "Një program nuk shkruhet një herë dhe pastaj harrohet. Gjatë viteve ai ndryshohet "
        "vazhdimisht: ndreqen gabime, shtohen veçori, ndryshojnë kërkesat. Kodi lexohet shumë "
        "më shpesh se sa shkruhet, zakonisht nga njerëz që nuk e shkruan vetë. Nëse kodi është "
        "i rregullt, ndryshimi është i shpejtë dhe i sigurt. Nëse është i ngatërruar, çdo "
        "ndryshim zgjat dhe rrezikon të prishë diçka tjetër. Kjo punë e vazhdueshme quhet "
        "mirëmbajtje (maintenance)."
    )
    _analogy(doc, "Një kuzhinë e rregullt dhe një kuzhinë kaotike mund të gatuajnë të njëjtën "
                  "gjellë. Por në atë kaotike, çdo herë që do të gatuash diçka të re, humbet "
                  "kohë duke kërkuar.")  # fmt: skip
    _idea(doc, "Programi që punon nuk është domosdoshmërisht program i mirë: rëndësi ka edhe sa "
               "lehtë ndryshohet.")  # fmt: skip
    _check(doc, [
        ("Çka është kodi burimor?", "Teksti që shkruan njeriu në një gjuhë programimi, si Java."),
        ("Çka është dallimi mes një fushe dhe një metode?",
         "Fusha është një e dhënë që mban klasa; metoda është një veprim që bën klasa."),
        ("Pse kodi duhet të jetë i rregullt, nëse programi punon?",
         "Sepse kodi ndryshohet vazhdimisht, dhe kodi i ngatërruar e bën çdo ndryshim të "
         "ngadaltë dhe të rrezikshëm."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 2, "Erërat e kodit",
             "Në këtë kapitull mëson çka është një «code smell», pse nuk është gabim, dhe katër "
             "erërat që mat punimi.")
    doc.add_paragraph(
        "Kur hap frigoriferin dhe ndien erë, nuk e di ende çka është prishur, por e di që duhet "
        "të shikosh. Në programim, «erë e kodit» (code smell) është po kjo: një shenjë që diçka "
        "në ndërtimin e kodit nuk shkon. Programi punon, por kodi është i ndërtuar në një mënyrë "
        "që e bën çdo ndryshim të ardhshëm më të vështirë. Termin e bënë të njohur Kent Beck dhe "
        "Martin Fowler te libri «Refactoring» (1999)."
    )
    _grid(doc, ["", "Gabim (bug)", "Erë e kodit (code smell)"], [
        ["Programi punon?", "Jo: jep rezultat të gabuar", "Po"],
        ["Ku dëmton?", "Te përdoruesi, sot", "Te programuesi, sa herë ndryshon kodin"],
        ["Si gjendet?", "Me teste", "Duke e lexuar kodin, ose me mjete si ky punim"],
        ["Shembull", "Shuma del gabim", "Një metodë 300 rreshta që askush nuk e kupton"],
    ])  # fmt: skip
    doc.add_heading("Tetë erërat që gjen mjeti", level=2)
    mlcq = {"GodClass", "DataClass", "FeatureEnvy", "LongMethod"}
    _grid(doc, ["Era", "Emri i thjeshtë", "Çka është", "Pse ka rëndësi", "Në MLCQ?"], [
        [_words(name), entry["name"], entry["what"], entry["why"], "po" if name in mlcq else "jo"]
        for name, entry in vocab["smells"].items()
    ])  # fmt: skip
    doc.add_paragraph(
        "Mjeti gjen tetë erëra, por vetëm katër prej tyre kanë etiketa nga njerëz në bazën MLCQ "
        "(kapitulli 8). Prandaj vetëm këto katër mund të krahasohen me mendimin e "
        "programuesve. «Blob» në MLCQ është e njëjta erë që libri i Lanza-s dhe Marinescu-t e "
        "quan «God Class»."
    )
    doc.add_heading("Katër erërat e punimit, me shembuj", level=2)
    _list(doc, [
        "Blob: një klasë që bën gjithçka. Si një punonjës që bën punën e gjithë zyrës: kur "
        "mungon ai, asgjë nuk funksionon, dhe askush nuk di çka bën saktësisht.",
        "Data Class: një klasë që vetëm mban të dhëna dhe nuk bën asgjë me to. Si një formular "
        "i plotësuar, me të cilin punojnë vetëm të tjerët.",
        f"Long Method: një metodë shumë e gjatë. Mjeti e quan të gjatë mbi {t['long_method_loc']:g} "
        "rreshta kodi. Si një recetë tri faqe pa asnjë nëntitull.",
        "Feature Envy («zili për veçoritë e tjetrit»): një metodë që punon më shumë me të dhënat "
        "e një klase tjetër se me të vetat. Si një kuzhinier që gatuan gjithmonë në kuzhinën e "
        "fqinjit: ndoshta i takon atje.",
    ])  # fmt: skip
    doc.add_paragraph(
        "Ja një shembull i vogël, i shkruar për këtë dokument, që i ka dy nga këto erëra "
        "njëkohësisht:"
    )
    _code(doc, EXAMPLE_INVOICE.splitlines())
    doc.add_paragraph(
        "Klasa Customer vetëm mban gjashtë të dhëna dhe i jep me metoda «get». Metoda label() "
        "e klasës Invoice merr gjashtë të dhëna nga Customer dhe vetëm dy nga klasa e vet "
        "(customer dhe amount). "
        "Kur mjeti e mat këtë kod, gjen:"
    )
    _list(doc, [
        f"{_words(smell.smell_type)}"
        + (f" te {smell.method}" if smell.method else f" te klasa {smell.class_name}")
        + ": " + ", ".join(c.describe() for c in smell.conditions)
        for smell in invoice_smells
    ])  # fmt: skip
    doc.add_paragraph(
        "Kuptimin e çdo shkurtese (WOC, ATFD, LAA…) e mëson te kapitulli 5. Për momentin mjafton "
        "kjo: mjeti nuk thotë vetëm «ka erë», por edhe pse, me numra."
    )
    _idea(doc, "Një erë nuk është gabim: është shenjë që kodi do të jetë i vështirë për "
               "t'u ndryshuar.")  # fmt: skip
    _check(doc, [
        ("A punon programi kur ka një erë kodi?", "Po. Era nuk e prish programin; e bën kodin të "
                                                  "vështirë për t'u ndryshuar."),
        ("Cilat janë katër erërat që mat punimi?", "Blob, Data Class, Long Method dhe Feature Envy."),
        ("Pse vetëm këto katër?", "Sepse vetëm për to ka etiketa nga programues profesionistë në "
                                  "bazën MLCQ."),
        ("Çka do të thotë Feature Envy?", "Një metodë që punon më shumë me të dhënat e një klase "
                                          "tjetër se me të vetat."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 3, "Refaktorimi: ndreqja pa e ndryshuar sjelljen",
             "Në këtë kapitull mëson çka do të thotë të refaktorosh, dhe cilat ndreqje i bën mjeti "
             "vetë.")
    doc.add_paragraph(
        "Refaktorimi (refactoring) është ndryshimi i ndërtimit të kodit pa ndryshuar atë që bën "
        "programi. Pas refaktorimit, për çdo hyrje programi duhet të japë saktësisht të njëjtin "
        "rezultat si më parë. Ndryshon vetëm sa lehtë lexohet dhe ndryshohet."
    )
    _analogy(doc, "Rregullon dollapin: rrobat janë po ato, por tani i gjen më shpejt. Nëse gjatë "
                  "rregullimit humbet një këmishë, nuk ishte rregullim: ishte dëm.")  # fmt: skip
    doc.add_paragraph(
        "Rreziku i refaktorimit është pikërisht ky: një ndryshim që duket i pafajshëm mund ta "
        "ndryshojë sjelljen pa e kuptuar askush. Prandaj motori i këtij punimi është shumë i "
        "kujdesshëm (kapitulli 12)."
    )
    doc.add_heading("Ndreqjet që propozon mjeti", level=2)
    automated = {s["automated"] for s in reference["strategies"] if s["automated"]}
    _grid(doc, ["Ndreqja", "Emri i thjeshtë", "Çka bën", "E zbaton mjeti vetë?"], [
        [_words(name), entry["name"], entry["what"],
         "po" if name in automated else "vetëm e propozon"]
        for name, entry in vocab["refactorings"].items()
    ])  # fmt: skip
    advisory = [s for s in reference["strategies"] if not s["automated"]]
    doc.add_paragraph(
        "Pse disa ndreqje vetëm propozohen: për "
        + _dhe([_words(s["smell"]) for s in advisory])
        + ", ndreqja do të duhej të zhvendoste metoda ose të ndryshonte fusha, pra të gjente "
        "çdo vend në gjithë projektin ku ato përdoren. Mjeti i njeh faktet e sintaksës, jo "
        "kuptimin e çdo emri në gjithë projektin (kapitulli 4), ndaj nuk mund t'i gjejë me "
        "siguri të gjitha këto vende. Më mirë të mos e bëjë, se sa ta bëjë gabim."
    )
    _idea(doc, "Refaktorimi e ndryshon formën e kodit, kurrë sjelljen.")
    _check(doc, [
        ("Çka është rregulli i artë i refaktorimit?",
         "Për çdo hyrje, programi duhet të japë saktësisht të njëjtin rezultat si më parë."),
        ("Cilat tri ndreqje i bën mjeti vetë?",
         "Extract Method, Guard Clauses dhe Introduce Parameter Object."),
        ("Pse Move Method vetëm propozohet?",
         "Sepse duhen gjetur të gjitha vendet ku thirret metoda, dhe mjeti nuk i zgjidh dot "
         "emrat në gjithë projektin."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 4, "Si e lexon mjeti kodin",
             "Në këtë kapitull mëson çka është një parser dhe një pemë sintakse, dhe pse mjeti "
             "nuk punon me tekst të thjeshtë.")
    doc.add_paragraph(
        "Për njeriun, kodi është tekst. Për mjetin, kodi duhet të bëhet strukturë: duhet të dijë "
        "ku nis një klasë, cilat janë metodat e saj, cili if është brenda cilit. Programi që e "
        "kthen tekstin në strukturë quhet parser. Rezultati quhet pemë e sintaksës (AST, "
        "Abstract Syntax Tree)."
    )
    doc.add_paragraph("Për shembull, rreshtat «if (p > 0) { sum += p; }» bëhen këtë pemë:")
    _code(doc, [
        "if_statement",
        "├── condition:  p > 0",
        "│   ├── p",
        "│   ├── >",
        "│   └── 0",
        "└── block",
        "    └── sum += p",
    ])  # fmt: skip
    _analogy(doc, "Analiza gramatikore në shkollë e ndan fjalinë «Ana lexon librin» në "
                  "kryefjalë, kallëzues dhe kundrinë. Parseri bën të njëjtën gjë me kodin.")  # fmt: skip
    doc.add_paragraph(
        "Mjeti përdor parserin tree-sitter. Ai është i shpejtë, punon edhe kur kodi ka gabime "
        "ose nuk kompilohet, dhe nuk kërkon që projekti të ndërtohet. Kjo ka rëndësi, sepse "
        "projektet e studimit vijnë pa skedarët e ndërtimit (kapitulli 8)."
    )
    doc.add_paragraph(
        "Çmimi është ky: tree-sitter sheh vetëm sintaksën. Ai e di që në kod shkruhet "
        "«customer.getName()», por nuk e di me siguri se cila klasë është customer në gjithë "
        "projektin. Mjeti mbledh vetëm këto fakte sintakse: emra të thjeshtë, «this.x» dhe "
        "«marrësi.anëtari». Me qëllim nuk përpiqet të zgjidhë çdo emër, sepse një zgjidhje e "
        "hamendësuar do të ishte më e rrezikshme se një refuzim."
    )
    doc.add_paragraph(
        "Nga pema, mjeti ndërton një model të thjeshtë të projektit: klasat, fushat, metodat, "
        "parametrat dhe çka përdor secila metodë. Çdo klasë lexohet vetëm një herë, dhe të "
        "gjitha matjet llogariten nga ky model."
    )
    _idea(doc, "Mjeti e kthen kodin në një pemë dhe punon mbi pemën, jo mbi tekstin.")
    _check(doc, [
        ("Çka bën një parser?", "E kthen tekstin e kodit në strukturë, në një pemë sintakse."),
        ("Pse tree-sitter?", "Sepse është i shpejtë, punon edhe me kod që nuk kompilohet, dhe "
                             "nuk kërkon ndërtimin e projektit."),
        ("Çka nuk di tree-sitter?", "Nuk i zgjidh emrat në gjithë projektin: nuk di me siguri "
                                    "se cilës klasë i përket një emër."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 5, "Metrikat: kodi i kthyer në numra",
             "Në këtë kapitull mëson çka është një metrikë, dhe llogarit me dorë disa prej tyre "
             "mbi një shembull të vogël.")
    doc.add_paragraph(
        "Një metrikë është një numër që përshkruan kodin. Mjeku të mat presionin, temperaturën "
        "dhe peshën: asnjë numër i vetëm nuk jep diagnozë, por bashkë po. Njësoj, mjeti mat "
        f"{len(reference['metrics']['class'])} metrika për çdo klasë dhe "
        f"{len(reference['metrics']['method'])} për çdo metodë, gjithsej {f['metrics']}."
    )
    for level, title in (("class", "Metrikat e klasës"), ("method", "Metrikat e metodës")):
        doc.add_heading(title, level=2)
        _grid(doc, ["Shkurtesa", "Emri", "Çka mat"], [
            [name, vocab["metrics"][name]["name"], vocab["metrics"][name]["what"]]
            for name in reference["metrics"][level]
        ])  # fmt: skip
    doc.add_heading("Shembull i llogaritur me dorë", level=2)
    doc.add_paragraph("Merr këtë klasë të vogël:")
    _code(doc, EXAMPLE_SHOP.splitlines())
    doc.add_paragraph("Për metodën total(), llogaritja me dorë jep:")
    _grid(doc, ["Metrika", "Si numërohet", "Vlera"], [
        ["MLOC (rreshtat)", "Numërohen rreshtat që kanë kod, bashkë me rreshtin e emrit të "
         "metodës. Rreshtat bosh, komentet dhe rreshtat që kanë vetëm «}» nuk numërohen.",
         f"{total['MLOC']:g}"],
        ["CC (rrugët)", "Nis nga 1 dhe shton 1 për çdo vendim: for (+1), if (p > 0) (+1), "
         "if (member && …) (+1), dhe && (+1).", f"{total['CC']:g}"],
        ["MAXNESTING (thellësia)", "for është niveli 1, if brenda tij niveli 2. If-i i dytë "
         "është vetëm niveli 1.", f"{total['MAXNESTING']:g}"],
        ["NP (parametrat)", "prices dhe member.", f"{total['NP']:g}"],
    ])  # fmt: skip
    doc.add_paragraph(
        "Këto janë saktësisht vlerat që jep mjeti kur e mat këtë kod. WMC e klasës është shuma "
        f"e CC-së së të gjitha metodave: këtu ka vetëm një metodë, ndaj WMC = "
        f"{shop_classes['Shop']['WMC']:g}."
    )
    doc.add_heading("Kohezioni (TCC): a lidhen metodat mes tyre?", level=2)
    _code(doc, EXAMPLE_ACCOUNT.splitlines())
    doc.add_paragraph(
        "TCC shikon çiftet e metodave publike dhe pyet: a përdorin të dyja të paktën një fushë "
        "të përbashkët? Këtu ka tri metoda, pra tri çifte. deposit dhe report përdorin të dyja "
        "balance: të lidhura. deposit dhe notify: jo. report dhe notify: jo. Pra 1 nga 3 çifte, "
        f"TCC = {account_classes['Account']['TCC']:.2f}. Vlera e ulët do të thotë që klasa bën "
        "gjëra që nuk lidhen me njëra-tjetrën."
    )
    doc.add_heading("Të dhënat e huaja (ATFD, LAA, FDP)", level=2)
    doc.add_paragraph(
        "Kthehu te metoda label() e kapitullit 2. Ajo lexon gjashtë të dhëna të Customer përmes "
        "metodave get, dhe dy fusha të klasës së vet (customer dhe amount). Mjeti mat:"
    )
    _list(doc, [
        f"ATFD = {label['ATFD']:g}: sa të dhëna të huaja lexon.",
        f"LAA = {label['LAA']:.2f}: pjesa e të dhënave që janë të klasës së vet (2 nga 8).",
        f"FDP = {label['FDP']:g}: nga sa klasa të huaja vijnë të dhënat (vetëm Customer).",
    ])  # fmt: skip
    doc.add_paragraph(
        "Burimet e metrikave: Chidamber & Kemerer (1994) për WMC, CBO, DIT, NOC, RFC dhe LCOM; "
        "McCabe (1976) për CC; Bieman & Kang (1995) për TCC; Lanza & Marinescu (2006) për ATFD, "
        "LAA, FDP, WOC dhe metrikat e tjera të rregullave të tyre."
    )
    _idea(doc, "Metrikat janë numra të thjeshtë për kodin; bashkë tregojnë ku ka probleme.")
    _check(doc, [
        ("Sa është CC për një metodë pa asnjë if, for ose while?", "1."),
        ("Pse numërohet && te CC?", "Sepse krijon një rrugë tjetër: kushti mund të ndalet pas "
                                    "pjesës së parë."),
        ("Çka do të thotë TCC = 0?", "Asnjë çift metodash nuk përdor fusha të përbashkëta: "
                                     "klasa bën gjëra krejt të palidhura."),
        ("Çka mat LAA?", "Sa nga të dhënat që përdor një metodë janë të klasës së vet."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 6, "Qasja A: rregullat e botuara",
             "Në këtë kapitull mëson si i kthejnë rregullat metrikat në një vendim «ka erë / nuk "
             "ka», dhe pse ky vendim humbet shumë raste.")
    doc.add_paragraph(
        "Një rregull gjetjeje (detection strategy) është një kombinim kushtesh mbi metrika. Kur "
        "të gjitha kushtet plotësohen, kodi shënohet. Shumica e rregullave vijnë nga libri i "
        "Lanza-s dhe Marinescu-t «Object-Oriented Metrics in Practice» (2006)."
    )
    _grid(doc, ["Era", "Rregulli", "Burimi"], [
        [_words(s["smell"]), s["formula"], s["title"]] for s in reference["strategies"]
    ])  # fmt: skip
    doc.add_paragraph(
        "Fjalët e mëdha në rregulla janë emra për vlerat më poshtë; te rregullat e klasës, HIGH "
        "dhe VERY_HIGH janë HIGH_WMC dhe VERY_HIGH_WMC."
    )
    _grid(doc, ["Emri", "Vlera"], [
        [name, f"{value:.2f}" if isinstance(value, float) and value != int(value) else f"{value:g}"]
        for name, value in q.items()
    ])  # fmt: skip
    doc.add_paragraph(
        f"Lexoje me fjalë rregullin e God Class: WMC ≥ {q['VERY_HIGH_WMC']} DHE TCC < 1/3 DHE "
        f"ATFD > {q['FEW']}. Pra: klasa bën shumë punë, pjesët e saj nuk lidhen me njëra-tjetrën, "
        "dhe merr shumë të dhëna nga klasat e tjera. Të tria duhet të jenë të vërteta."
    )
    doc.add_heading("Mjeti thotë edhe pse", level=2)
    doc.add_paragraph(
        "Një rregull nuk kthen vetëm «po» ose «jo». Kthen kushtet që kontrolloi, me vlerën e "
        "matur, si te shembulli i kapitullit 2: «ATFD = 6 (> 5)». Kështu ndërfaqja i shpjegon "
        "programuesit pse, dhe punimi mund të tregojë cili kusht e shkaktoi gjetjen."
    )
    doc.add_heading("Sa e rëndë është era", level=2)
    doc.add_paragraph(
        "Për çdo kusht, mjeti llogarit sa herë e kalon vlera kufirin, si raport. Merr mesataren "
        f"e raporteve, secili me tavan {t['excess_cap']:g}, që një vlerë e vetme ekstreme të mos "
        f"e bëjë gjithçka të rëndë. Nën {t['severity_major']:g} era është minor (e lehtë), nën "
        f"{t['severity_critical']:g} major (e mesme), dhe ndryshe critical (e rëndë). Shembull: "
        f"një metodë me 60 rreshta, me kufi {t['long_method_loc']:g}, ka raport 60 / "
        f"{t['long_method_loc']:g} = {60 / t['long_method_loc']:.1f}, pra major. Kjo shkallë është "
        "e njëjtë me atë të MLCQ-së, që rezultatet të krahasohen drejtpërdrejt."
    )
    doc.add_heading("Vlerat pa burim të botuar", level=2)
    doc.add_paragraph(
        f"Katër kufij nuk kanë burim të botuar: {t['long_method_loc']:g} rreshta për Long Method, "
        f"CC {t['brain_method_cc']:g} dhe thellësia {t['brain_method_nesting']:g} për Brain "
        f"Method, dhe thellësia {t['deep_nesting']:g} për Deep Nesting. Edhe tri vlerat e "
        "rëndësisë janë të autorit. Punimi e thotë hapur. Kufiri i Long Method u provua me vlera "
        "të ndryshme dhe u kalibrua kundrejt MLCQ-së; të tjerët nuk mund të provohen, sepse MLCQ "
        "nuk ka etiketa për Brain Method dhe Deep Nesting."
    )
    doc.add_heading("Problemi i rregullave", level=2)
    doc.add_paragraph(
        "Numrat e rregullave janë fiksë dhe të caktuar duke parë kod tjetër. Rezultati, që do ta "
        f"shohësh te kapitulli 10: kur alarmojnë zakonisht kanë të drejtë (precision "
        f"{min(precision):.0%}–{max(precision):.0%}), por humbasin shumicën e rasteve: nga 100 "
        f"erëra të vërteta gjejnë vetëm {lowest:.0f} deri {highest:.0f}."
    )
    _analogy(doc, "Një roje që lë të hyjë vetëm kë ka mbi dy metra gjatësi do të ndalë vërtet "
                  "disa njerëz të rrezikshëm, por do t'i lërë të hyjnë shumicën e tyre.")  # fmt: skip
    _idea(doc, "Rregullat janë të sakta kur alarmojnë, por të rrepta: i humbasin shumicën e "
               "erërave.")  # fmt: skip
    _check(doc, [
        ("Kur fiton God Class?", f"Kur WMC ≥ {q['VERY_HIGH_WMC']}, TCC < 1/3 dhe ATFD > {q['FEW']}, "
                                 "të tria njëkohësisht."),
        ("Pse rregulli kthen kushtet dhe jo vetëm po/jo?",
         "Që programuesi të shohë pse, dhe punimi të tregojë cili kusht e shkaktoi gjetjen."),
        ("Një metodë ka 90 rreshta. Sa e rëndë është, sipas Long Method?",
         f"Raporti është 90 / {t['long_method_loc']:g} = {90 / t['long_method_loc']:.1f}, mbi "
         f"{t['severity_critical']:g}, pra critical."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 7, "Qasja B: mësimi makinerik nga e para",
             "Në këtë kapitull mëson çka do të thotë që një kompjuter «mëson», dhe si punojnë "
             "katër modelet e punimit.")
    doc.add_paragraph(
        "Si e dallon një shalqi të pjekur? Mund të kesh një rregull: «nëse tingëllon bosh, është "
        "i pjekur». Ose mund të provosh një mijë shalqinj, t'i shënosh cilët ishin të pjekur, dhe "
        "të mësosh vetë kufirin nga përvoja. Kjo e dyta është ideja e mësimit makinerik "
        "(machine learning)."
    )
    _grid(doc, ["Fjala", "Në punim"], [
        ["Shembull (sample)", "Një klasë ose metodë e gjykuar nga një njeri"],
        ["Tipare (features)", f"{f['metrics']} metrikat e saj"],
        ["Etiketë (label)", "Përgjigjja e njeriut: ka erë apo jo"],
        ["Mësim (training)", "Modeli sheh shumë shembuj me etiketë dhe gjen kufirin"],
        ["Provë (testing)", "Modeli gjykon shembuj që nuk i ka parë, dhe krahasohet me njerëzit"],
    ])  # fmt: skip
    doc.add_paragraph(
        "Rreziku më i madh quhet mbipërshtatje (overfitting): modeli i mëson përmendësh shembujt "
        "e mësimit në vend që të mësojë rregullin. Del i shkëlqyer mbi ata shembuj dhe i dobët "
        "mbi të rinjtë. Prandaj modeli provohet gjithmonë mbi shembuj që nuk i ka parë "
        "(kapitulli 9)."
    )
    _analogy(doc, "Një student që mëson përmendësh përgjigjet e provimit të vjetër merr notë të "
                  "lartë vetëm nëse provimi përsëritet. Me pyetje të reja, ngec.")  # fmt: skip
    doc.add_heading("Katër modelet", level=2)
    explain = {
        "majority": "Thotë gjithmonë përgjigjen më të shpeshtë («i pastër»). Nuk mëson asgjë: "
        "është pika e krahasimit që çdo model tjetër duhet ta kalojë.",
        "logistic": "I jep çdo metrike një peshë, i mbledh, dhe e kthen shumën në një gjasë nga "
        "0 deri në 1. Mbi gjysmën thotë «ka erë». I thjeshtë dhe i shpjegueshëm.",
        "random_forest": "Shumë pemë vendimi, secila si loja «20 pyetje» («a ka MLOC mbi 30? "
        "a ka CC mbi 4?…»), dhe secila e mëson pak ndryshe. Në fund votojnë. Punimi përdor 300 "
        "pemë.",
        "gradient_boosting": "Pemë njëra pas tjetrës: çdo pemë e re përpiqet të ndreqë gabimet "
        "e të mëparshmeve.",
    }
    _grid(doc, ["Modeli", "Si punon", "Fitoi te"], [
        [MODEL_EN[m], text, _dhe([SMELL_SQ[s] for s in smells if winners[s] == m]) or "–"]
        for m, text in explain.items()
    ])  # fmt: skip
    doc.add_paragraph(
        "Erërat janë të rralla: shumica e kodit është i pastër. Nëse modeli nuk e di këtë, mëson "
        "të thotë gjithmonë «i pastër». Prandaj çdo model që e lejon i jep më shumë peshë "
        "shembujve të rrallë (class weight «balanced»). Punimi nuk i shumëfishoi shembujt e "
        "rrallë, sepse kopjet mund të kalonin nga grupi i mësimit te grupi i provës dhe ta "
        "prishnin ndarjen sipas projektit (kapitulli 9)."
    )
    doc.add_paragraph(
        "Pse jo deep learning: të dhënat janë të pakta dhe janë një tabelë numrash, modelet duhet "
        "të shpjegohen, dhe gjithçka duhej të punonte në një laptop pa kartë grafike dhe pa "
        "kosto. Modelet e thjeshta mjaftuan për t'i mundur rregullat."
    )
    doc.add_paragraph("Metrikat që ndikojnë më shumë te secili model:")
    _list(doc, [
        f"{SMELL_SQ[s]}: " + ", ".join(_metric_name(x) for x in ml["per_smell"][s]["top_features"][:4])
        for s in smells
    ])  # fmt: skip
    doc.add_paragraph(
        f"Numrat e rastit që përdorin modelet nisin gjithmonë nga e njëjta farë ({ml['seed']}), "
        "ndaj kushdo që i rindërton merr saktësisht të njëjtat rezultate."
    )
    _idea(doc, "Modeli mëson kufirin nga shembujt e gjykuar nga njerëzit, në vend që ta marrë "
               "të gatshëm nga një libër.")  # fmt: skip
    _check(doc, [
        ("Çka është mbipërshtatja?", "Kur modeli i mëson përmendësh shembujt në vend të "
                                     "rregullit, dhe dështon me shembuj të rinj."),
        ("Pse ka një model «majority»?", "Si pikë krahasimi: një model i vërtetë duhet të bëjë më "
                                         "mirë se ai që thotë gjithmonë «i pastër»."),
        ("Si punon random forest?", "Shumë pemë vendimi mësojnë secila pak ndryshe dhe në fund "
                                    "votojnë."),
        ("Pse nuk u shumëfishuan shembujt e rrallë?", "Sepse kopjet mund ta kalonin ndarjen mes "
                                                      "mësimit dhe provës."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 8, "Të dhënat: MLCQ dhe korpusi",
             "Në këtë kapitull mëson nga vijnë përgjigjet e njerëzve, dhe çka u desh për t'i "
             "lidhur me kodin.")
    doc.add_paragraph(
        "Për ta provuar një detektor, të duhet një «çelës përgjigjesh»: kod për të cilin njerëz "
        "me përvojë kanë thënë «ka erë» ose «nuk ka». Punimi përdor MLCQ (Madeyski & Lewowski, "
        "2020), një bazë publike të dhënash."
    )
    _list(doc, [
        "26 programues profesionistë shqyrtuan kod të vërtetë Java me burim të hapur dhe dhanë "
        "rreth 15 mijë vlerësime.",
        "Çdo vlerësim është në shkallën none, minor, major, critical: pa erë, e lehtë, e mesme, "
        "e rëndë.",
        "Shpesh e njëjta copë kodi ka disa vlerësues. Etiketa e punimit merret nga mesatarja e "
        "vlerësimeve të tyre.",
        f"MLCQ tregon vetëm ku ndodhet kodi. Punimi shkarkoi projektet dhe e gjeti çdo mostër në "
        f"kod: nga {dataset['samples_considered']:,} mostra u gjetën {f['samples']:,}. Të tjerat "
        "humbën sepse skedari mungonte, kishte gabime sintakse, ose rreshtat nuk përputheshin.",
        f"MLCQ përmend {dataset['repositories']} projekte; disa nuk ishin më të arritshme, ndaj "
        f"mostrat vijnë nga {f['repositories']}.",
        "Projektet e shkarkuara quhen korpus. Korpusi mban vetëm skedarët .java, pa skedarët e "
        "ndërtimit dhe pa libraritë. Kjo ka rëndësi te kontrolli me javac (kapitulli 12).",
    ])  # fmt: skip
    _idea(doc, "Rezultatet krahasohen me gjykimin e programuesve profesionistë, jo me mendimin e "
               "autorit.")  # fmt: skip
    _check(doc, [
        ("Kush i etiketoi të dhënat?", "26 programues profesionistë, në bazën publike MLCQ."),
        ("Pse u humbën disa mostra?", "Sepse skedari mungonte, kishte gabime sintakse, ose "
                                      "rreshtat nuk përputheshin me kodin."),
        ("Si bëhet etiketa kur ka disa vlerësues?", "Merret mesatarja e vlerësimeve të tyre."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 9, "Si matet nëse një detektor është i mirë",
             "Në këtë kapitull mëson matricën e konfuzionit, precision, recall dhe MCC, me "
             "llogaritje të plota, dhe pse të dhënat ndahen sipas projektit.")
    doc.add_heading("Matrica e konfuzionit", level=2)
    blob_total = blob["tp"] + blob["fp"] + blob["tn"] + blob["fn"]
    doc.add_paragraph(
        f"Merr rregullat për Blob mbi {blob_total:,} mostra. Çdo mostër bie në njërën nga katër "
        "kutitë:"
    )
    _grid(doc, ["", "Njerëzit: erë", "Njerëzit: e pastër"], [
        ["Mjeti: erë", f"{blob['tp']} (TP, saktë)", f"{blob['fp']} (FP, alarm i rremë)"],
        ["Mjeti: e pastër", f"{blob['fn']} (FN, e humbur)", f"{blob['tn']:,} (TN, saktë)"],
    ])  # fmt: skip
    f1 = 2 * blob["precision"] * blob["recall"] / (blob["precision"] + blob["recall"])
    accuracy = (blob["tp"] + blob["tn"]) / blob_total
    _list(doc, [
        f"Precision = TP / (TP + FP) = {blob['tp']} / ({blob['tp']} + {blob['fp']}) = "
        f"{blob['precision']:.1%}. Kur thotë «erë», ka të drejtë kaq shpesh.",
        f"Recall = TP / (TP + FN) = {blob['tp']} / ({blob['tp']} + {blob['fn']}) = "
        f"{blob['recall']:.1%}. Nga erërat e vërteta, gjen kaq.",
        f"F1 = 2 · P · R / (P + R) = {f1:.3f}. Një notë që i bashkon të dyja.",
        f"Saktësia (accuracy) = (TP + TN) / gjithsej = {accuracy:.1%}. Duket e mirë, por mashtron.",
    ])  # fmt: skip
    doc.add_heading("Kurthi i saktësisë", level=2)
    doc.add_paragraph(
        f"Te Feature Envy vetëm {envy['support_positive']} nga {envy_total:,} mostra kanë erë. "
        f"Një «detektor» që thotë gjithmonë «i pastër» del {majority:.1%} i saktë, edhe pse nuk "
        "ka gjetur asnjë erë. Prandaj saktësia nuk mjafton kur erërat janë të rralla."
    )
    doc.add_heading("MCC: nota që nuk mashtrohet", level=2)
    numerator = blob["tp"] * blob["tn"] - blob["fp"] * blob["fn"]
    product = ((blob["tp"] + blob["fp"]) * (blob["tp"] + blob["fn"])
               * (blob["tn"] + blob["fp"]) * (blob["tn"] + blob["fn"]))  # fmt: skip
    mcc = numerator / product**0.5
    doc.add_paragraph(
        "MCC (koeficienti i korrelacionit Matthews) përdor të katër kutitë e matricës. Shkon nga "
        "−1 në +1: +1 është e përsosur, 0 është sa hamendja, −1 është gjithmonë e kundërta."
    )
    _code(doc, [
        "MCC = (TP·TN − FP·FN) / √((TP+FP)·(TP+FN)·(TN+FP)·(TN+FN))",
        "",
        f"    = ({blob['tp']}·{blob['tn']} − {blob['fp']}·{blob['fn']}) / "
        f"√({blob['tp'] + blob['fp']}·{blob['tp'] + blob['fn']}·{blob['tn'] + blob['fp']}·"
        f"{blob['tn'] + blob['fn']})",
        f"    = {numerator:,} / {product**0.5:,.0f}",
        f"    = {mcc:.3f}",
    ])  # fmt: skip
    doc.add_paragraph(
        f"Pra rregullat për Blob kanë MCC {mcc:.2f}, edhe pse saktësia e tyre është "
        f"{accuracy:.0%}. Detektori që thotë gjithmonë «i pastër» merr MCC 0."
    )
    doc.add_heading("Ndarja sipas projektit", level=2)
    doc.add_paragraph(
        "Kodi i një projekti ka të njëjtin stil, të njëjtat emra dhe shpesh të njëjtët autorë. "
        "Nëse modeli mëson nga një pjesë e projektit dhe provohet te një pjesë tjetër e të "
        "njëjtit projekt, ai mund ta njohë projektin, jo erën. Ky gabim quhet rrjedhje e të "
        "dhënave (data leakage), dhe është gjëja e parë që kontrollon një recensent."
    )
    doc.add_paragraph(
        f"Prandaj projektet ndahen në {ml['folds']} grupe. Modeli mëson nga katër grupe dhe "
        "provohet te i pesti; pastaj rrotullohet, derisa çdo grup të jetë provuar një herë. Çdo "
        "mostër parashikohet nga një model që nuk e ka parë kurrë projektin e saj."
    )
    _analogy(doc, "Nuk e provon studentin me pyetjet që i ke dhënë për ushtrim. Merr pyetje nga "
                  "një libër që nuk e ka hapur.")  # fmt: skip
    doc.add_heading("Bootstrap: a është ndryshimi i vërtetë apo fat?", level=2)
    resamples = f["resamples"]
    doc.add_paragraph(
        f"Modeli del më i mirë se rregullat. Por a do të dilte i tillë edhe me projekte të "
        f"tjera? Për ta parë, projektet hidhen në një thes dhe nxirren rastësisht, me kthim, "
        f"aq sa janë, dhe ndryshimi model − rregulla rillogaritet. Kjo përsëritet {resamples:,} "
        "herë. 95% e mesme e këtyre vlerave quhet interval besimi. Nëse i gjithë intervali është "
        "mbi zero, përmirësimi nuk është fat."
    )
    _idea(doc, "MCC e mat saktë kur erërat janë të rralla; ndarja sipas projektit dhe bootstrap "
               "e bëjnë provën të drejtë.")  # fmt: skip
    _check(doc, [
        ("Pse saktësia mashtron te Feature Envy?",
         f"Sepse kodi është pothuajse gjithë i pastër: «gjithmonë i pastër» del {majority:.1%} i "
         "saktë pa gjetur asgjë."),
        ("Çka do të thotë MCC 0?", "Sa hamendja: detektori nuk ka mësuar asgjë."),
        ("Çka është rrjedhja e të dhënave?", "Kur modeli provohet mbi diçka që e ka parë gjatë "
                                             "mësimit, p.sh. mbi të njëjtin projekt."),
        ("Kur themi që përmirësimi nuk është fat?", "Kur intervali 95% i ndryshimit është i "
                                                    "gjithi mbi zero."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 10, "Rezultatet e gjetjes së erërave",
             "Në këtë kapitull sheh çka gjetën rregullat dhe modelet, erë për erë, dhe si "
             "krahasohen me mjete dhe punime të tjera.")
    rows = []
    for s in smells:
        rule, model = f["strategy"][s], f["best"][s]
        diff = intervals[s]["intervals"]["difference_model_minus_rules"]
        rows.append([SMELL_SQ[s], f"{rule['precision']:.2f} / {rule['recall']:.2f} / {rule['mcc']:.2f}",
                     MODEL_EN[winners[s]],
                     f"{model['precision']:.2f} / {model['recall']:.2f} / {model['mcc']:.2f}",
                     f"{diff['low']:+.2f} … {diff['high']:+.2f}",
                     "po" if diff["low"] > 0 else "jo, përfshin zeron"])  # fmt: skip
    _grid(doc, ["Era", "Rregullat P / R / MCC", "Modeli", "Modeli P / R / MCC", "B − A, 95%",
                "I sigurt?"], rows)  # fmt: skip
    doc.add_heading("Si ta lexosh tabelën", level=2)
    _list(doc, [
        f"Nga çdo 100 erëra të vërteta, rregullat gjejnë {lowest:.0f} deri {highest:.0f}; "
        f"modelet {min(found):.0f} deri {max(found):.0f}.",
        "Modelet gjejnë shumë më tepër, por zakonisht me më shumë alarme të rreme: te Blob, "
        f"rregullat kanë precision {blob['precision']:.0%}, modeli "
        f"{f['best']['blob']['precision']:.0%}.",
        f"Përmirësimi është i sigurt te {_dhe(f['sure'])}. Te {_dhe(f['unsure'])} intervali "
        "përfshin zeron, ndaj mund të jetë fat.",
        f"Edhe algoritmi më i dobët i mund rregullat te {f['weakest']} nga {len(smells)} "
        "erërat: përfundimi nuk varet nga zgjedhja e modelit më të mirë.",
    ])  # fmt: skip
    doc.add_heading(f"Krahasimi me PMD {pmd['pmd_version']}", level=2)
    doc.add_paragraph(
        "PMD është një mjet falas dhe shumë i përdorur për analizën e kodit. U ekzekutua me "
        "pragjet e veta mbi të njëjtat mostra. Ndryshimi është yni minus i PMD-së."
    )
    _grid(doc, ["Era", "PMD MCC", "MCC ynë", "Ndryshimi, 95%"], [
        [PMD_NAMES.get(key, key), f"{e['pmd']['mcc']:.2f}",
         "–" if e.get("ours") is None else f"{e['ours']['mcc']:.2f}",
         "–" if e.get("difference") is None
         else f"{e['difference']['low']:+.2f} … {e['difference']['high']:+.2f}"]
        for key, e in pmd["by_smell"].items()
    ])  # fmt: skip
    closer = [PMD_NAMES.get(key, key) for key, e in pmd["by_smell"].items()
              if e.get("ours") and e["pmd"]["mcc"] > e["ours"]["mcc"]]  # fmt: skip
    doc.add_paragraph(
        ("Me intervalin e besimit, PMD nuk del kurrë përpara. " if not f["pmd_ahead"] else "")
        + (f"Te {_dhe(closer)} PMD ka numër pak më të lartë, por intervali i ndryshimit e "
           "përfshin zeron, ndaj nuk mund të thuhet që është përpara. " if closer else "")
        + f"Te {_dhe(f['pmd_behind'])} rregullat tona janë qartë më të mira. PMD nuk ka rregull "
        "për Feature Envy; LawOfDemeter është më i afërti. Me SonarQube nuk u krahasua."
    )  # fmt: skip
    doc.add_heading("Krahasimi me literaturën", level=2)
    _grid(doc, ["Era", "Rregullat (A)", "Modeli (B)", "Madeyski & Lewowski 2023"], [
        [SMELL_SQ[s], f"{f['strategy'][s]['mcc']:.2f}", f"{f['best'][s]['mcc']:.2f}",
         MADEYSKI_2023_DS1[s]] for s in smells
    ])  # fmt: skip
    doc.add_paragraph(
        "Ata përdorën ndarje rastësore dhe modele të akorduara, që zakonisht japin numra më të "
        "lartë. Ky punim i ndau sipas projektit dhe nuk i akordoi. Rezultatet janë në të njëjtin "
        "rang."
    )
    _idea(doc, "Kufijtë e mësuar i mundin pragjet fikse, qartë te tri nga katër erërat.")
    _check(doc, [
        ("Kush gjen më shumë erëra, rregullat apo modelet?",
         f"Modelet: {min(found):.0f}–{max(found):.0f} nga 100, kundrejt {lowest:.0f}–"
         f"{highest:.0f} të rregullave."),
        ("Cili është çmimi i modeleve?", "Zakonisht më shumë alarme të rreme."),
        ("Te cila erë përmirësimi mund të jetë fat?", f"Te {_dhe(f['unsure'])}."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 11, "Kur as njerëzit nuk pajtohen",
             "Në këtë kapitull mëson pse disa erëra mbeten të vështira për çdo detektor.")
    doc.add_paragraph(
        "Te mostrat me më shumë se një vlerësim, punimi krahasoi çdo çift vlerësimesh të së "
        "njëjtës mostër: a thonë të dy «ka erë» apo jo? Pajtimi u mat me të njëjtën MCC."
    )
    _grid(doc, ["Era", "MCC mes dy rishikuesve", "Kuptimi"], [
        [SMELL_SQ[s], f"{f['reviewers'][s]:.2f}",
         "sa hamendja" if s in f["chance"] else "pajtim i pjesshëm"]
        for s in smells
    ])  # fmt: skip
    doc.add_paragraph(
        "Nëse dy mësues nuk pajtohen për çelësin e përgjigjeve, nxënësi nuk mund ta mësojë "
        "përgjigjen e saktë. Kufiri nuk është vetëm algoritmi: është vetë përkufizimi i erës. "
        f"Prandaj te {_dhe(f['unsure'])} modeli nuk i mund qartë rregullat."
    )
    _idea(doc, "Një erë që ekspertët nuk e njohin njësoj nuk mund të mësohet mirë nga asnjë mjet.")
    _check(doc, [
        ("Si u mat pajtimi i njerëzve?", "Me MCC mbi çiftet e vlerësimeve të së njëjtës mostër."),
        ("Çka tregon −0.07 te Feature Envy?", "Që dy ekspertë pajtohen jo më shumë se me hamendje."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 12, "Motori i refaktorimit: si e ndreq kodin pa e prishur",
             "Në këtë kapitull mëson hapat e motorit, tri ndreqjet, pse refuzon, dhe si kontrollohet "
             "çdo ndryshim.")
    doc.add_paragraph(
        "Motori është pjesa e vetme e mjetit që mund ta prishë punën e dikujt, sepse e rishkruan "
        "kodin. Prandaj ndjek rregullin «më mirë refuzo, se sa prish». Ai nuk përdor model "
        "gjuhësor (LLM): ndreqjet janë transformime të caktuara të pemës, gjithmonë me të njëjtin "
        "rezultat për të njëjtin kod. Një LLM mund të shkruajë kod që duket i saktë, por nuk e "
        "vërteton dot sigurinë, dhe mund të japë përgjigje tjetër çdo herë."
    )
    doc.add_heading("Pesë hapat", level=2)
    _list(doc, [
        "Gjej vendin: ku ka një erë që një ndreqje mund ta heqë.",
        "Kontrollo kushtet paraprake: a mund të vërtetohet nga pema që ndryshimi është i sigurt?",
        "Llogarit ndryshimin nga pema: ku nis dhe ku mbaron çdo pjesë e kodit e thotë pema, jo "
        "një kërkim «gjej dhe zëvendëso» në tekst.",
        "Kontrollo me javac, kompajlerin e Java-s.",
        "Shkruaj një skedar ndryshimesh (patch). Kodi origjinal nuk preket pa një zgjedhje të qartë.",
    ])  # fmt: skip
    doc.add_heading("Tri ndreqjet", level=2)
    doc.add_paragraph("Guard Clauses, nga shembulli i vetë motorit:")
    _code(doc, [
        "void post(int amount) {              void post(int amount) {",
        "    if (amount > 0) {                    if (!(amount > 0)) {",
        "        total += amount;      →              return;",
        "        log(amount);                     }",
        "    }                                    total += amount;",
        "}                                        log(amount);",
        "                                     }",
    ])  # fmt: skip
    doc.add_paragraph(
        "Pranohet vetëm kjo formë: metodë void, i gjithë trupi një if, pa else. Kushti "
        "mbështillet me !( … ) dhe nuk kthehet së prapthi me dorë, ndaj mbetet i saktë për çdo "
        "kusht, edhe për numrat me presje (NaN), ku a > b dhe a <= b mund të jenë të dy të rremë."
    )
    doc.add_paragraph(
        "Extract Method (shembull ilustrues, jo dalja e saktë e motorit): një pjesë e metodës "
        "kalon në një metodë të re, dhe hyrjet e saj bëhen parametra. Motori e quan metodën e re "
        "«extracted»; emrin e mirë e zgjedh programuesi."
    )
    _code(doc, [
        "void report(List<Order> orders) {        void report(List<Order> orders) {",
        "    double total = 0;                        double total = extracted(orders);",
        "    for (Order o : orders) {        →        System.out.println(total);",
        "        total += o.price();                }",
        "    }                                      // + metoda e re extracted(orders)",
        "    System.out.println(total);",
        "}",
    ])  # fmt: skip
    doc.add_paragraph(
        "Introduce Parameter Object (ilustrues): metoda(a, b, c, d, e, f) bëhet metoda(params), "
        "me një klasë të re që i mban të gjashtë."
    )
    doc.add_heading("Pse refuzon", level=2)
    doc.add_paragraph(
        "Mendo për një kirurg: nëse nuk je i sigurt, mos opero. Nëse një kusht nuk vërtetohet "
        f"dot, motori refuzon dhe thotë pse. Gjithsej u refuzuan {refused:,} vende "
        f"({_pct(f['refused'])}):"
    )
    _grid(doc, ["Arsyeja", "Vende", "Pjesa", "Çka do të thotë"], [
        [reason, f"{count:,}", _pct(count / detected), REASON_SQ.get(reason, "")]
        for reason, count in sorted(refactoring["refused_by_reason"].items(), key=lambda p: -p[1])
    ])  # fmt: skip
    doc.add_paragraph(
        "Pse një return e ndalon Extract Method: nëse pjesa që nxirret ka brenda një return, në "
        "metodën e re ai return do të dilte vetëm nga metoda e re, jo nga metoda origjinale. "
        "Programi do të vazhdonte aty ku më parë ndalej. Kjo do të ishte ndryshim sjelljeje."
    )
    doc.add_heading("Kontrolli me javac", level=2)
    doc.add_paragraph(
        "Korpusi nuk ka skedarë ndërtimi, ndaj shumica e skedarëve nuk kompilohen të vetëm as "
        "para ndryshimit: u mungojnë klasat fqinje. Prandaj kontrolli është: kompilo para dhe "
        f"pas, dhe shiko nëse ndryshimi shtoi një lloj të ri gabimi. Nga {applied:,} ndryshimet:"
    )
    _grid(doc, ["Rezultati", "Ndryshime", "Kuptimi"], [
        [verdict, f"{count:,}", VERDICT_SQ.get(verdict, "")]
        for verdict, count in sorted(refactoring["verdicts"].items(), key=lambda p: -p[1])
    ])  # fmt: skip
    inside = verify["compiled_in_project"]
    doc.add_paragraph(
        f"Kur {verify['rewrites']} ndryshime u kompiluan brenda projektit të tyre, ato që "
        f"kompilohen plotësisht u rritën nga {verify['compiled_alone'].get('compiles', 0)} në "
        f"{inside.get('compiles', 0)}, dhe vetëm {inside.get('new_errors', 0)} shtuan gabim të ri. "
        "Pra shumica e rasteve «nuk kompilohet» vjen nga skedarët e vetmuar, jo nga ndryshimi."
    )
    _idea(doc, "Motori ndryshon vetëm atë që vërteton si të sigurt, dhe çdo ndryshim e kontrollon "
               "me kompajlerin.")  # fmt: skip
    _check(doc, [
        ("Pse motori nuk përdor LLM?", "Sepse një LLM nuk e vërteton dot sigurinë dhe nuk jep "
                                       "gjithmonë të njëjtin rezultat."),
        ("Pse refuzimi është rezultat i mirë?", "Sepse është më mirë të mos ndryshosh kodin, se "
                                                "sa ta prishësh."),
        ("Çka do të thotë «nuk shton lloj të ri gabimi»?",
         "Skedari mund të mos kompilohej as më parë, por ndryshimi nuk shtoi gabim tjetër."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 13, "Rezultatet e motorit",
             "Në këtë kapitull sheh sa ndryshime bëri motori, sa erëra hoqi, dhe sa ishin të "
             "pranueshme për një programues.")
    resolution = refactoring["resolution"]
    _list(doc, [
        f"{detected:,} vende ku mund të zbatohej një ndreqje; {applied:,} u ndryshuan "
        f"({_pct(applied / detected)}): " + ", ".join(f"{_words(k)} {v:,}" for k, v in by_type.items())
        + ".",
        f"Era u zhduk në {resolution['resolved']:,} raste ({_pct(f['resolved'])}), mbeti në "
        f"{resolution['persists']:,}, dhe në {resolution.get('unknown', 0):,} nuk u mat dot.",
        f"{sum(f['introduced'].values()):,} erëra të reja u shfaqën: "
        + ", ".join(f"{_words(k)} {v}" for k, v in sorted(f["introduced"].items(), key=lambda p: -p[1]))
        + ". Më shpesh Long Parameter List, sepse metoda e nxjerrë i merr të gjitha hyrjet si "
        "parametra.",
    ])  # fmt: skip
    doc.add_heading("Rishikimi me dorë", level=2)
    doc.add_paragraph(
        f"Autori lexoi {quality['reviewed']} ndryshime, {quality['per_refactoring']} për çdo "
        "ndreqje, dhe për secilin vendosi: i pranueshëm siç është, i pranueshëm pas një "
        "ndreqjeje të vogël, ose i refuzuar."
    )
    rows, weighted = [], []
    for name, entry in quality["by_refactoring"].items():
        accept = entry["acceptance"]
        rows.append([_words(name), f"{entry['population']:,}", str(accept["as_is"]),
                     str(accept["after_edit"]), str(accept["reject"]), _pct(entry["acceptable"], 0),
                     str(entry["behaviour"]["changed"])])  # fmt: skip
        weighted.append(f"{entry['population']:,} × {entry['acceptable']:.2f}")
    _grid(doc, ["Ndreqja", "Ndryshime gjithsej", "Siç është", "Pas ndreqjes", "Refuzuar",
                "E pranueshme", "Sjellja ndryshoi"], rows)  # fmt: skip
    doc.add_paragraph(
        "Pse përqindja e përgjithshme është «e peshuar»: llojet nuk janë njësoj të shpeshta. "
        "Extract Method është shumica e ndryshimeve, ndaj numëron më shumë: "
        f"({' + '.join(weighted)}) / {applied:,} = {_pct(quality['acceptable_reweighted'])}."
    )
    changed = {_words(name): e["behaviour"]["changed"]
               for name, e in quality["by_refactoring"].items() if e["behaviour"]["changed"]}  # fmt: skip
    doc.add_paragraph(
        "Dije këtë: në rishikim, sjellja dukej se ndryshonte te "
        + _dhe([f"{count} {'ndryshim' if count == 1 else 'ndryshime'} {name}"
                for name, count in changed.items()])
        + ". Prandaj punimi thotë «kompilohet» dhe jo «sillet njësoj»."
    )  # fmt: skip
    doc.add_heading("Shembulli i prezantimit", level=2)
    doc.add_paragraph(
        "Te sllajdi 11 i prezantimit është një ndryshim i vërtetë nga projekti Apache Geode, "
        "metoda checkForLimitedOrNoAccess. E gjithë metoda ishte brenda një if-i të madh. Pas "
        "Guard Clauses, kushti bëhet një dalje e hershme me return, dhe pjesa tjetër ngjitet një "
        "nivel lart. Ndryshimi kompilohet, dhe në rishikim ishte i pranueshëm ashtu siç është."
    )
    _idea(doc, "Ndreqja automatike e sigurt është e mundshme, por vetëm për pak raste; Guard "
               "Clauses del më e besueshmja, Introduce Parameter Object më pak e dobishmja.")  # fmt: skip
    _check(doc, [
        ("Sa vende u ndryshuan?", f"{applied:,} nga {detected:,}, pra {_pct(applied / detected)}."),
        ("Pse Introduce Parameter Object merr 0%?",
         "Sepse lista e gjatë e parametrave thjesht kalon në një klasë të re."),
        ("Pse rishikimi është «i peshuar»?", "Sepse Extract Method është shumica e ndryshimeve, "
                                             "ndaj numëron më shumë te totali."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 14, "Mjeti si sistem",
             "Në këtë kapitull mëson nga cilat pjesë përbëhet mjeti, si lidhen, dhe si mbrohet.")
    doc.add_paragraph(
        "Mjeti është ndërtuar si një zinxhir pjesësh. Çdo pjesë përdor vetëm atë që vjen para "
        "saj, kurrë atë që vjen pas. Kështu një ndryshim në ndërfaqe nuk mund ta prishë "
        "matjen, dhe çdo pjesë provohet më vete."
    )
    _grid(doc, ["Pjesa", "Çka bën"], [
        ["parsing", "E lexon kodin me tree-sitter dhe mbledh faktet e sintaksës."],
        ["model", "Të dhëna të thjeshta: klasa, metoda, fusha. Pa logjikë."],
        ["metrics", f"Llogarit {f['metrics']} metrikat, duke e lexuar çdo klasë një herë."],
        ["detectors", "Rregullat. Çdo prag ndodhet në një skedar të vetëm, që të ndryshohet "
                      "lehtë gjatë provave."],
        ["ml", "Modelet e mësimit makinerik."],
        ["refactor", "Motori i ndreqjeve."],
        ["api", "Lidhja mes ndërfaqes dhe mjetit: vetëm kontroll dhe transport, pa logjikë."],
        ["frontend", "Ndërfaqja web, në React."],
        ["glossary", "Fjalori i thjeshtë shqip, një skedar i vetëm për ndërfaqen dhe për CLI-në."],
    ])  # fmt: skip
    _list(doc, [
        "Ndërfaqja web tregon erërat, kushtet që i shkaktuan dhe ndryshimin e propozuar, para se "
        "të preket ndonjë skedar.",
        "Mënyra me gjuhë të thjeshtë, shqip, vjen nga fjalori; e përdor edhe CLI-ja me "
        "--thjeshte, që të dyja të shpjegojnë njësoj.",
        "Mjeti punon vetëm në kompjuterin tënd (localhost). Nuk pranon kërkesa nga faqe të huaja, "
        "dhe çdo dosje që analizohet duhet të jetë brenda një dosjeje të lejuar.",
        "Niset me Nis-JavaSmell.bat ose me Docker. Licenca është MIT.",
        "Ka një suitë të madhe testesh automatike: çdo detektor ka një shembull pozitiv, një "
        "«gati por jo», dhe një rast të veçantë; çdo ndreqje provohet se zbatohet, se refuzon kur "
        "duhet, dhe se rezultati kompilohet.",
    ])  # fmt: skip
    _idea(doc, "Pjesë të ndara, me një drejtim të vetëm varësie: secila provohet më vete.")
    _check(doc, [
        ("Pse pjesët varen vetëm nga ato para tyre?", "Që një ndryshim në një pjesë të mos prishë "
                                                      "pjesët e tjera, dhe secila të provohet "
                                                      "më vete."),
        ("Pse fjalori është një skedar i vetëm?", "Që ndërfaqja dhe CLI-ja ta shpjegojnë çdo term "
                                                  "njësoj."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 15, "Shkenca e mirë: pse rezultatet janë të besueshme",
             "Në këtë kapitull mëson parimet që e bëjnë punimin të verifikueshëm nga dikush tjetër.")
    _list(doc, [
        "I riprodhueshëm: çdo numër prodhohet nga një skript dhe ruhet te data/results/. Një "
        "anëtar i komisionit me një kopje të pastër mund t'i rikrijojë të gjitha tabelat.",
        "Fara të fiksuara dhe versione të fiksuara të librarive: i njëjti ekzekutim jep të njëjtin "
        "rezultat.",
        "Zero kosto: asnjë shërbim me pagesë; gjithçka punon në një laptop pa kartë grafike.",
        "Pa LLM në produkt: ndreqjet janë transformime të pemës, jo tekst i gjeneruar.",
        "Çdo prag ka një burim, ose thuhet hapur që nuk ka.",
        "Asnjë prag nuk ndryshohet që të kalojë një test: vlerat e pritura në teste llogariten "
        "me dorë.",
        "Pa rrjedhje të dhënash: ndarja sipas projektit në çdo provë.",
        "Rezultatet negative raportohen: Feature Envy, refuzimet, erërat e reja, Introduce "
        "Parameter Object.",
        "Vendimet shënohen kur merren, te docs/DECISIONS.md, secili me numrin e vet VD.",
    ])  # fmt: skip
    _idea(doc, "Një rezultat që dikush tjetër nuk e rikrijon dot nuk hyn në punim.")
    _check(doc, [
        ("Çka do të thotë «i riprodhueshëm»?", "Që dikush tjetër, me një kopje të pastër, merr "
                                               "të njëjtat numra me të njëjtat skripte."),
        ("Pse raportohen rezultatet negative?", "Sepse fshehja e tyre do ta bënte punimin më pak "
                                                "të ndershëm dhe më pak të dobishëm."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 16, "Si u ndërtua punimi, hap pas hapi",
             "Në këtë kapitull sheh rrugën e plotë nga projektet e shkarkuara te numrat e "
             "prezantimit, dhe cili skript bën secilin hap.")
    doc.add_paragraph(
        "Çdo hap është një skript te dosja scripts/ dhe e shkruan rezultatin te data/results/. "
        "Hapat varen nga njëri-tjetri me këtë radhë: pa korpus nuk ka mostra, pa mostra nuk ka "
        "vlerësim, pa vlerësim nuk ka intervale."
    )
    _grid(doc, ["Hapi", "Skripti", "Çka bën", "Rezultati"], [
        ["1", "fetch_corpus.py", "Shkarkon projektet që përmend MLCQ, vetëm skedarët .java",
         "korpusi, jashtë git-it"],
        ["2", "report_matching.py", "Kontrollon sa mostra të MLCQ-së gjenden në kod, dhe pse "
         "humbasin të tjerat", "mlcq_matching.json"],
        ["3", "build_dataset.py", "Ndërton tabelën: një rresht për mostër, me metrikat dhe "
         "etiketat", "mlcq_dataset.json"],
        ["4", "evaluate_rules.py", "Qasja A: rregullat mbi mostrat, krahasuar me njerëzit",
         "rules_evaluation.json"],
        ["5", "train_models.py", "Qasja B: modelet, me grupe sipas projektit",
         "ml_evaluation.json"],
        ["6", "bootstrap_intervals.py", "Intervalet e besimit, duke rizgjedhur projektet",
         "bootstrap_intervals.json"],
        ["7", "reviewer_agreement.py", "Sa pajtohen rishikuesit me njëri-tjetrin",
         "reviewer_agreement.json"],
        ["8", "fetch_pmd.py, compare_with_pmd.py", "Krahasimi me PMD-në mbi të njëjtat mostra",
         "pmd_comparison.json"],
        ["9", "sweep_thresholds.py, calibrate_thresholds.py", "Sa ndikojnë pragjet te "
         "rezultatet", "threshold_sweep.json, threshold_calibration.json"],
        ["10", "evaluate_refactorings.py", "Motori mbi gjithë korpusin, me kontroll nga javac",
         "refactoring_evaluation.json"],
        ["11", "verify_with_project.py", "Kompilimi i një mostre ndryshimesh brenda projektit",
         "verify_with_project.json"],
        ["12", "review_rewrites.py", "Rishikimi me dorë i një mostre ndryshimesh",
         "rewrite_quality.json"],
        ["13", "export_system_reference.py, build_figures.py", "Konfigurimi i sistemit dhe "
         "figurat e punimit", "system_reference.json, figurat"],
    ])  # fmt: skip
    doc.add_paragraph(
        "Në fund, ndërtuesit te docs/thesis/ i lexojnë këto skedarë dhe prodhojnë punimin, "
        "prezantimin dhe këtë dokument. Asnjë numër nuk shkruhet me dorë: nëse ndryshon një "
        "rezultat, ndryshon kudo njëkohësisht."
    )
    _idea(doc, "Çdo numër ka një skript që e prodhon dhe një skedar ku ruhet.")
    _check(doc, [
        ("Ku ruhen rezultatet?", "Te data/results/, si skedarë JSON ose CSV."),
        ("Pse nuk shkruhet asnjë numër me dorë?", "Që punimi, prezantimi dhe dokumentet të "
                                                  "thonë gjithmonë të njëjtin numër, dhe që "
                                                  "dikush tjetër t'i rikrijojë."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 17, "Kufizimet dhe puna e ardhshme",
             "Në këtë kapitull mëson çka nuk vërteton punimi, pse, dhe çka do të bëhej më pas.")
    doc.add_paragraph("Formula për çdo kufizim: pranoje → thuaj pse → thuaj çka do të bëje më pas.")
    _grid(doc, ["Kufizimi", "Pse", "Çka do të bëhej"], [
        ["Kompilohet ≠ sillet njësoj", "Korpusi nuk ka skedarë ndërtimi, ndaj testet e "
         "projekteve nuk u ekzekutuan", "Testet e projekteve pas çdo ndryshimi"],
        ["Një bazë, katër erëra, vetëm Java", "Etiketat vijnë vetëm nga MLCQ", "Baza dhe gjuhë të "
         "tjera; tree-sitter i mbështet"],
        ["Modeli më i mirë u zgjodh mbi të njëjtat parashikime", "Krahasimi bëhet mbi të njëjtat "
         "grupe", f"Raportohen të gjithë; edhe më i dobëti fiton te {f['weakest']} nga "
         f"{len(smells)}"],
        ["Disa pragje pa burim", "Literatura nuk jep vlerë", "U provuan me vlera të ndryshme"],
        ["Rishikimin me dorë e bëri vetëm autori", "Punim individual", "Rishikues të pavarur"],
        ["Etiketat e Feature Envy nuk janë të besueshme", "Njerëzit nuk pajtohen",
         "Përkufizim më i qartë ose më shumë rishikues"],
        ["Emrat nuk zgjidhen në gjithë projektin", "Tree-sitter sheh vetëm sintaksën",
         "Një zgjidhës i vogël emrash, për më shumë ndreqje"],
    ])  # fmt: skip
    _check(doc, [
        ("Cili është kufizimi kryesor?", "Kompilimi nuk vërteton që sjellja mbetet e njëjtë, "
                                         "sepse testet e projekteve nuk u ekzekutuan."),
        ("Çka do të bëje më pas?", "Testet e projekteve pas çdo ndryshimi, një zgjidhës i vogël "
                                   "emrash, dhe rishikues të pavarur."),
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 18, "Si lidhet me prezantimin",
             "Në këtë kapitull sheh cili kapitull shpjegon secilin sllajd.")
    _grid(doc, ["Sllajdi", "Kapitujt"], [
        ["1 · Titulli", "–"],
        ["2 · Problemi dhe tri pyetjet", "1, 2"],
        ["3 · Çka shton punimi", "2, 15, 16"],
        ["4 · Katër erërat", "2"],
        ["5 · JavaSmell", "4, 5, 6, 7, 12, 14"],
        ["6 · Të dhënat dhe prova e drejtë", "8, 9"],
        ["7 · Rregullat kundrejt modeleve", "6, 7, 9, 10"],
        ["8 · Pajtimi i njerëzve", "11"],
        ["9 · Motori i refaktorimit", "12, 13"],
        ["10 · Si u verifikuan rishkrimet", "12, 13, 17"],
        ["11 · Shembulli nga Apache Geode", "12, 13"],
        ["12 · Gjetjet dhe çka vjen më pas", "10, 13, 17"],
        ["13 · Referencat", "21"],
    ])  # fmt: skip

    # ------------------------------------------------------------------
    _chapter(doc, 19, "Pyetjet që mund të të bëjnë",
             "Të gjitha pyetjet e mundshme të konferencës, nga më e thjeshta te më e vështira, "
             "me përgjigjen shqip. Përgjigjet në anglisht janë te udhëzuesi i fjalimit.")
    for level, (name, who) in LEVELS.items():
        doc.add_heading(f"Niveli {level}: {name}", level=2)
        paragraph = doc.add_paragraph()
        paragraph.add_run(who).italic = True
        _qa(doc, [(q.pyetja, q.pergjigjja) for q in _question_bank(f) if q.level == level])

    # ------------------------------------------------------------------
    _chapter(doc, 20, "Fjalorth", "Çdo term i punimit, me emrin shqip dhe një shpjegim të shkurtër.")
    _grid(doc, ["Termi", "Shqip", "Shpjegimi"], [list(row) for row in TERMS])

    _chapter(doc, 21, "Burimet kryesore", "Librat dhe artikujt mbi të cilët mbështetet punimi.")
    _list(doc, SOURCES)

    doc.save(BOOK)
    return BOOK


# ----------------------------------------------------------------------
# Dokumenti i pyetjeve të shpjeguara
# ----------------------------------------------------------------------
@dataclass
class DeepQuestion:
    """Një pyetje që nuk mbyllet me një fjali: përgjigjja e shkurtër, dhe shpjegimi pas saj.

    Tabela, kur ka, vihet pas paragrafit që e paralajmëron me dy pika; kodi, pas paragrafit
    të parë, që shembulli të jetë para syve kur lexohet numërimi.
    """

    question: Question
    explanation: list[str]
    table: tuple[list[str], list[list[str]]] | None = None
    code: str = ""


def _deep_questions(f: dict) -> list[DeepQuestion]:
    """Pyetjet që salla bën më shpesh dhe që kërkojnë më shumë se një fjali.

    Shifrat lexohen nga `data/results/`, dhe shembulli i Long Method matet me vetë mjetin,
    si te libri, që dokumenti të mos thotë një numër që mjeti nuk e jep.
    """
    smells, t = f["smells"], f["reference"]["thresholds"]
    limit = t["long_method_loc"]
    sweep_rows = _load("threshold_sweep.json")["per_smell"]["long method"]["long_method_loc"]
    sweep = {row["value"]: row for row in sweep_rows}
    shortest, longest = min(sweep), max(sweep)
    calibration = _load("threshold_calibration.json")["per_smell"]
    chosen = sorted(float(v) for v in calibration["long method"]["chosen"]["long_method_loc"])
    chosen_en, chosen_sq = _and([f"{v:g}" for v in chosen]), _dhe([f"{v:g}" for v in chosen])
    still_ahead = sum(f["best"][s]["mcc"] > calibration[s]["calibrated"]["mcc"] for s in smells)
    intervals = _load("bootstrap_intervals.json")["per_smell"]
    quality = f["quality"]
    as_is = sum(e["acceptance"]["as_is"] for e in quality["by_refactoring"].values())
    after_edit = sum(e["acceptance"]["after_edit"] for e in quality["by_refactoring"].values())
    _, shop_methods, _ = _measured(EXAMPLE_SHOP)
    total = shop_methods["Shop.total"]
    # Metoda total() nis te rreshti me emrin e saj dhe mbaron para «}» së fundit të klasës.
    shop = EXAMPLE_SHOP.splitlines()
    method = shop[next(i for i, line in enumerate(shop) if "int total(" in line) : -1]
    braces = sum(line.strip() == "}" for line in method)
    n = len(smells)
    return [
        DeepQuestion(
            Question(1, "Why do we need this tool at all?", "Pse na duhet fare ky mjet?",
                     "Because smelly code costs time every time it changes, and finding these "
                     "places by hand in thousands of files is slow. Existing tools use fixed "
                     "numbers and fix nothing; this one explains every finding in plain words, "
                     "fixes the safe cases automatically, and is free.",
                     "Sepse kodi me erë kushton kohë sa herë ndryshohet, dhe gjetja e këtyre "
                     "vendeve me dorë në mijëra skedarë është e ngadaltë. Mjetet ekzistuese "
                     "përdorin numra fiksë dhe nuk ndreqin asgjë; ky e shpjegon çdo gjetje me "
                     "fjalë të thjeshta, i ndreq vetë rastet e sigurta, dhe është falas.",
                     "The problem"),
            [
                "Një program mirëmbahet për vite. Sa herë ndryshohet, programuesi duhet ta "
                "lexojë dhe ta kuptojë kodin ekzistues. Kodi me erë e bën këtë më të ngadaltë "
                "dhe më të rrezikshëm, edhe pse programi punon.",
                "Në një projekt me qindra klasa, askush nuk i lexon dot të gjitha për të gjetur "
                "ku janë problemet. Një mjet që i gjen automatikisht ia kursen këtë kohë "
                "programuesit dhe i tregon ku të shikojë së pari, sipas rëndësisë.",
                "Mjetet e sotme kanë tri mangësi që ky punim i mat: përdorin numra fiksë, të "
                f"caktuar mbi kod tjetër, që gjejnë vetëm {min(f['recall']) * 100:.0f}–"
                f"{max(f['recall']) * 100:.0f} nga çdo 100 erëra; nuk shpjegojnë pse një gjetje "
                "ka rëndësi; dhe nuk ndreqin asgjë. JavaSmell e shpjegon çdo gjetje me numra dhe "
                "me gjuhë të thjeshtë, mëson kufijtë nga njerëzit, dhe ndreq vetë rastet që mund "
                "t'i vërtetojë si të sigurta.",
            ],
        ),
        DeepQuestion(
            Question(1, "Couldn't an AI like ChatGPT do all of this?",
                     "A nuk mund ta bënte gjithë këtë një AI si ChatGPT?",
                     "It can give suggestions, but not the same guarantees. Its answer can change "
                     "each time, it cannot prove that a change is safe, your code would leave your "
                     "computer, and it usually costs money. This tool gives the same answer every "
                     "time, explains it with numbers, and runs offline for free. A good next step "
                     "would be to let an AI suggest and let this engine check.",
                     "Mund të japë sugjerime, por jo të njëjtat garanci. Përgjigjja e saj mund të "
                     "ndryshojë çdo herë, nuk e vërteton dot që një ndryshim është i sigurt, kodi "
                     "yt do të dilte nga kompjuteri yt, dhe zakonisht kushton. Ky mjet jep të "
                     "njëjtën përgjigje çdo herë, e shpjegon me numra, dhe punon pa internet e "
                     "falas. Një hap i mirë më pas: AI të sugjerojë, dhe ky motor të kontrollojë.",
                     "What we built"),
            [
                "Një model gjuhësor (LLM) si ChatGPT mund ta lexojë një metodë dhe të thotë «kjo "
                "është shumë e gjatë», ose të propozojë një version më të pastër. Kjo është e "
                "dobishme, por për një punim shkencor dhe për një mjet që e rishkruan kodin e "
                "dikujt ka katër probleme.",
                "E para, përsëritshmëria: e njëjta pyetje mund të marrë përgjigje të ndryshme në "
                "dy ekzekutime. Një rezultat që nuk përsëritet nuk mund të raportohet si numër "
                "në punim.",
                "E dyta, siguria: LLM-ja shkruan kod që duket i saktë, por nuk e vërteton që "
                "sjellja mbetet e njëjtë. Motori i këtij punimi ndryshon vetëm kur kushtet "
                "vërtetohen nga pema e sintaksës, dhe pastaj e kontrollon me javac.",
                "E treta, privatësia dhe kostoja: kodi i një kompanie do t'i dërgohej një "
                "shërbimi në internet, dhe çdo analizë kushton. Punimi kishte kusht zero kosto "
                "dhe punë pa internet.",
                "E katërta, shpjegimi: rregullat e punimit thonë pse, me numra: çdo kusht, si "
                f"«MLOC > {limit:g}», tregohet me vlerën e matur pranë. Një LLM jep një arsye "
                "me fjalë, që nuk mund të kontrollohet njësoj.",
                "Kjo nuk do të thotë që AI nuk ka vend. Një kombinim i arsyeshëm për të ardhmen: "
                "AI propozon një ndryshim, dhe motori i këtij punimi e kontrollon para se të "
                "pranohet. Vendimi për të mos përdorur LLM në produkt është VD-03 te "
                "docs/DECISIONS.md.",
            ],
        ),
        DeepQuestion(
            Question(1, "Why Java, and not another language?", "Pse Java, dhe jo një gjuhë tjetër?",
                     "Three reasons. The only large dataset of smells judged by professional "
                     "developers, MLCQ, is Java code. The classic smell rules were written for "
                     "object-oriented languages like Java. And Java has a strict compiler, javac, "
                     "that I use to check every change. The parser supports other languages, so "
                     "the approach could be extended later.",
                     "Tri arsye. E vetmja bazë e madhe me erëra të gjykuara nga programues "
                     "profesionistë, MLCQ, është kod Java. Rregullat klasike u shkruan për gjuhë "
                     "të orientuara nga objektet, si Java. Dhe Java ka një kompajler të rreptë, "
                     "javac, me të cilin kontrolloj çdo ndryshim. Parseri i mbështet edhe gjuhë të "
                     "tjera, ndaj qasja mund të zgjerohet më vonë.", "Data and evaluation"),
            [
                "Arsyeja e parë është e dhëna. Për të matur sa i saktë është një detektor, duhen "
                "përgjigje nga njerëz. Në MLCQ, 26 programues profesionistë dhanë rreth 15 mijë "
                "vlerësime, dhe të gjitha janë për kod Java. Për gjuhë të tjera nuk ka një bazë "
                "të tillë me këtë madhësi.",
                "Arsyeja e dytë janë rregullat. Libri i Lanza-s dhe Marinescu-t dhe katalogu i "
                "Fowler-it i përshkruajnë erërat me klasa, metoda dhe fusha: koncepte të "
                "gjuhëve si Java. Në një gjuhë pa klasa, God Class ose Data Class nuk kanë të "
                "njëjtin kuptim.",
                "Arsyeja e tretë është kontrolli. Java ka tipe të rrepta dhe një kompajler që e "
                "refuzon kodin e gabuar. Kjo e bën të mundur kontrollin e çdo ndryshimi të "
                "motorit. Në një gjuhë pa kompajler të tillë, si Python, ky kontroll do të ishte "
                "shumë më i dobët.",
                "Parseri tree-sitter i mbështet shumë gjuhë, ndaj leximi i kodit mund të "
                "zgjerohet. Metrikat dhe motori i ndreqjeve janë shkruar për Java, dhe për një "
                "gjuhë tjetër do të duheshin edhe të dhëna të reja me etiketa.",
            ],
        ),
        DeepQuestion(
            Question(1, "Who would use this tool?", "Kush do ta përdorte këtë mjet?",
                     "A developer or a team before changing old code, to see where the problems "
                     "are and fix the simple ones. And students or teachers, thanks to the "
                     "plain-language mode in Albanian.",
                     "Një programues ose një ekip para se të ndryshojë kod të vjetër, që të shohë "
                     "ku janë problemet dhe t'i ndreqë të thjeshtat. Dhe studentë ose mësues, "
                     "falë mënyrës me gjuhë të thjeshtë shqip.", "What we built"),
            [
                "Rasti kryesor: një ekip merr një projekt të vjetër dhe duhet të shtojë diçka. "
                "Mjeti i tregon cilat klasa dhe metoda janë më të rrezikshme, sipas rëndësisë, "
                "dhe pse. Për disa raste përgatit vetë ndryshimin, që programuesi ta shikojë dhe "
                "ta pranojë.",
                "Rasti i dytë: mësimi. Mënyra me gjuhë të thjeshtë e shpjegon çdo erë dhe çdo "
                "metrikë shqip, pa zhargon. Një student mund ta përdorë mbi kodin e vet për të "
                "parë çka nuk shkon dhe pse.",
                "Mjeti punon në kompjuterin e përdoruesit, pa internet dhe falas, ndaj mund ta "
                "përdorë edhe një kompani që nuk do ta nxjerrë kodin jashtë.",
            ],
        ),
        DeepQuestion(
            Question(1, "Is the tool better than a human reviewer?",
                     "A është mjeti më i mirë se një rishikues njeri?",
                     "It is not meant to replace one. It shows where to look and explains why, so "
                     "a developer can decide faster. For the fixes, a developer still reviews the "
                     "change before it is applied.",
                     "Nuk është menduar ta zëvendësojë. Tregon ku të shikosh dhe pse, që "
                     "programuesi të vendosë më shpejt. Edhe për ndreqjet, një programues e shikon "
                     "ndryshimin para se të zbatohet.", "What we built"),
            [
                "Mjeti mëson nga njerëzit dhe krahasohet me ta, ndaj nuk ka kuptim të thuhet se "
                "është «më i mirë» se ata. Dobia e tij është shpejtësia dhe qëndrueshmëria: e "
                "kontrollon gjithë projektin në sekonda, dhe e gjykon çdo klasë me të njëjtin "
                "rregull.",
                "Për ndreqjet, mjeti përgatit një skedar ndryshimesh dhe nuk e prek kodin vetë. "
                f"Në rishikimin me dorë të {quality['reviewed']} ndryshimeve, vetëm {as_is} u "
                f"pranuan ashtu siç ishin, dhe {after_edit} pas një ndreqjeje të vogël, si një "
                "emër më i mirë për metodën e re. Prandaj fjala e fundit mbetet te programuesi.",
            ],
        ),
        DeepQuestion(
            Question(1, "How does the tool decide that a method is a Long Method?",
                     "Si vendos mjeti që një metodë është Long Method?",
                     f"It counts the lines of the method that contain code. The first line, with "
                     f"the method's name, counts; empty lines, comments and lines with only a "
                     f"closing brace do not. If there are more than {limit:g} such lines, it is a "
                     "Long Method.",
                     f"Numëron rreshtat e metodës që kanë kod. Rreshti i parë, me emrin e metodës, "
                     f"numërohet; rreshtat bosh, komentet dhe rreshtat me vetëm një «}}» jo. Nëse "
                     f"ka më shumë se {limit:g} rreshta të tillë, është Long Method.", "The problem"),
            [
                "Metrika quhet MLOC (rreshtat e kodit të metodës). Ajo nuk numëron rreshtat e "
                "ekranit, por rreshtat që kanë një udhëzim: kështu një metodë nuk bëhet «e gjatë» "
                "vetëm sepse ka shumë rreshta bosh ose komente. Shembull:",
                f"Metoda total() ka {len(method)} rreshta në ekran, por MLOC = {total['MLOC']:g}. "
                "Numërohen emri i metodës, int sum, for, if, sum += p, if-i i dytë, sum = …, dhe "
                f"return. {braces} rreshtat me vetëm «}}» nuk numërohen. Këtë vlerë e jep vetë "
                "mjeti kur e mat këtë kod.",
                f"Rregulli është i thjeshtë: MLOC > {limit:g}. Pse pikërisht {limit:g}, e "
                "shpjegon pyetja e Nivelit 3 më poshtë.",
            ],
            code=EXAMPLE_SHOP,
        ),
        DeepQuestion(
            Question(2, "How does the tool decide that a class is a Blob or a Data Class?",
                     "Si vendos mjeti që një klasë është Blob ose Data Class?",
                     f"With rules from Lanza and Marinescu. A Blob does a lot of work (WMC at least "
                     f"{t['god_class_wmc']}), its methods do not share data (TCC below one third), "
                     f"and it uses data of other classes (ATFD above {t['god_class_atfd']}). A Data "
                     "Class has many public fields or get and set methods, and almost no real "
                     "work.",
                     f"Me rregullat e Lanza-s dhe Marinescu-t. Një Blob bën shumë punë (WMC së "
                     f"paku {t['god_class_wmc']}), metodat e saj nuk ndajnë të dhëna (TCC nën një "
                     f"të tretën), dhe përdor të dhëna të klasave të tjera (ATFD mbi "
                     f"{t['god_class_atfd']}). Një Data Class ka shumë fusha publike ose metoda "
                     "get e set, dhe pothuajse asnjë punë të vërtetë.", "Four smells"),
            [
                "Rregulli i Blob-it (God Class) ka tri kushte, dhe duhet të plotësohen të tria. "
                f"WMC ≥ {t['god_class_wmc']}: shuma e kompleksitetit të metodave është shumë e "
                "lartë, pra klasa bën shumë punë. TCC < 1/3: më pak se një e treta e çifteve të "
                "metodave përdorin fusha të përbashkëta, pra pjesët e klasës nuk lidhen me "
                f"njëra-tjetrën. ATFD > {t['god_class_atfd']}: klasa lexon shumë të dhëna nga "
                "klasa të tjera.",
                "Rregulli i Data Class ka dy pjesë. WOC < 1/3: më pak se një e treta e pjesës "
                "publike të klasës bën punë të vërtetë; pjesa tjetër vetëm jep ose vendos "
                f"vlera. Dhe: ose klasa ka më shumë se {t['data_class_public_members']} fusha "
                f"publike dhe metoda get/set me pak logjikë (WMC < {t['data_class_wmc_low']}), "
                f"ose më shumë se {t['data_class_public_members_many']} me pak më shumë logjikë "
                f"(WMC < {t['data_class_wmc_high']}).",
                "Mjeti nuk thotë vetëm «Blob»: tregon secilin kusht me vlerën e matur, që "
                "programuesi ta shohë vetë pse.",
            ],
        ),
        DeepQuestion(
            Question(2, "What did you build yourself, and what did you reuse?",
                     "Çka ndërtove vetë, dhe çka ripërdore?",
                     "I wrote the analysis, the metrics, the rules, the training and the "
                     "evaluation, the refactoring engine and the interface. I reused the "
                     "tree-sitter parser, the scikit-learn library for the models, the MLCQ "
                     "labels, the Java compiler, and PMD only for comparison.",
                     "Shkrova vetë analizën, metrikat, rregullat, trajnimin dhe vlerësimin, motorin "
                     "e ndreqjeve dhe ndërfaqen. Ripërdora parserin tree-sitter, librarinë "
                     "scikit-learn për modelet, etiketat e MLCQ-së, kompajlerin e Java-s, dhe "
                     "PMD-në vetëm për krahasim."),
            [
                "Ripërdorimi është i qëllimshëm: një parser, një librari modelesh dhe një "
                "kompajler janë mjete të provuara nga mijëra njerëz, dhe rishkrimi i tyre do "
                "të sillte gabime pa fitim. Etiketat e MLCQ-së janë ripërdorur sepse etiketimi "
                "nga vetë autori do ta bënte provën të njëanshme.",
                "Pjesa e re e punimit është mënyra si lidhen: metrikat e llogaritura nga faktet "
                "e sintaksës, rregullat që kthejnë kushtet me vlerat, krahasimi i drejtë me "
                "ndarje sipas projektit, dhe motori që refuzon kur nuk vërteton dot sigurinë.",
            ],
        ),
        DeepQuestion(
            Question(2, "What is the biggest weakness of the work?",
                     "Cila është dobësia më e madhe e punimit?",
                     "That compiling is not the same as behaving the same. The projects come "
                     "without their build files, so I could not run their own tests after each "
                     "change. That is the first thing I would add.",
                     "Që kompilimi nuk është njësoj si sjellja e njëjtë. Projektet vijnë pa "
                     "skedarët e ndërtimit, ndaj nuk munda t'i ekzekutoj testet e tyre pas çdo "
                     "ndryshimi. Kjo është gjëja e parë që do ta shtoja.",
                     "Conclusions"),
            [
                "Një ndryshim që kompilohet mund ta ndryshojë prapë sjelljen, p.sh. nëse e "
                "ndryshon radhën e dy veprimeve. Testet e projektit do ta kapnin këtë, por "
                "korpusi mban vetëm skedarët .java, pa skedarët që e ndërtojnë projektin dhe pa "
                "libraritë.",
                "Punimi nuk e fsheh: thotë «kompilohet, ose nuk shton lloj të ri gabimi», jo "
                "«sillet njësoj». Në rishikimin me dorë, një ndryshim Extract Method dukej se e "
                "ndryshonte sjelljen, dhe pikërisht kjo e tregon pse kufizimi ka rëndësi.",
                "Dobësi të tjera, më të vogla: një bazë të dhënash, katër erëra, vetëm Java, dhe "
                "rishikimi me dorë i bërë vetëm nga autori.",
            ],
        ),
        DeepQuestion(
            Question(3, f"Why exactly {limit:g} lines? How did you arrive at that limit?",
                     f"Pse saktësisht {limit:g} rreshta? Si arrite te ky kufi?",
                     f"No published source gives a number, so {limit:g} is my choice, and the "
                     f"thesis says so. Then I measured what happens when it moves, from "
                     f"{shortest:g} to {longest:g} lines: the MCC goes from "
                     f"{sweep[shortest]['mcc']:.2f} down to {sweep[longest]['mcc']:.2f}. When I "
                     f"calibrate it inside each training group, the best value is {chosen_en} "
                     f"lines: the developers in MLCQ call a method long earlier than {limit:g}.",
                     f"Asnjë burim i botuar nuk jep një numër, ndaj {limit:g} është zgjedhja ime, "
                     f"dhe punimi e thotë hapur. Pastaj mata çka ndodh kur lëviz, nga "
                     f"{shortest:g} deri në {longest:g} rreshta: MCC-ja bie nga "
                     f"{sweep[shortest]['mcc']:.2f} në {sweep[longest]['mcc']:.2f}. Kur e "
                     f"kalibroj brenda çdo grupi mësimi, vlera më e mirë del {chosen_sq} rreshta: "
                     f"programuesit e MLCQ-së e quajnë një metodë të gjatë më herët se "
                     f"{limit:g}.", "The problem"),
            [
                "Fowler e përshkruan Long Method pa dhënë një numër, dhe mjetet e ndryshme "
                f"përdorin vlera të ndryshme. Punimi zgjodhi {limit:g} rreshta dhe e thotë hapur "
                "që kjo është zgjedhje e autorit, jo një vlerë nga literatura (vendimi VD-133).",
                "Një zgjedhje e tillë nuk mund të mbetet pa u matur. Prandaj kufiri u provua me "
                f"{len(sweep_rows)} vlera, nga gjysma deri në dyfishin e tij, dhe për secilën u "
                "mat sa metoda shënon dhe sa pajtohet me njerëzit:",
                "Kufiri i ulët kap shumicën e metodave të gjata (recall i lartë), por shënon edhe "
                "disa që njerëzit nuk i quajnë të gjata (precision më i ulët). Kufiri i lartë "
                "bën të kundërtën: shënon pak, por ato pak janë pothuajse gjithmonë të sakta.",
                f"Kur kufiri zgjidhet vetëm mbi grupet e mësimit, dhe provohet te grupi që nuk u "
                f"pa, vlera më e mirë del {chosen_sq} rreshta. Përfundimi: programuesit e "
                f"MLCQ-së e quajnë një metodë «të gjatë» më herët se {limit:g} rreshta.",
            ],
            (["Kufiri (rreshta)", "Sa metoda shënon", "Precision", "Recall", "MCC"],
             [[f"{row['value']:g}", f"{row['fired']}", f"{row['precision']:.0%}",
               f"{row['recall']:.0%}", f"{row['mcc']:.2f}"] for row in sweep_rows]),
        ),
        DeepQuestion(
            Question(3, "How reliable are your results?", "Sa të besueshme janë rezultatet e tua?",
                     f"As reliable as a fair test allows, and I say where they are not. Every "
                     f"model is tested on projects it has never seen, every difference is checked "
                     f"with {f['resamples']:,} resamples, and every number can be recreated by a "
                     "script. The limits are the labels themselves for Feature Envy, one dataset, "
                     "and that compiling is not the same as behaving the same.",
                     f"Aq të besueshme sa e lejon një provë e drejtë, dhe them hapur ku nuk janë. "
                     f"Çdo model provohet mbi projekte që nuk i ka parë kurrë, çdo ndryshim "
                     f"kontrollohet me {f['resamples']:,} rizgjedhje, dhe çdo numër rikrijohet me "
                     "një skript. Kufijtë janë vetë etiketat te Feature Envy, një bazë e vetme "
                     "të dhënash, dhe që kompilimi nuk është njësoj si sjellja e njëjtë.",
                     "Rules vs. machine learning", "Confidence intervals"),
            [
                "Punimi ka katër mbrojtje kundër rezultateve të fryra. E para: etiketat vijnë nga "
                "26 programues profesionistë, jo nga autori. E dyta: projektet ndahen në "
                f"{f['ml']['folds']} grupe, dhe asnjë model nuk provohet mbi një projekt nga i "
                "cili mësoi. E treta: nota është MCC, që nuk mashtrohet kur erërat janë të "
                f"rralla. E katërta: çdo ndryshim model − rregulla rillogaritet {f['resamples']:,} "
                "herë mbi projekte të rizgjedhura.",
                "Ja intervalet 95% të ndryshimit model − rregulla. Kur i gjithë intervali është "
                "mbi zero, përmirësimi nuk është fat:",
                "Ku rezultatet janë më pak të forta, punimi e thotë: te Feature Envy edhe "
                "njerëzit nuk pajtohen, ndaj as etiketat nuk janë të sigurta; modeli më i mirë "
                "u zgjodh mbi të njëjtat parashikime, ndaj është pak optimist (por edhe më i "
                f"dobëti i mund rregullat te {f['weakest']} nga {n}); dhe ndreqjet kontrollohen "
                "me kompajlerin, jo me testet e projekteve.",
            ],
            (["Era", "Ndryshimi model − rregulla, 95%", "Mbi zero?"],
             [[SMELL_SQ[s], "{low:+.2f} … {high:+.2f}".format(
                 **intervals[s]["intervals"]["difference_model_minus_rules"]),
               "po" if intervals[s]["intervals"]["difference_model_minus_rules"]["low"] > 0
               else "jo"] for s in smells]),
        ),
        DeepQuestion(
            Question(3, "Why these three refactorings, and not others?",
                     "Pse këto tri ndreqje, dhe jo të tjera?",
                     "Because they can be proven safe from the code of one file. Others, like "
                     "Move Method or Extract Class, need every place in the project that uses the "
                     "moved code, and the tool cannot find those places with certainty.",
                     "Sepse mund të vërtetohen si të sigurta nga kodi i një skedari. Të tjerat, si "
                     "Move Method ose Extract Class, kërkojnë çdo vend në projekt që e përdor "
                     "kodin e zhvendosur, dhe mjeti nuk i gjen dot me siguri ato vende.",
                     "Automatic refactoring"),
            [
                "Extract Method, Guard Clauses dhe Introduce Parameter Object ndryshojnë vetëm "
                "brenda një metode ose rreth saj. Kushtet e tyre mund të kontrollohen nga pema e "
                "sintaksës: a del një return nga blloku, a ka vlera që nuk janë caktuar, a ndryshon "
                "pjesa më shumë se një variabël.",
                "Move Method ose Extract Class e zhvendosin kodin në një klasë tjetër. Për t'i bërë "
                "të sigurta, duhet gjetur çdo vend në gjithë projektin që e thërret metodën ose e "
                "përdor fushën. Kjo kërkon zgjidhjen e emrave në gjithë projektin, që mjeti me "
                "qëllim nuk e bën. Prandaj ato vetëm propozohen.",
                "Ky është parimi i motorit: më mirë një ndreqje më pak, se sa një ndreqje që "
                "prish kodin.",
            ],
        ),
        DeepQuestion(
            Question(4, "If you change the thresholds, does your conclusion change?",
                     "Nëse i ndryshon pragjet, a ndryshon përfundimi yt?",
                     f"The rules change a lot: for Long Method the MCC moves between "
                     f"{sweep[longest]['mcc']:.2f} and {sweep[shortest]['mcc']:.2f}. But even when "
                     f"I calibrate the rules' thresholds inside each training group, the models "
                     f"stay ahead on {_count_word(still_ahead, n)} smells, so the conclusion "
                     "holds. The rule catches up only for Long Method, and only if I pick the "
                     "limit after looking at the test data.",
                     f"Rregullat ndryshojnë shumë: për Long Method MCC-ja lëviz mes "
                     f"{sweep[longest]['mcc']:.2f} dhe {sweep[shortest]['mcc']:.2f}. Por edhe kur "
                     f"i kalibroj pragjet e rregullave brenda çdo grupi mësimi, modelet mbeten "
                     f"përpara te {still_ahead} nga {n} erërat, ndaj përfundimi qëndron. Rregulla "
                     "e arrin modelin vetëm te Long Method, dhe vetëm nëse kufirin e zgjedh pasi "
                     "i shoh të dhënat e provës.",
                     "Rules vs. machine learning"),
            [
                "Kjo është pyetja më e fortë që mund të bëjë një recensent: «modelet fitojnë vetëm "
                "sepse rregullat kanë pragje të këqija?». Për t'u përgjigjur, punimi i kalibroi "
                "pragjet e rregullave mbi MLCQ, pa e shkelur ndarjen: vlera zgjidhet vetëm mbi "
                "grupet e mësimit dhe provohet te grupi që nuk u pa.",
                "Tabela krahason rregullat me pragjet e botuara, rregullat me pragjet e "
                "kalibruara, dhe modelin më të mirë:",
                "Kalibrimi i ndihmon rregullat, sidomos te Long Method, por modeli mbetet përpara "
                f"te {still_ahead} nga {n} erërat. Pra përparësia e mësimit makinerik nuk vjen "
                "vetëm nga pragjet e këqija.",
                f"Një kujdes që duhet thënë hapur: te Long Method, me kufirin {shortest:g} "
                f"rregulla arrin MCC {sweep[shortest]['mcc']:.3f}, dhe modeli "
                f"{f['best']['long method']['mcc']:.3f}, pra janë të barabartë. Por {shortest:g} "
                "u gjet duke i parë të gjitha të dhënat, edhe ato të provës. Kur kufiri zgjidhet "
                "vetëm mbi grupet e mësimit, rregulla bie në "
                f"{calibration['long method']['calibrated']['mcc']:.3f}. Krahasimi i drejtë është "
                "kolona e kalibruar.",
            ],
            (["Era", "Rregullat, pragje të botuara", "Rregullat, pragje të kalibruara",
              "Modeli më i mirë"],
             [[SMELL_SQ[s], f"{calibration[s]['published']['mcc']:.2f}",
               f"{calibration[s]['calibrated']['mcc']:.2f}", f"{f['best'][s]['mcc']:.2f}"]
              for s in smells]),
        ),
        DeepQuestion(
            Question(4, "Why not just tune the rules' thresholds on your data?",
                     "Pse thjesht nuk i akordove pragjet e rregullave mbi të dhënat e tua?",
                     "Because Approach A measures the published rules as people actually use them. "
                     "Tuning them on the same data I test on would be tuning to the test. I report "
                     "the calibration separately, done only on the training groups.",
                     "Sepse qasja A i mat rregullat e botuara ashtu siç i përdorin njerëzit. T'i "
                     "akordoja mbi të njëjtat të dhëna me të cilat provoj, do të ishte akordim "
                     "sipas provës. Kalibrimin e raportoj veçmas, të bërë vetëm mbi grupet e "
                     "mësimit.", "Rules vs. machine learning"),
            [
                "Pyetja e parë e punimit (Q1) është: sa të sakta janë rregullat e botuara? Për "
                "t'iu përgjigjur, rregullat duhet të maten me vlerat që përdorin njerëzit, jo me "
                "vlera të zgjedhura që të dalin mirë mbi këto të dhëna.",
                "Nëse pragjet do të zgjidheshin mbi të gjitha të dhënat dhe pastaj do të "
                "provoheshin mbi të njëjtat, rezultati do të ishte i fryrë: do të ishte si një "
                "student që i sheh pyetjet para provimit. Prandaj kalibrimi bëhet vetëm mbi "
                "grupet e mësimit, dhe raportohet si rezultat më vete.",
            ],
        ),
    ]  # fmt: skip


def _explanation(doc, entry: DeepQuestion) -> None:
    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = DocPt(2)
    run = heading.add_run("Shpjegimi")
    run.bold = True
    run.font.color.rgb = DocRGB(*ACCENT)
    for index, text in enumerate(entry.explanation):
        doc.add_paragraph(text)
        if entry.code and index == 0:
            _code(doc, entry.code.splitlines())
        if entry.table and text.endswith(":"):
            _grid(doc, *entry.table)
    doc.add_paragraph()


def build_questions(f: dict, backups: list[str], main: int) -> str:
    """Pyetjet më të shpeshta me shpjegim të gjerë, dhe pas tyre të gjitha të tjerat.

    Një dokument më vete, që udhëzuesi i fjalimit dhe libri të mbeten siç janë provuar.
    """
    deep = _deep_questions(f)
    doc = _document()
    doc.add_heading("Pyetjet e konferencës, të shpjeguara", level=0)
    doc.add_paragraph(f"{TITLE}\n{AUTHORS}")
    _list(doc, [
        "Pjesa e parë ka pyetjet që salla bën më shpesh dhe që nuk mbyllen me një fjali: pse na "
        "duhet mjeti, pse jo me AI, pse Java, si vendoset Long Method dhe pse "
        f"{f['reference']['thresholds']['long_method_loc']:g} rreshta, sa të besueshme janë "
        "rezultatet.",
        "Për secilën: pyetja anglisht dhe shqip, përgjigjja e shkurtër që e thua në sallë (me "
        "të zeza anglisht, me gri shqip), dhe pastaj «Shpjegimi», që e lexon ti për ta kuptuar "
        "deri në fund.",
        "Pjesa e dytë ka pyetjet e udhëzuesit të fjalimit, sipas nivelit, me përgjigjen e "
        "shkurtër, që të gjitha të jenë në një vend.",
        "Në sallë thuaj vetëm përgjigjen e shkurtër. Nëse dikush pyet më tej, shpjegimi të jep "
        "çka të shtosh.",
    ])  # fmt: skip
    doc.add_heading("Pjesa e parë: pyetjet që kërkojnë shpjegim", level=1)
    for level, (name, _) in LEVELS.items():
        entries = [entry for entry in deep if entry.question.level == level]
        if entries:
            doc.add_heading(f"Niveli {level}: {name}", level=2)
        for entry in entries:
            _answer(doc, entry.question, backups, main)
            _explanation(doc, entry)
    doc.add_page_break()
    doc.add_heading("Pjesa e dytë: pyetjet e udhëzuesit të fjalimit, sipas nivelit", level=1)
    bank = _question_bank(f)
    for level, (name, who) in LEVELS.items():
        doc.add_heading(f"Niveli {level}: {name}", level=2)
        paragraph = doc.add_paragraph()
        paragraph.add_run(who).italic = True
        for question in bank:
            if question.level == level:
                _answer(doc, question, backups, main)
    doc.save(QUESTIONS)
    return QUESTIONS


def build_spoken_text(script: list[tuple[str, list[Line]]], long: bool) -> str:
    """Vetëm fjalitë angleze që thuhen, sllajd për sllajd, pa përkthim e pa udhëzime.

    Për t'u mbajtur në dorë ose për t'u lexuar në provë: udhëzuesi i fjalimit ka gjithçka
    tjetër. Çdo fjali në rresht të vet, që syri ta gjejë vendin pas një pauze, dhe çdo
    sllajd i mbajtur në një faqe.
    """
    doc = _document()
    words = sum(_word_count(_spoken(parts, long)) for _, parts in script)
    version = "9-minute version" if long else "short version"
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"UBT 2026 · {version} · about "
                            f"{_clock(round(words / WORDS_PER_MINUTE * 60))} at a calm pace")  # fmt: skip
    run.font.size = DocPt(10)
    run.font.color.rgb = DocRGB(*SOFT)
    for index, (key, parts) in enumerate(script, 1):
        lines = _spoken(parts, long)
        if not lines:
            continue
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = DocPt(14)
        heading.paragraph_format.space_after = DocPt(4)
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(f"Slide {index} · {key}")
        run.bold = True
        run.font.size = DocPt(11)
        run.font.color.rgb = DocRGB(*ACCENT)
        for position, line in enumerate(lines):
            sentence = doc.add_paragraph()
            sentence.paragraph_format.space_after = DocPt(5)
            sentence.paragraph_format.keep_with_next = position < len(lines) - 1
            sentence.add_run(str(line)).font.size = DocPt(14)
    path = SPOKEN_LONG if long else SPOKEN
    doc.save(path)
    return path


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    facts = _facts()
    for long in (False, True):
        deck, script, backups = build_deck(facts, long)
        words = sum(_word_count(_spoken(parts, long)) for _, parts in script)
        print(f"U gjenerua: {deck}")
        print(f"  Fjalë në tekstin folës: {words}, rreth {words / WORDS_PER_MINUTE:.1f} minuta "
              f"me ritëm të qetë, {words / FAST_WORDS_PER_MINUTE:.1f} me të shpejtë")  # fmt: skip
    for path in (
        build_sheet(facts, script),
        build_speech_guide(facts, script, backups),
        build_book(facts),
        build_questions(facts, backups, len(script)),
        build_spoken_text(script, long=False),
        build_spoken_text(script, long=True),
    ):
        print(f"U gjenerua: {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
